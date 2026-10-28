"""Geometria, ancoraggi, carichi, elementi e squadra del ponteggio; PiMUS con disegni e allegati."""

import json
import sqlite3
from datetime import date

import pytest
from docx import Document

from sicurcantiere import pimus
from sicurcantiere import ponteggio as P
from sicurcantiere.db import MIGRAZIONI, Progetto
from sicurcantiere.repository import Repository


@pytest.fixture
def repo(tmp_path):
    progetto = Progetto(tmp_path / "p.spm", crea=True)
    yield Repository(progetto)
    progetto.chiudi()


def _ponteggio(**geometria):
    g = {"lati": [{"nome": "Nord", "lunghezza": 18}, {"nome": "Est", "lunghezza": 9, "altezza": 8}], **geometria}
    return {"id": 1, "nome": "Facciate", "altezza": 12, "lunghezza": 0, "tipo": "Telai prefabbricati",
            "classe_carico": P.CLASSI_CARICO[3][0], "conforme_schemi": 1, "autorizzazione": "x", "impresa_id": 1,
            "preposto": "", "progettista": "", "geometria": json.dumps(g)}


def test_geometria_predefinita_e_lati():
    p = _ponteggio()
    g = P.geometria(p)
    assert g["campata"] == 1.8 and g["interpiano"] == 2.0 and len(g["tratti"]) == 2  # i «lati» diventano tratti
    nord, est = P.lati(p, g)
    assert (nord.campate, nord.livelli) == (10, 6)  # 18 / 1,80 e 12 / 2,00
    assert (est.campate, est.livelli, est.altezza) == (5, 4, 8)
    assert P.sviluppo(g) == 27
    # senza geometria: un lato con lo sviluppo e il nome del ponteggio
    vecchio = P.geometria({"nome": "Sud", "lunghezza": 14, "geometria": "{}"})
    assert vecchio["tratti"][0]["nome"] == "Sud" and vecchio["tratti"][0]["lunghezza"] == 14
    assert P.geometria({"geometria": "non json"})["campata"] == 1.8


def test_ancoraggi_schema_e_modifiche_a_mano():
    p = _ponteggio()
    g = P.geometria(p)
    nord = P.lati(p, g)[0]
    schema = P.schema_ancoraggi(nord, g)
    assert (0, 2) in schema and (10, 2) in schema  # testate
    assert (1, 4) in schema and (1, 2) not in schema  # file sfalsate a rombo
    assert all(l in (2, 4, 6) for _, l in schema)  # ogni 2 livelli, sempre l'ultimo
    g = P.cambia_ancoraggio(g, p, 0, 3, 1)  # aggiunto
    g = P.cambia_ancoraggio(g, p, 0, 0, 2)  # tolto
    p["geometria"] = P.testo_geometria(g)
    g = P.geometria(p)
    nord = P.lati(p, g)[0]
    posizioni = P.posizioni_ancoraggi(nord, g)
    assert (3, 1) in posizioni and (0, 2) not in posizioni
    g = P.cambia_ancoraggio(g, p, 0, 0, 2)  # rimesso
    assert g["tratti"][0]["tolti"] == []
    a = P.ancoraggi(p)
    assert a.numero > 10 and 0 < a.area_per_ancoraggio <= P.AREA_PER_ANCORAGGIO


def test_carico_sul_piano_di_appoggio():
    p = _ponteggio(terreno=100, area_appoggio=0.15, peso_proprio=0.3)
    s = P.appoggio(p)
    # peso 0,30 × 1,80 × 13 / 2 = 3,51; servizio 3,00 × 1,80 × 1,05 / 2 × 1,5 = 4,25
    assert s.carico_montante == pytest.approx(7.76, abs=0.01)
    assert s.pressione == pytest.approx(51.8, abs=0.1) and s.verificato
    debole = _ponteggio(terreno=20, area_appoggio=0.05)
    assert not P.appoggio(debole).verificato
    assert any("Piano di appoggio" in m for m in P.avvisi_geometria(debole))
    assert P.sovraccarico("Classe 2") == 1.5 and P.sovraccarico("boh") == 0


def test_elementi_per_tipologia():
    p = _ponteggio()
    voci = dict((n, q) for n, q, _ in P.elementi(p))
    assert voci["Telai (stilate × livelli)"] == 11 * 6 + 6 * 4
    assert voci["Basette regolabili"] == (11 + 6) * 2
    assert next(q for n, q in voci.items() if n.startswith("Ancoraggi (")) == P.ancoraggi(p).numero
    assert voci["Spinotti di collegamento tra i telai"] == 11 * 5 * 2 + 6 * 3 * 2
    assert "Mantovana parasassi" in voci
    p["tipo"] = "Tubi e giunti"
    assert "Giunti ortogonali" in dict((n, q) for n, q, _ in P.elementi(p))
    p["geometria"] = json.dumps({"mantovana": False, "lati": [{"nome": "A", "lunghezza": 5}]})
    assert "Mantovana parasassi" not in dict((n, q) for n, q, _ in P.elementi(p))


def test_avvisi_geometria():
    assert any("parapetto anche sul lato interno" in m for m in P.avvisi_geometria(_ponteggio(distanza=0.35)))
    assert any("teli" in m for m in P.avvisi_geometria(_ponteggio(teli=True)))
    assert any("Ancoraggi" in m for m in P.avvisi_geometria(_ponteggio(ancoraggi_campate=6, ancoraggi_livelli=4)))
    assert P.avvisi_geometria(_ponteggio()) == []
    # un lato più alto di 20 m richiede il progetto anche se l'altezza del ponteggio è minore
    alto = _ponteggio(lati=[{"nome": "Torre", "lunghezza": 6, "altezza": 24}])
    assert any("art. 133" in m for m in pimus.avvisi(alto))


def test_tratti_collegati_staccati_e_ruotati():
    g = {"tratti": [
        {"nome": "Nord", "lunghezza": 10},
        {"nome": "Strettoia", "lunghezza": 3, "tipo": "Tubi e giunti", "larghezza": 0.6, "svolta": 0},
        {"nome": "Est", "lunghezza": 8, "svolta": 90},
        {"nome": "Corpo ruotato", "lunghezza": 6, "collegato": False, "x": -5, "y": 2, "angolo": 30,
         "quota": 3.5, "altezza": 4, "verso": "interno"},
    ]}
    p = {**_ponteggio(), "geometria": json.dumps(g)}
    nord, strettoia, est, ruotato = P.tratti(p)
    assert strettoia.origine == (10, 0) and strettoia.angolo == 0  # prosegue sulla stessa facciata
    assert strettoia.tipo == "Tubi e giunti" and strettoia.larghezza == 0.6 and nord.larghezza == 1.05
    assert est.origine == (13, 0) and est.angolo == 90
    assert ruotato.origine == (-5, 2) and ruotato.angolo == 30 and ruotato.quota == 3.5 and ruotato.livelli == 2
    assert ruotato.normale == pytest.approx((-0.5, 3 ** 0.5 / 2))  # a sinistra della facciata
    assert P.tipologie(p) == ["Telai prefabbricati", "Tubi e giunti"]
    voci = [n for n, _, _ in P.elementi(p)]
    assert "Giunti ortogonali (tubi e giunti)" in voci and "Telai (stilate × livelli) (telai prefabbricati)" in voci
    assert any("tipologie diverse" in m for m in P.avvisi_geometria(p))
    assert not P.chiuso(P.tratti(p))
    quadrato = {**p, "geometria": json.dumps({"tratti": [{"lunghezza": 6, "svolta": 90}] * 4})}
    assert P.chiuso(P.tratti(quadrato))
    assert "Parapetti di testata" not in [n for n, _, _ in P.elementi(quadrato)]


def test_elementi_inseriti_e_ancoraggio_su_tubo():
    p = _ponteggio()
    g = P.geometria(p)
    prima = P.ancoraggi(p).numero
    g = P.cambia_elemento(g, "ancoraggio_tubo", 0, 3, 2)
    g = P.cambia_elemento(g, "trave_carraia", 0, 4, 3)  # a terra: il livello diventa 0
    g = P.cambia_elemento(g, "mensola_esterna", 1, 1, 4)
    assert [(e["tipo"], e["livello"]) for e in g["elementi"]] == [
        ("ancoraggio_tubo", 2), ("trave_carraia", 0), ("mensola_esterna", 4)]
    p["geometria"] = P.testo_geometria(g)
    assert P.ancoraggi(p).numero == prima + 1
    voci = dict((n, q) for n, q, _ in P.elementi(p))
    assert voci["Ancoraggio su tubo e giunto"] == 1 and voci["Giunti ortogonali per ancoraggi su tubo"] == 2
    assert voci["Mensola esterna (sbalzo)"] == 2  # due stilate
    avvisi = P.avvisi_geometria(p)
    assert any("tubo e giunto" in m for m in avvisi) and any("trave carraia" in m for m in avvisi)
    g = P.cambia_elemento(P.geometria(p), "ancoraggio_tubo", 0, 3, 2)  # secondo clic: tolto
    assert [e["tipo"] for e in g["elementi"]] == ["trave_carraia", "mensola_esterna"]
    # elementi di tratti che non esistono più vengono scartati
    g["elementi"].append({"tipo": "telo", "tratto": 9})
    assert len(P.geometria({"geometria": P.testo_geometria(g)})["elementi"]) == 2


def test_unione_di_due_ponteggi():
    p = _ponteggio()
    g = P.geometria(p)
    altro = P.geometria({"geometria": json.dumps({
        "tratti": [{"nome": "Cortile", "lunghezza": 6, "x": 0, "y": 0}, {"nome": "Retro", "lunghezza": 4}],
        "elementi": [{"tipo": "telo", "tratto": 1, "da": 0, "a": 1, "livello": 1}]})})
    unito = P.unisci(g, altro, P.tratti(p, g))
    p["geometria"] = P.testo_geometria(unito)
    nomi = [t.nome for t in P.tratti(p)]
    assert nomi == ["Nord", "Est", "Cortile", "Retro"]
    cortile = P.tratti(p)[2]
    assert cortile.origine == (20.0, 9.0)  # 2 m dopo la fine dell'ultimo tratto
    assert P.geometria(p)["elementi"][0]["tratto"] == 3


def test_linee_elettriche_e_autorizzazione():
    assert pimus.distanza_minima("da 30 a 132 kV") == 5
    assert pimus.avvisi_linee({"linee_tensione": ""}) == []
    assert "3,5 m" in pimus.avvisi_linee({"linee_tensione": "da 1 a 30 kV", "linee_distanza": 2})[0]
    assert pimus.avvisi_linee({"linee_tensione": "da 1 a 30 kV", "linee_distanza": 4}) == []
    assert pimus.scadenza_autorizzazione({"aut_data": "2015-06-01", "aut_rinnovo": "2024-01-10"}) == date(2034, 1, 10)
    assert "scaduta" in pimus.avvisi_autorizzazione({"aut_data": "2010-01-01"}, date(2026, 10, 28))[0]
    assert pimus.avvisi_autorizzazione({"aut_data": "2020-01-01"}, date(2026, 10, 28)) == []


def test_formazione_della_squadra():
    oggi = date(2026, 10, 26)
    valido = {"nome": "Luca", "ruolo": P.PREPOSTO, "corso_data": "2020-03-01", "aggiornamento_data": "2024-02-01"}
    assert P.scadenza_formazione(valido) == date(2028, 2, 1)
    assert P.stato_formazione(valido, oggi) == "Valida"
    scaduto = {"nome": "Mario", "ruolo": P.ADDETTO, "corso_data": "2021-01-10", "aggiornamento_data": None}
    assert P.stato_formazione(scaduto, oggi) == "Aggiornamento scaduto"
    assert P.stato_formazione({"corso_data": "2022-12-01"}, oggi) == "In scadenza"
    assert P.stato_formazione({"corso_data": None}, oggi) == "Corso da indicare"
    messaggi = P.avvisi_squadra([valido, scaduto], oggi)
    assert len(messaggi) == 1 and "Mario" in messaggi[0] and "10/01/2025" in messaggi[0]
    assert any("preposto" in m for m in P.avvisi_squadra([scaduto], oggi))


def test_migrazione_porta_preposto_e_squadra(tmp_path):
    percorso = tmp_path / "vecchio.spm"
    conn = sqlite3.connect(percorso)
    for n, sql in enumerate(MIGRAZIONI[:-3], 1):  # file della versione 0.23
        conn.executescript(f"BEGIN;{sql};PRAGMA user_version = {n};COMMIT;")
    conn.execute("INSERT INTO ponteggi (nome, preposto, squadra) VALUES ('Nord', 'Luca Bianchi', 'Mario, Gino')")
    conn.commit()
    conn.close()
    progetto = Progetto(percorso)
    repo = Repository(progetto)
    addetti = repo.elenco("addetti_ponteggio")
    assert [(a["nome"], a["ruolo"]) for a in addetti] == [("Luca Bianchi", P.PREPOSTO), ("Mario, Gino", P.ADDETTO)]
    assert repo.elenco("ponteggi")[0]["geometria"] == "{}"
    progetto.chiudi()


def test_pimus_con_squadra_disegni_e_allegati(repo, tmp_path):
    from PySide6.QtGui import QColor, QImage

    from sicurcantiere.documenti.pimus import genera_pimus

    repo.aggiorna_cantiere(nome="Scuola", indirizzo="Via Roma 1", comune="Bari")
    impresa = repo.inserisci("imprese", ragione_sociale="Edil Ponteggi", datore_lavoro="Anna Verdi")
    p = repo.inserisci("ponteggi", nome="Facciate", impresa_id=impresa, altezza=10,
                       classe_carico=P.CLASSI_CARICO[3][0], geometria=_ponteggio()["geometria"])
    repo.inserisci("addetti_ponteggio", ponteggio_id=p, nome="Luca Bianchi", ruolo=P.PREPOSTO,
                   corso_data="2024-05-10", corso_ente="Scuola Edile")
    repo.inserisci("addetti_ponteggio", ponteggio_id=p, nome="Mario Rossi", corso_data="2023-01-15")
    immagine = tmp_path / "disegno esecutivo.png"
    img = QImage(40, 30, QImage.Format.Format_RGB32)
    img.fill(QColor("#336699"))
    img.save(str(immagine))
    for descrizione, probabilita, danno, misure, dpi in pimus.RISCHI_MONTAGGIO[:2]:
        repo.inserisci("rischi_ponteggio", ponteggio_id=p, descrizione=descrizione, probabilita=probabilita,
                       danno=danno, misure=misure, dpi=dpi)
    repo.inserisci("allegati_ponteggio", ponteggio_id=p, tipo="Disegno esecutivo", titolo="Tavola 1",
                   file=str(immagine))
    repo.inserisci("allegati_ponteggio", ponteggio_id=p, tipo="Libretto del fabbricante", titolo="Libretto",
                   file=str(tmp_path / "manca.pdf"))
    file = genera_pimus(repo, p, tmp_path / "pimus.docx")
    doc = Document(file)
    testo = "\n".join([x.text for x in doc.paragraphs]
                      + [c.text for t in doc.tables for r in t.rows for c in r.cells])
    assert "Luca Bianchi" in testo and "Scuola Edile" in testo and "10/05/2028" in testo
    assert "Telai (stilate × livelli)" in testo and "ancoraggi" in testo
    assert "Verifica dell'appoggio" in testo and "verificata" in testo
    assert "Prospetto · Nord" in testo and "Sezione tipo" in testo
    assert "Allegato 1 · Disegno esecutivo" in testo and "File non trovato" in testo
    assert "Valutazione dei rischi delle operazioni" in testo and "Caduta dall'alto degli addetti" in testo
    assert "f) Misure in presenza di linee elettriche aeree" in testo and "Progetto del ponteggio" in testo
    # vista d'insieme, pianta, 2 prospetti, sezione e il disegno allegato
    assert len(doc.inline_shapes) >= 6


def test_testi_base_per_tipologia():
    telai = pimus.testo_base("montaggio", ["Telai prefabbricati"])
    assert "spinotti" in telai and "tubi e giunti" not in telai
    assert "giunti di testa" in pimus.testo_base("montaggio", ["Tubi e giunti"])
    assert "rosette" in pimus.testo_base("montaggio", ["Multidirezionale"])
    misto = pimus.testo_base("montaggio", ["Telai prefabbricati", "Tubi e giunti"])
    assert misto.startswith(telai) and "raccordi in tubi e giunti" in misto
    # sezioni senza testo specifico: il testo generico
    generico = dict((k, t) for k, _, t in pimus.SEZIONI)["meteo"]
    assert pimus.testo_base("meteo", ["Multidirezionale"]) == generico
    assert pimus.testo_base("uso", []) == dict((k, t) for k, _, t in pimus.SEZIONI)["uso"]


def test_archivio_dei_sistemi_di_ponteggio(repo, tmp_path):
    from sicurcantiere.archivio_utente import SISTEMA, ArchivioAnagrafiche

    archivio = ArchivioAnagrafiche(tmp_path / "archivio.db")
    p = repo.inserisci("ponteggi", nome="Nord", fabbricante="Acme Ponteggi", sistema="Telaio 105", modello="T5",
                       autorizzazione="n. 12345", aut_data="2019-05-10", classe_carico=P.CLASSI_CARICO[3][0],
                       geometria=json.dumps({"campata": 2.5, "larghezza": 0.75, "peso_proprio": 0.27,
                                             "tratti": [{"nome": "Nord", "lunghezza": 20}]}))
    repo.inserisci("allegati_ponteggio", ponteggio_id=p, tipo="Libretto del fabbricante", titolo="Libretto",
                   file="C:/Libretti/acme.pdf")
    repo.inserisci("allegati_ponteggio", ponteggio_id=p, tipo="Foto", titolo="Facciata", file="C:/foto.jpg")
    archivio.salva_sistema(repo, p)
    voce = archivio.elenco(SISTEMA)[0]
    assert voce["nome"] == "Acme Ponteggi Telaio 105 T5"
    assert voce["dati"]["geometria"]["campata"] == 2.5 and len(voce["dati"]["allegati"]) == 1
    # in un altro ponteggio: dati del sistema e valori comuni, la geometria del cantiere resta
    altro = repo.inserisci("ponteggi", nome="Sud", geometria=json.dumps({"tratti": [{"nome": "Sud", "lunghezza": 7}]}))
    archivio.applica_sistema(repo, altro, voce)
    r = repo.leggi("ponteggi", altro)
    assert r["fabbricante"] == "Acme Ponteggi" and r["autorizzazione"] == "n. 12345" and r["aut_data"] == "2019-05-10"
    g = P.geometria(r)
    assert g["campata"] == 2.5 and g["peso_proprio"] == 0.27 and g["tratti"][0]["nome"] == "Sud"
    assert [a["titolo"] for a in repo.elenco("allegati_ponteggio", ponteggio_id=altro)] == ["Libretto"]
    archivio.applica_sistema(repo, altro, voce)  # una seconda volta il libretto non si duplica
    assert len(repo.elenco("allegati_ponteggio", ponteggio_id=altro)) == 1
    with pytest.raises(ValueError):
        archivio.salva_sistema(repo, repo.inserisci("ponteggi", nome="Senza dati"))
    archivio.chiudi()


def _dxf(file):
    righe = file.read_text(encoding="cp1252").splitlines()
    coppie = list(zip(righe[::2], righe[1::2]))
    return coppie, [v for c, v in coppie if c == "0"]


def test_esportazione_dxf(tmp_path):
    from sicurcantiere.documenti import dxf_ponteggio as X

    p = {**_ponteggio(), "geometria": json.dumps({
        "tratti": [{"nome": "Nord", "lunghezza": 9}, {"nome": "Est", "lunghezza": 5, "svolta": 90}],
        "elementi": [{"tipo": "ancoraggio_tubo", "tratto": 0, "da": 1, "a": 1, "livello": 2}]})}
    coppie, entita = _dxf(X.esporta(p, tmp_path / "m.dxf"))
    assert ("1", "AC1009") in coppie and entita[-1] == "EOF"
    assert entita.count("LINE") > 100 and entita.count("3DFACE") > 10
    layer = {v for c, v in coppie if c == "2"}
    assert {"PONT_MONTANTI", "PONT_IMPALCATI", "PONT_ANCORAGGI", "PONT_TUBI_GIUNTI", "EDIFICIO"} <= layer
    # quote z nel modello 3D: i montanti arrivano oltre i 12 m
    assert max(float(v) for c, v in coppie if c in ("30", "31")) > 12
    coppie, entita = _dxf(X.esporta(p, tmp_path / "t.dxf", X.TAVOLE_2D))
    testi = [v for c, v in coppie if c == "1"]
    assert "PIANTA DEL PONTEGGIO" in testi and "PROSPETTO - Nord" in testi and "PROSPETTO - Est" in testi
    assert any(t.startswith("SEZIONE TIPO") for t in testi) and "+12,00" in testi
    assert "CARTIGLIO DEL DISEGNO ESECUTIVO" in testi and "S1" not in testi
    assert all(float(v) == 0 for c, v in coppie if c in ("30", "31"))  # tavole piane


def test_contenuti_dal_materiale_pimus(repo, tmp_path):
    """Persona competente, contesto, conformità al PSC, sistema 7.1, situazioni particolari, codici degli
    ancoraggi, cartiglio, presa visione e verbale di consegna."""
    from sicurcantiere.documenti.pimus import genera_consegna, genera_pimus

    repo.aggiorna_cantiere(nome="Scuola", indirizzo="Via Roma 1", comune="Bari")
    impresa = repo.inserisci("imprese", ragione_sociale="Edil Ponteggi", datore_lavoro="Anna Verdi")
    p = repo.inserisci("ponteggi", nome="Facciate", impresa_id=impresa, altezza=10, autorizzazione="n. 1",
                       geometria=json.dumps({"tratti": [{"nome": "Nord", "lunghezza": 9}],
                                             "elementi": [{"tipo": "ancoraggio_tubo", "tratto": 0, "da": 1, "a": 1,
                                                           "livello": 2, "note": "Tubo tra i montanti: pilastro "
                                                                                 "arretrato"}]}))
    avvisi = pimus.avvisi(repo.leggi("ponteggi", p))
    assert any("persona competente" in m for m in avvisi) and any("sistema adottato" in m for m in avvisi)
    repo.aggiorna("ponteggi", p, redattore="Ing. Paolo Neri", redattore_qualifica=pimus.QUALIFICHE_REDATTORE[1],
                  opera_servita="Ristrutturazione della facciata", impresa_utilizzatrice="Restauri Sud srl",
                  conformita_psc=pimus.CONFORMITA_PSC[2], sistema_anticaduta=pimus.SISTEMI_ANTICADUTA[3][0],
                  situazioni=";".join(pimus.SITUAZIONI[:2]), situazioni_note="Basette snodate sulla rampa.")
    avvisi = pimus.avvisi(repo.leggi("ponteggi", p))
    assert any("difformità" in m for m in avvisi) and not any("persona competente" in m for m in avvisi)
    repo.inserisci("addetti_ponteggio", ponteggio_id=p, nome="Luca Bianchi", ruolo=P.PREPOSTO,
                   corso_data="2024-01-10", note="Dieci anni di montaggio di ponteggi a telai")
    codici = P.elenco_ancoraggi(repo.leggi("ponteggi", p))
    assert codici[0]["codice"] == "A1.01" and codici[-1]["stilata"] == "campata 2"
    assert pimus.caso_progetto(repo.leggi("ponteggi", p)).startswith("A - ")
    doc = Document(genera_pimus(repo, p, tmp_path / "pimus.docx"))
    testo = "\n".join([x.text for x in doc.paragraphs] + [c.text for t in doc.tables for r in t.rows for c in r.cells])
    for atteso in ("Ing. Paolo Neri", "Restauri Sud srl", "Ristrutturazione della facciata", "Difforme dal PSC",
                   "d.1) Sistema adottato", "Parapetto provvisorio montato dal basso", "Partenza su piani inclinati",
                   "Basette snodate sulla rampa", "pilastro arretrato", "A1.01", "Cartiglio del disegno esecutivo",
                   "Presa visione del piano", "Dieci anni di montaggio", "Interdizione delle parti",
                   "Cessione del ponteggio", "La persona competente che ha redatto il piano", "Circolare MLPS"):
        assert atteso in testo, atteso
    consegna = Document(genera_consegna(repo, p, tmp_path / "consegna.docx"))
    testo = "\n".join([x.text for x in consegna.paragraphs]
                      + [c.text for t in consegna.tables for r in t.rows for c in r.cells])
    assert "Verbale di consegna del ponteggio" in testo and "Restauri Sud srl" in testo and "T1 Nord" in testo
