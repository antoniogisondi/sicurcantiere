from datetime import date

import pytest
from docx import Document

from sicurcantiere import analisi
from sicurcantiere.archivio import ATTREZZATURE, CATEGORIE, LAVORAZIONI
from sicurcantiere.db import MIGRAZIONI, ErroreProgetto, Progetto, VERSIONE_SCHEMA
from sicurcantiere.documenti.psc import genera_psc
from sicurcantiere.repository import Repository


def lav(codice):
    return next(l for l in LAVORAZIONI if l["codice"] == codice)


@pytest.fixture
def repo(tmp_path):
    progetto = Progetto(tmp_path / "prova.scs", crea=True)
    yield Repository(progetto)
    progetto.chiudi()


def test_nuovo_progetto_ha_schema_aggiornato(tmp_path):
    p = Progetto(tmp_path / "a.scs", crea=True)
    assert p.query("PRAGMA user_version")[0][0] == VERSIONE_SCHEMA
    p.chiudi()
    # Riapertura: crea il backup e non riapplica le migrazioni.
    p = Progetto(tmp_path / "a.scs")
    assert not (tmp_path / "a.scs.bak").exists()  # nessuna copia accanto al file
    from sicurcantiere.db import COPIE_TENUTE, cartella_copie

    for _ in range(COPIE_TENUTE + 2):
        Progetto(tmp_path / "a.scs").chiudi()
    copie = list(cartella_copie().glob("a-*.scs"))
    assert len(copie) == COPIE_TENUTE
    p.chiudi()


def test_file_non_valido(tmp_path):
    f = tmp_path / "rotto.scs"
    f.write_bytes(b"non sono un database" * 100)
    with pytest.raises(ErroreProgetto):
        Progetto(f)


def test_versione_futura_rifiutata(tmp_path):
    p = Progetto(tmp_path / "f.scs", crea=True)
    p.conn.execute(f"PRAGMA user_version = {VERSIONE_SCHEMA + 1}")
    p.chiudi()
    with pytest.raises(ErroreProgetto):
        Progetto(tmp_path / "f.scs")


def test_campi_non_validi_rifiutati(repo):
    with pytest.raises(ValueError):
        repo.inserisci("imprese", **{"ragione_sociale; DROP TABLE imprese": "x"})
    with pytest.raises(ValueError):
        repo.elenco("sqlite_master")


def test_crud_e_cascata(repo):
    impresa = repo.inserisci("imprese", ragione_sociale="Edil Srl")
    lavorazione = lav("AL.14")
    fase = repo.importa_lavorazione(lavorazione)
    repo.aggiorna("fasi", fase, impresa_id=impresa)
    assert repo.leggi("fasi", fase)["impresa_id"] == impresa
    attese = len(lavorazione["rischi"]) + sum(len(ATTREZZATURE[a]["rischi"]) for a in lavorazione["attrezzature"])
    assert len(repo.elenco("rischi", fase_id=fase)) == attese
    assert {a["nome"] for a in repo.elenco("attrezzature", fase_id=fase)} == set(lavorazione["attrezzature"])
    repo.elimina("imprese", impresa)
    assert repo.leggi("fasi", fase)["impresa_id"] is None
    repo.elimina("fasi", fase)
    assert repo.elenco("rischi") == []
    assert repo.elenco("attrezzature") == []


def test_salva_con_nome(repo, tmp_path):
    repo.aggiorna_cantiere(nome="Scuola")
    repo.progetto.salva_con_nome(tmp_path / "copia.scs")
    copia = Progetto(tmp_path / "copia.scs")
    assert Repository(copia).cantiere()["nome"] == "Scuola"
    copia.chiudi()


def test_classi_rischio():
    assert analisi.classe_rischio(analisi.livello_rischio(1, 2)) == "Basso"
    assert analisi.classe_rischio(analisi.livello_rischio(2, 2)) == "Medio"
    assert analisi.classe_rischio(analisi.livello_rischio(2, 3)) == "Alto"
    assert analisi.classe_rischio(analisi.livello_rischio(3, 4)) == "Molto alto"


def test_giorni_lavorativi():
    # 6-19 ottobre 2025: due settimane complete, lunedì -> domenica.
    assert analisi.giorni_lavorativi(date(2025, 10, 6), date(2025, 10, 19)) == 10
    assert analisi.giorni_lavorativi(date(2025, 10, 10), date(2025, 10, 13)) == 2  # ven -> lun
    assert analisi.giorni_lavorativi(date(2025, 10, 11), date(2025, 10, 12)) == 0
    assert analisi.giorni_lavorativi(date(2025, 10, 13), date(2025, 10, 10)) == 0


def _fase(id_, impresa, inizio, fine, lavoratori=1):
    return {"id": id_, "codice": f"F{id_}", "descrizione": "", "impresa_id": impresa,
            "data_inizio": inizio, "data_fine": fine, "lavoratori": lavoratori}


def test_interferenze():
    fasi = [
        _fase(1, 1, "2025-10-06", "2025-10-17"),
        _fase(2, 2, "2025-10-15", "2025-10-24"),  # si sovrappone a 1, altra impresa
        _fase(3, 1, "2025-10-16", "2025-10-16"),  # stessa impresa di 1, sovrapposta a 2
        _fase(4, 2, "2025-11-01", "2025-11-05"),  # isolata
        _fase(5, None, None, None),
    ]
    coppie = {(i.fase_a["id"], i.fase_b["id"]) for i in analisi.interferenze(fasi)}
    assert coppie == {(1, 2), (2, 3)}
    prima = analisi.interferenze(fasi)[0]
    assert (prima.inizio, prima.fine) == (date(2025, 10, 15), date(2025, 10, 17))


def test_obblighi():
    imprese = [{"id": 1, "tipo": "Affidataria"}, {"id": 2, "tipo": "Esecutrice"},
               {"id": 3, "tipo": "Lavoratore autonomo"}]
    una = [_fase(1, 1, "2025-10-06", "2025-10-10", 3)]
    o = analisi.obblighi(imprese, una)
    assert (o.imprese, o.uomini_giorno, o.psc, o.notifica) == (1, 15, False, False)

    grande = [_fase(1, 1, "2025-01-06", "2025-03-28", 4)]  # 60 gg lavorativi x 4
    o = analisi.obblighi(imprese, grande)
    assert o.uomini_giorno == 240 and not o.psc and o.notifica

    due = una + [_fase(2, 2, "2025-12-01", "2025-12-02")]
    o = analisi.obblighi(imprese, due)
    assert o.psc and o.notifica


def test_genera_psc(repo, tmp_path):
    repo.aggiorna_cantiere(nome="Ristrutturazione scuola", comune="Bari", provincia="BA",
                           importo_lavori=250000)
    repo.inserisci("soggetti", ruolo="Committente", nome="Comune di Bari")
    a = repo.inserisci("imprese", ragione_sociale="Alfa Costruzioni", tipo="Affidataria")
    b = repo.inserisci("imprese", ragione_sociale="Beta Impianti")
    f1 = repo.importa_lavorazione(lav("MT.01"))
    f2 = repo.importa_lavorazione(lav("AL.22"))
    repo.aggiorna("fasi", f1, impresa_id=a, data_inizio="2025-10-06", data_fine="2025-10-31")
    repo.aggiorna("fasi", f2, impresa_id=b, data_inizio="2025-10-20", data_fine="2025-11-14")
    repo.inserisci("costi", descrizione="Recinzione", unita="m", quantita=120, prezzo_unitario=12.5)
    repo.salva_testo("emergenze", "Presidio di primo soccorso nel box.")

    file = genera_psc(repo, tmp_path / "psc.docx")
    doc = Document(file)
    testo = "\n".join(p.text for p in doc.paragraphs)
    celle = "\n".join(c.text for t in doc.tables for r in t.rows for c in r.cells)
    assert "Piano di sicurezza e coordinamento" in testo
    assert "Presidio di primo soccorso" in testo
    # capitoli dell'Allegato XV punto 2.1.2 e dati mancanti segnalati
    for titolo in ("Identificazione e descrizione dell'opera", "Soggetti con compiti di sicurezza",
                   "Area di cantiere", "Organizzazione del cantiere", "Lavorazioni: rischi e misure per fase",
                   "Interferenze tra le lavorazioni e prescrizioni operative", "Cronoprogramma dei lavori",
                   "Stima dei costi della sicurezza"):
        assert titolo in testo
    assert "[Da compilare: coordinatore per la progettazione (CSP)]" in testo
    assert "Comune di Bari" in celle
    assert "Diagramma di Gantt dei lavori" in testo and doc.inline_shapes  # diagramma come immagine
    assert any(s.orientation == 1 for s in doc.sections)  # su pagine orizzontali
    assert "Alfa Costruzioni" in celle
    assert "€ 1.500,00" in celle
    assert "Scavo di sbancamento" in celle  # interferenza e cronoprogramma
    assert "Sfasamento spaziale" in celle  # prescrizione generica


def test_genera_psc_progetto_vuoto(repo, tmp_path):
    assert genera_psc(repo, tmp_path / "vuoto.docx").exists()


def test_migrazione_da_v1(tmp_path):
    import sqlite3

    f = tmp_path / "vecchio.scs"
    conn = sqlite3.connect(f)
    conn.executescript(MIGRAZIONI[0] + ";PRAGMA user_version = 1;")
    conn.execute("INSERT INTO fasi (codice, descrizione) VALUES ('X', 'Vecchia fase')")
    conn.commit()
    conn.close()
    p = Progetto(f)
    r = Repository(p)
    assert p.query("PRAGMA user_version")[0][0] == VERSIONE_SCHEMA
    fase = r.elenco("fasi")[0]
    assert fase["descrizione"] == "Vecchia fase" and fase["zona"] == ""
    assert r.cantiere()["costo_uomo_giorno"] == 0  # vuoto: vale il valore tipico di 250 €
    p.chiudi()


def test_archivio_coerente():
    codici = [l["codice"] for l in LAVORAZIONI]
    assert len(codici) == len(set(codici))
    for l in LAVORAZIONI:
        assert l["categoria"] in CATEGORIE
        assert 0 < l["incidenza_manodopera"] <= 100
        for r in l["rischi"]:
            assert 1 <= r["probabilita"] <= 4 and 1 <= r["danno"] <= 4
            assert r["misure"].strip(), (l["codice"], r["descrizione"])
    for nome, a in ATTREZZATURE.items():
        for r in a["rischi"]:
            assert r["misure"].strip(), (nome, r["descrizione"])


def test_durata_stimata():
    # 100.000 € x 40% / 250 €/ug = 160 ug; squadra di 4 -> 40 giorni lavorativi
    assert analisi.durata_stimata(100_000, 40, 250, 4) == 40
    assert analisi.durata_stimata(1_000, 40, 250, 4) == 1
    assert analisi.durata_stimata(0, 40, 250, 4) is None
    assert analisi.durata_stimata(1_000, 40, 250, 0) is None


def test_fine_dopo_giorni_lavorativi():
    assert analisi.fine_dopo_giorni_lavorativi(date(2025, 10, 6), 1) == date(2025, 10, 6)
    assert analisi.fine_dopo_giorni_lavorativi(date(2025, 10, 6), 5) == date(2025, 10, 10)
    assert analisi.fine_dopo_giorni_lavorativi(date(2025, 10, 6), 6) == date(2025, 10, 13)
    # inizio di sabato: si parte dal lunedì
    assert analisi.fine_dopo_giorni_lavorativi(date(2025, 10, 11), 1) == date(2025, 10, 13)
    d = analisi.fine_dopo_giorni_lavorativi(date(2025, 10, 6), 40)
    assert analisi.giorni_lavorativi(date(2025, 10, 6), d) == 40


def test_interferenze_zona_e_gravita():
    fasi = [
        {**_fase(1, 1, "2025-10-06", "2025-10-17"), "zona": "Edificio A"},
        {**_fase(2, 2, "2025-10-06", "2025-10-17"), "zona": "edificio a "},  # stessa zona
        {**_fase(3, 3, "2025-10-06", "2025-10-17"), "zona": "Edificio B"},
        {**_fase(4, 4, "2025-10-06", "2025-10-17"), "zona": ""},  # zona non indicata
    ]
    livelli = {1: 4, 2: 12, 4: 2}
    elenco = analisi.interferenze(fasi, livelli)
    assert {i.chiave for i in elenco} == {(1, 2), (1, 4), (2, 4), (3, 4)}
    assert elenco[0].chiave == (1, 2) and elenco[0].gravita == 12


def test_sovrapposizione_solo_nel_fine_settimana():
    fasi = [_fase(1, 1, "2025-10-06", "2025-10-11"), _fase(2, 2, "2025-10-11", "2025-10-12")]
    assert analisi.interferenze(fasi) == []


def test_andamento():
    fasi = [_fase(1, 1, "2025-10-06", "2025-10-10", 3), _fase(2, 2, "2025-10-09", "2025-10-14", 2)]
    giornate = analisi.andamento(fasi, {1: 4, 2: 9})
    assert [g.giorno.day for g in giornate] == [6, 7, 8, 9, 10, 13, 14]
    assert [g.lavoratori for g in giornate] == [3, 3, 3, 5, 5, 2, 2]
    assert [g.pericolosita for g in giornate] == [4, 4, 4, 9, 9, 9, 9]
    assert analisi.picco_lavoratori(fasi) == 5


def test_prescrizioni(repo):
    a = repo.inserisci("fasi")
    b = repo.inserisci("fasi")
    repo.salva_prescrizione(b, a, "Lavorare in orari diversi")
    assert repo.prescrizioni() == {(a, b): "Lavorare in orari diversi"}
    repo.salva_prescrizione(a, b, "  ")
    assert repo.prescrizioni() == {}


def test_controlli(repo):
    from sicurcantiere import controlli

    messaggi = {s.messaggio for s in controlli.controlla(repo)}
    assert "Indicare il committente." in messaggi
    assert "Inserire le fasi di lavoro." in messaggi
    repo.inserisci("imprese", ragione_sociale="A")
    repo.inserisci("imprese", ragione_sociale="B")
    errori = [s for s in controlli.controlla(repo) if s.livello == controlli.ERRORE]
    assert any("CSE" in s.messaggio for s in errori)


def _cantiere_completo(repo):
    repo.aggiorna_cantiere(nome="Scuola", indirizzo="Via Roma 1", comune="Bari", provincia="BA",
                           natura_opera="Ristrutturazione", importo_lavori=250000)
    repo.inserisci("soggetti", ruolo="Committente", nome="Comune di Bari")
    repo.inserisci("soggetti", ruolo="Coordinatore per l'esecuzione (CSE)", nome="Ing. Rossi")
    a = repo.inserisci("imprese", ragione_sociale="Alfa Costruzioni", tipo="Affidataria",
                       datore_lavoro="Mario Bianchi")
    repo.inserisci("figure", impresa_id=a, ruolo="RSPP", nome="Luca Verdi")
    repo.inserisci("lavoratori", impresa_id=a, nome="Giuseppe Neri", mansione="Muratore",
                   formazione="Formazione generale e specifica 16 ore")
    f = repo.importa_lavorazione(lav("ST.07"))
    repo.aggiorna("fasi", f, impresa_id=a, data_inizio="2025-10-06", data_fine="2025-10-31", lavoratori=4)
    repo.inserisci("interventi", parte_opera="Copertura", intervento="Pulizia dei canali di gronda",
                   periodicita="Annuale", misure_in_dotazione="Linea vita")
    repo.inserisci("elaborati", titolo="Progetto strutturale", autore="Ing. Blu")
    repo.inserisci("revisioni", documento="PSC", numero=1, data="2025-09-01", descrizione="Prima emissione")
    return a


def _testo_tabelle(tabelle):
    for t in tabelle:
        for r in t.rows:
            for c in r.cells:
                yield c.text
                yield from _testo_tabelle(c.tables)


def _testo_docx(file):
    doc = Document(file)
    return "\n".join([p.text for p in doc.paragraphs] + list(_testo_tabelle(doc.tables)))


def test_genera_pos(repo, tmp_path):
    from sicurcantiere.documenti.pos import chiave_testo, genera_pos

    a = _cantiere_completo(repo)
    repo.salva_testo(chiave_testo(a, "rumore"), "Esposizione inferiore a 80 dB(A).")
    testo = _testo_docx(genera_pos(repo, a, tmp_path / "pos.docx"))
    assert "Piano operativo di sicurezza" in testo
    assert "Mario Bianchi" in testo and "Giuseppe Neri" in testo
    assert "Gru a torre" in testo
    assert "80 dB(A)" in testo
    assert "Otoprotettori" in testo  # DPI raccolti dalle schede


def test_consegna_e_verifica_del_pos(repo, tmp_path):
    """Copia al RLS 10 giorni prima (art. 100), trasmissione e inizio dopo le verifiche (art. 101)."""
    from sicurcantiere import controlli
    from sicurcantiere.documenti.pos import genera_pos

    a = _cantiere_completo(repo)  # lavori dal 06/10/2025

    def messaggi():
        return [x.messaggio for x in controlli.controlla_pos(repo)]

    assert any("messa a disposizione del RLS" in m for m in messaggi())
    repo.aggiorna("imprese", a, pos_rls_data="2025-10-01")
    assert any("solo 5 giorni prima" in m for m in messaggi())
    repo.aggiorna("imprese", a, pos_rls_data="2025-09-20", pos_affidataria_data="2025-09-22",
                  pos_cse_data="2025-09-24", pos_esito_data="2025-10-10")
    assert any("prima dell'esito positivo" in m for m in messaggi())
    repo.aggiorna("imprese", a, pos_esito_data="2025-10-01")
    assert not any("art. 100" in m or "art. 101" in m for m in messaggi())
    testo = _testo_docx(genera_pos(repo, a, tmp_path / "pos.docx"))
    assert "Consegna, verifica e aggiornamento del piano" in testo and "20/09/2025" in testo
    assert "almeno dieci giorni prima dell'inizio dei lavori" in testo
    # solo un lavoratore autonomo: niente POS
    repo2 = Repository(Progetto(tmp_path / "b.spo", crea=True))
    repo2.inserisci("imprese", ragione_sociale="Mario Rossi", tipo="Lavoratore autonomo")
    assert any("non redige il POS" in x.messaggio for x in controlli.controlla_pos(repo2))
    repo2.progetto.chiudi()


def test_valutazione_del_rumore_vecchia(repo):
    from sicurcantiere import controlli

    _cantiere_completo(repo)
    repo.inserisci("valutazioni", metodo="rumore", titolo="Rumore 2019", data="2019-03-01")
    assert any("più vecchia di quattro anni" in x.messaggio for x in controlli.controlla_pos(repo))


def test_genera_fascicolo_e_notifica(repo, tmp_path):
    from sicurcantiere.documenti.fascicolo import genera_fascicolo
    from sicurcantiere.documenti.notifica import genera_notifica

    _cantiere_completo(repo)
    testo = _testo_docx(genera_fascicolo(repo, tmp_path / "f.docx"))
    assert "Pulizia dei canali di gronda" in testo and "Linea vita" in testo
    assert "Progetto strutturale" in testo
    testo = _testo_docx(genera_notifica(repo, tmp_path / "n.docx"))
    assert "Comune di Bari" in testo and "Ing. Rossi" in testo
    assert "06/10/2025" in testo and "20 giorni lavorativi" in testo


def test_revisioni_nel_frontespizio(repo, tmp_path):
    _cantiere_completo(repo)
    assert "Prima emissione" in _testo_docx(genera_psc(repo, tmp_path / "p.docx"))


def test_elenco_dpi():
    from sicurcantiere.documenti.pos import elenco_dpi

    assert elenco_dpi(["Casco, guanti", "casco; otoprotettori.", ""]) == ["Casco", "Guanti", "Otoprotettori"]


def test_anteprima(tmp_path):
    from sicurcantiere.db import anteprima

    p = Progetto(tmp_path / "a.scs", crea=True)
    r = Repository(p)
    r.aggiorna_cantiere(nome="Scuola", comune="Bari", provincia="BA")
    p.chiudi()
    assert anteprima(tmp_path / "a.scs") == {"nome": "Scuola", "comune": "Bari", "provincia": "BA",
                                             "committente": "", "impresa": ""}
    # con committente, impresa e un documento composto: la pagina iniziale ne mostra lo stato
    from sicurcantiere import composizioni

    p = Progetto(tmp_path / "a.scs")
    r = Repository(p)
    r.inserisci("soggetti", ruolo="Committente", nome="Comune di Bari")
    r.inserisci("imprese", ragione_sociale="Alfa srl")
    composizioni.componi(r, "PSC")
    p.chiudi()
    dati = anteprima(tmp_path / "a.scs")
    assert dati["committente"] == "Comune di Bari" and dati["impresa"] == "Alfa srl"
    assert dati["stato"] == "bozza" and dati["documento"] == "PSC" and dati["mancanti"] > 0
    import shutil

    shutil.rmtree(tmp_path / "copie", ignore_errors=True)
    anteprima(tmp_path / "a.scs")
    assert not (tmp_path / "copie").exists()  # sola lettura: niente copia di sicurezza
    (tmp_path / "rotto.scs").write_bytes(b"x" * 200)
    assert anteprima(tmp_path / "rotto.scs") is None
    assert anteprima(tmp_path / "manca.scs") is None


def test_valutazione_e_pos(repo, tmp_path):
    from sicurcantiere.documenti.pos import genera_pos
    from sicurcantiere.documenti.valutazione import genera_valutazione

    a = _cantiere_completo(repo)
    fase = repo.elenco("fasi")[0]["id"]
    v = repo.inserisci("valutazioni", metodo="niosh", titolo="Sacchi di cemento", mansione="Manovale")
    altra = repo.inserisci("fasi", codice="Z", descrizione="Massetti", impresa_id=None)
    repo.salva_fasi_valutazione(v, [fase, altra, fase])
    assert repo.fasi_valutazione(v) == sorted([fase, altra])
    repo.salva_dati_valutazione(v, {"peso": 25, "h": 45, "v": 10, "d": 90, "a": 45, "frequenza": 2,
                                    "durata": "lunga", "presa": "scarsa", "lavoratori": "M18-45",
                                    "descrizione": "Carico sacchi in betoniera",
                                    "ipotesi_ai": ["Sacchi presi da terra"]})
    testo = _testo_docx(genera_valutazione(repo, v, tmp_path / "v.docx"))
    assert "Sacchi di cemento" in testo and "UNI ISO 11228-1" in testo
    assert "Carico sacchi in betoniera" in testo and "Sacchi presi da terra" in testo
    assert "Rischio elevato" in testo and "Z Massetti" in testo
    pos = _testo_docx(genera_pos(repo, a, tmp_path / "pos.docx"))
    assert "Valutazioni dei rischi specifici" in pos and "Sacchi di cemento" in pos


def test_pimus(repo, tmp_path):
    from sicurcantiere import pimus
    from sicurcantiere.documenti.pimus import genera_pimus

    a = _cantiere_completo(repo)
    p = repo.inserisci("ponteggi", nome="Facciata nord", impresa_id=a, altezza=24, preposto="Luca Bianchi",
                       autorizzazione="n. 123 del 2020")
    avvisi = pimus.avvisi(repo.leggi("ponteggi", p))
    assert any("Progetto del ponteggio obbligatorio" in m for m in avvisi)
    assert not any("preposto" in m for m in avvisi)
    repo.salva_testo(pimus.chiave_testo(p, "uso"), "Vietato rimuovere i parapetti.")
    for momento, voce in pimus.VERIFICHE[:2]:
        repo.inserisci("verifiche_ponteggio", ponteggio_id=p, momento=momento, voce=voce, esito="Conforme")
    testo = _testo_docx(genera_pimus(repo, p, tmp_path / "pimus.docx"))
    assert "PiMUS" in testo and "Facciata nord" in testo and "Luca Bianchi" in testo
    assert "Vietato rimuovere i parapetti." in testo
    assert "Conforme" in testo and "art. 133" in testo
    basso = repo.inserisci("ponteggi", nome="Basso", altezza=8, conforme_schemi=1)
    assert not any("Progetto" in m for m in pimus.avvisi(repo.leggi("ponteggi", basso)))


def test_controlli_pos_e_consegna_dpi(repo, tmp_path):
    from sicurcantiere import controlli
    from sicurcantiere.documenti.moduli import genera_consegna_dpi

    messaggi = [s.messaggio for s in controlli.controlla_pos(repo)]
    assert "Inserire l'impresa che redige il POS." in messaggi
    a = _cantiere_completo(repo)
    repo.aggiorna("lavoratori", repo.elenco("lavoratori")[0]["id"], idoneita_scadenza="2000-01-01")
    messaggi = [s.messaggio for s in controlli.controlla_pos(repo)]
    assert any("idoneità sanitaria scaduta" in m for m in messaggi)
    assert any("medico competente" in m for m in messaggi)
    testo = _testo_docx(genera_consegna_dpi(repo, a, tmp_path / "dpi.docx"))
    assert "Verbale di consegna dei DPI" in testo and "Giuseppe Neri" in testo and "Casco" in testo


def test_sostanze_nel_pos_e_elenco(repo, tmp_path):
    from sicurcantiere import controlli
    from sicurcantiere.documenti.pos import genera_pos
    from sicurcantiere.documenti.sostanze import genera_elenco_sostanze, sostanze_impresa

    a = _cantiere_completo(repo)
    fase = repo.elenco("fasi")[0]["id"]
    altra = repo.inserisci("fasi", codice="X.01", descrizione="Massetti", impresa_id=a)
    cemento = repo.inserisci("sostanze", nome="Cemento Portland", codici_h="H315, H317, H318, H335",
                             avvertenza="Pericolo", data_sds="2024-03-15", dpi="Guanti nitrile, occhiali a tenuta")
    repo.salva_fasi_sostanza(cemento, [fase, altra, fase])
    assert repo.fasi_sostanza(cemento) == sorted([fase, altra])
    repo.inserisci("sostanze", impresa_id=a, nome="Disarmante", codici_h="H350")
    repo.inserisci("sostanze", nome="Orfano")
    assert [s["nome"] for s in sostanze_impresa(repo, a)] == ["Cemento Portland", "Disarmante"]
    pos = _testo_docx(genera_pos(repo, a, tmp_path / "pos.docx"))
    assert "Sostanze e miscele pericolose" in pos and "Cemento Portland" in pos
    assert "GHS05 Corrosivo" in pos and "Provoca gravi lesioni oculari" in pos
    assert "Guanti nitrile" in pos
    assert "X.01 Massetti" in pos  # tutte le lavorazioni in cui si usa
    elenco = _testo_docx(genera_elenco_sostanze(repo, a, tmp_path / "sostanze.docx"))
    assert "Cancerogeno o mutageno" in elenco
    messaggi = [s.messaggio for s in controlli.controlla_pos(repo)]
    assert any("Disarmante: cancerogeno" in m for m in messaggi)
    assert any("Disarmante: indicare la data della scheda" in m for m in messaggi)
    assert any("Orfano: collegare la sostanza" in m for m in messaggi)


def test_rumore_e_vibrazioni_nel_pos(repo, tmp_path):
    from sicurcantiere import controlli
    from sicurcantiere.documenti.pos import genera_pos

    a = _cantiere_completo(repo)
    messaggi = [s.messaggio for s in controlli.controlla_pos(repo)]
    assert any("eseguire la valutazione del rumore" in m for m in messaggi)
    r = repo.inserisci("valutazioni", metodo="rumore", impresa_id=a, titolo="Rumore muratori", mansione="Muratore")
    repo.salva_dati_valutazione(r, {"snr": 0, "compiti": [{"nome": "Betoniera", "laeq": 88, "picco": 120,
                                                            "durata": 240}]})
    v = repo.inserisci("valutazioni", metodo="vibrazioni", impresa_id=a, titolo="Vibrazioni demolizioni",
                       mansione="Manovale")
    repo.salva_dati_valutazione(v, {"sistema": "hav", "compiti": [{"nome": "Demolitore", "accelerazione": 8,
                                                                    "durata": 120}]})
    messaggi = [s.messaggio for s in controlli.controlla_pos(repo)]
    assert not any("eseguire la valutazione del rumore" in m for m in messaggi)
    testo = _testo_docx(genera_pos(repo, a, tmp_path / "pos.docx"))
    assert "Esito della valutazione del rumore" in testo and "85,0 dB(A)" in testo
    assert "Esito della valutazione delle vibrazioni" in testo and "4,00 m/s²" in testo
    assert "Tra il valore inferiore e il valore superiore di azione" in testo  # 84,99 dB(A)


def test_migrazione_figure_da_v4(tmp_path):
    import sqlite3

    f = tmp_path / "v4.spo"
    conn = sqlite3.connect(f)
    conn.executescript(";".join(MIGRAZIONI[:4]) + ";PRAGMA user_version = 4;")
    conn.execute("INSERT INTO imprese (ragione_sociale, rspp, rls, addetti_emergenze, posizione_inail_inps) "
                 "VALUES ('Alfa', 'Luca Verdi', '', 'Anna Blu', 'PZ 150413')")
    conn.commit()
    conn.close()
    p = Progetto(f)
    r = Repository(p)
    impresa = r.elenco("imprese")[0]
    assert impresa["inail"] == "PZ 150413"
    ruoli = {x["ruolo"]: x["nome"] for x in r.elenco("figure", impresa_id=impresa["id"])}
    assert ruoli == {"RSPP": "Luca Verdi", "Addetto al primo soccorso": "Anna Blu",
                     "Addetto antincendio ed evacuazione": "Anna Blu"}
    p.chiudi()


def test_committente_responsabile_dei_lavori(repo, tmp_path):
    from sicurcantiere import anagrafiche
    from sicurcantiere.documenti.notifica import genera_notifica

    _cantiere_completo(repo)
    assert not any(s["ruolo"] == "Responsabile dei lavori" for s in anagrafiche.soggetti(repo))
    repo.aggiorna_cantiere(committente_rl=1)
    elenco = anagrafiche.soggetti(repo)
    assert [s["ruolo"] for s in elenco][:2] == ["Committente", "Responsabile dei lavori"]
    assert elenco[1]["nome"] == "Comune di Bari"
    doc = Document(genera_notifica(repo, tmp_path / "n.docx"))
    voce = next(r for r in doc.tables[0].rows if r.cells[0].text.startswith("5."))
    assert voce.cells[1].text == "Comune di Bari"


def test_figure_e_dati_impresa_nel_pos(repo, tmp_path):
    from sicurcantiere import controlli
    from sicurcantiere.documenti.pos import genera_pos

    a = _cantiere_completo(repo)
    messaggi = [s.messaggio for s in controlli.controlla_pos(repo)]
    assert any(m.startswith("Alfa Costruzioni: indicare il rappresentante dei lavoratori") for m in messaggi)
    assert not any("responsabile del servizio" in m for m in messaggi)
    repo.aggiorna("imprese", a, inail="BA 123", cassa_edile="Cassa edile di Bari 7236",
                  ccnl="Edilizia - industria", durc_scadenza="2001-01-01", rspp_datore=1,
                  tipologia_lavori="Opere strutturali")
    b = repo.inserisci("imprese", ragione_sociale="Beta Impianti", tipo="Subappaltatrice", affidataria_id=a)
    repo.inserisci("figure", impresa_id=a, ruolo="RLST", nome="Paolo Gialli")
    messaggi = [s.messaggio for s in controlli.controlla_pos(repo)]
    assert "Alfa Costruzioni: DURC scaduto, acquisire il DURC in corso di validità." in messaggi
    assert not any(m.startswith("Alfa Costruzioni: indicare il rappresentante") for m in messaggi)
    testo = _testo_docx(genera_pos(repo, a, tmp_path / "pos.docx"))
    assert "BA 123" in testo and "Cassa edile di Bari 7236" in testo and "Edilizia - industria" in testo
    assert "Beta Impianti" in testo and "Opere strutturali" in testo and "Comune di Bari" in testo
    # Il datore di lavoro sostituisce l'RSPP inserito come figura.
    assert "svolti direttamente dal datore di lavoro" in testo and "Luca Verdi" not in testo
    assert "Paolo Gialli" in testo and "art. 50" in testo
    assert repo.leggi("imprese", b)["affidataria_id"] == a


def test_figure_svolte_dal_datore_o_non_presenti(repo, tmp_path):
    from sicurcantiere import anagrafiche, controlli
    from sicurcantiere.documenti.pos import genera_pos

    a = _cantiere_completo(repo)

    def figure_mancanti():
        return [s for s in controlli.controlla_pos(repo)
                if s.sezione == "Figure della sicurezza" and s.messaggio.startswith("Alfa Costruzioni: indicare")]

    prima = {s.messaggio.split(": indicare ")[1].split(" ")[1] for s in figure_mancanti()}
    assert "direttore" in prima and "capocantiere," in prima
    # ruolo scritto a mano in modo diverso dall'elenco: vale lo stesso
    repo.inserisci("figure", impresa_id=a, ruolo="RLS", nome="Anna Neri")
    repo.inserisci("figure", impresa_id=a, ruolo="addetto primo soccorso", nome="Anna Neri")
    repo.aggiorna("imprese", a, ruoli_datore=f"{anagrafiche.ANTINCENDIO}; {anagrafiche.PREPOSTO}",
                  figure_non_presenti=f"{anagrafiche.DIRETTORE_TECNICO}; {anagrafiche.CAPOCANTIERE}; "
                                      f"{anagrafiche.MEDICO}")
    assert not [s for s in figure_mancanti() if s.livello == controlli.ERRORE]
    assert not figure_mancanti()
    testo = _testo_docx(genera_pos(repo, a, tmp_path / "pos.docx"))
    assert "art. 34, comma 1-bis" in testo and "non è nominato" in testo
    assert "direzione dei lavori dell'impresa è svolta dal datore di lavoro" in testo


def test_archivio_anagrafiche(repo, tmp_path):
    from sicurcantiere.archivio_utente import IMPRESA, PERSONA, ArchivioAnagrafiche

    a = _cantiere_completo(repo)
    archivio = ArchivioAnagrafiche(tmp_path / "archivio.db")
    cse = next(s for s in repo.elenco("soggetti") if s["ruolo"].endswith("(CSE)"))
    archivio.salva_persona({**cse, "telefono": "333"})
    archivio.salva_persona({**cse, "telefono": "444"})  # stesso ruolo e nome: aggiornato
    voci = archivio.elenco(PERSONA)
    assert len(voci) == 1 and voci[0]["dati"]["telefono"] == "444"
    with pytest.raises(ValueError):
        archivio.salva_persona({"ruolo": "RSPP", "nome": " "})
    archivio.salva_impresa(repo, a)
    assert archivio.elenco(testo="alfa")[0]["tipo"] == IMPRESA

    # In un altro cantiere: la persona come figura dell'impresa, l'impresa con figure e lavoratori.
    altro = Repository(Progetto(tmp_path / "altro.spo", crea=True))
    nuova = archivio.inserisci_impresa(altro, archivio.elenco(IMPRESA)[0])
    assert altro.leggi("imprese", nuova)["datore_lavoro"] == "Mario Bianchi"
    assert [f["nome"] for f in altro.elenco("figure", impresa_id=nuova)] == ["Luca Verdi"]
    assert [l["nome"] for l in altro.elenco("lavoratori", impresa_id=nuova)] == ["Giuseppe Neri"]
    archivio.inserisci_persona(altro, voci[0], "figure", impresa_id=nuova, ruolo="RLST")
    figura = altro.elenco("figure", ruolo="RLST")[0]
    assert figura["nome"] == "Ing. Rossi" and figura["telefono"] == "444"
    archivio.elimina(voci[0]["id"])
    assert not archivio.elenco(PERSONA)
    archivio.chiudi()
    altro.progetto.chiudi()


def test_mansioni_tesserini_e_schede(repo, tmp_path):
    from sicurcantiere.documenti.maestranze import genera_schede_lavoratori, genera_tesserini, mansioni_lavoratore
    from sicurcantiere.documenti.pos import genera_pos

    a = _cantiere_completo(repo)
    repo.aggiorna("imprese", a, aut_subappalto="Prot. 123 del 01/10/2025")
    l = repo.elenco("lavoratori")[0]["id"]
    repo.aggiorna("lavoratori", l, cognome="Neri", nome="Giuseppe", matricola="M7", data_assunzione="2020-03-01",
                  dpi="casco; guanti antitaglio", form_base=1, form_dpi3=1)
    fase = repo.elenco("fasi")[0]
    attrezzature = repo.elenco("attrezzature", fase_id=fase["id"])
    gru = next(x for x in attrezzature if x["nome"] == "Gru a torre")
    repo.salva_mansioni(l, {(fase["id"], None), (fase["id"], gru["id"])})
    autonomo = repo.inserisci("lavoratori", impresa_id=a, autonomo=1, nome="Sara", cognome="Gialli",
                              mansione="Elettricista", partita_iva="0123")

    m = mansioni_lavoratore(repo, repo.leggi("lavoratori", l))
    assert [x["nome"] for x in m[0]["attrezzature"]] == ["Gru a torre"]
    fonti = {r["fonte"] for r in m[0]["rischi"]}
    assert "Gru a torre" in fonti and not fonti - {"", "Gru a torre"}
    assert [x["id"] for x in repo.lavoratori_fase(fase["id"])] == [l]

    testo = _testo_docx(genera_tesserini(repo, a, tmp_path / "t.docx"))
    assert "Giuseppe Neri" in testo and "M7" in testo and "01/03/2020" in testo
    assert "Prot. 123" in testo and "Lavoratore autonomo" in testo and "Sara Gialli" in testo
    testo = _testo_docx(genera_schede_lavoratori(repo, a, tmp_path / "s.docx"))
    assert "Scheda del lavoratore\nGiuseppe Neri" in testo and "Gru a torre" in testo
    assert "casco" in testo and "guanti antitaglio" in testo and "terza categoria" in testo
    testo = _testo_docx(genera_pos(repo, a, tmp_path / "pos.docx"))
    assert "Lavoratori impegnati" in testo and "Giuseppe Neri (Muratore)" in testo
    assert "Lavoratori autonomi: 1" in testo and "0123" in testo and "Prot. 123" in testo
    assert "casco; guanti antitaglio" in testo
    repo.elimina("lavoratori", autonomo)
    repo.elimina("fasi", fase["id"])
    assert not repo.mansioni(l)  # cancellate a cascata con la fase


def test_archivio_tipologie_e_lavorazioni_utente(repo, tmp_path):
    from sicurcantiere.archivio import TIPOLOGIE, per_categoria
    from sicurcantiere.archivio_utente import LAVORAZIONE, ArchivioAnagrafiche

    tutte = sum(len(v) for v in per_categoria().values())
    assert tutte == len(LAVORAZIONI)
    for t in TIPOLOGIE:
        gruppi = per_categoria(t)
        assert gruppi and all(t in l["tipologie"] for v in gruppi.values() for l in v)
    assert "Lavori stradali" not in per_categoria("Restauri")

    fase = repo.importa_lavorazione(lav("ST.07"))
    repo.aggiorna("fasi", fase, descrizione="Pilastri e travi in c.a. (personalizzata)")
    repo.inserisci("rischi", fase_id=fase, descrizione="Rischio aggiunto dall'utente", misure="Misura")
    archivio = ArchivioAnagrafiche(tmp_path / "a.db")
    archivio.salva_lavorazione(repo, fase)
    voce = archivio.elenco(LAVORAZIONE)[0]
    assert voce["nome"] == "ST.07 Pilastri e travi in c.a. (personalizzata)"
    altro = Repository(Progetto(tmp_path / "x.spo", crea=True))
    nuova = archivio.inserisci_lavorazione(altro, voce)
    assert altro.leggi("fasi", nuova)["descrizione"].endswith("(personalizzata)")
    assert len(altro.elenco("rischi", fase_id=nuova)) == len(repo.elenco("rischi", fase_id=fase))
    assert {a["nome"] for a in altro.elenco("attrezzature", fase_id=nuova)} == \
        {a["nome"] for a in repo.elenco("attrezzature", fase_id=fase)}
    archivio.chiudi()
    altro.progetto.chiudi()


def test_struttura_fase_e_schede_d_uso(repo, tmp_path):
    from sicurcantiere.archivio_attrezzature import SCHEDE
    from sicurcantiere.documenti.pos import genera_pos
    from sicurcantiere.struttura import struttura_fase

    assert set(SCHEDE) <= set(ATTREZZATURE)
    schede = [a["scheda"] for a in ATTREZZATURE.values()]
    assert all(s["prima"] and s["durante"] and s["dopo"] and s["dpi"] for s in schede)
    a = _cantiere_completo(repo)
    fase = repo.importa_lavorazione(lav("AL.11"))
    repo.aggiorna("fasi", fase, impresa_id=a, modalita="Squadra di due addetti", turni="7.30-16.30")
    f = repo.leggi("fasi", fase)
    assert f["lavorazione"] == "Preparazione delle aree di cantiere"
    assert f["addetto"].startswith("Addetto al taglio") and f["dettaglio"]
    motosega = next(x for x in repo.elenco("attrezzature", fase_id=fase) if x["nome"] == "Motosega")
    assert motosega["misure"].startswith("Prima dell'uso: 1)") and "pantaloni antitaglio" in motosega["dpi"]

    s = struttura_fase(repo, fase)
    assert [m["nome"] for m in s["macchine"]] == ["Trattore"]
    assert {x["nome"] for x in s["attrezzi"]} == {"Motosega", "Utensili manuali"}
    assert any("abbattimento" in r["descrizione"] for r in s["rischi_lavorazione"])
    assert all(r["fonte"] == "Motosega" for x in s["attrezzi"] if x["nome"] == "Motosega" for r in x["rischi"])

    testo = _testo_docx(genera_pos(repo, a, tmp_path / "pos.docx"))
    assert "La lavorazione è suddivisa nelle seguenti fasi" in testo
    assert "Taglio di alberi ed estirpazione delle ceppaie (fase)" in testo
    assert "Squadra di due addetti" in testo and "7.30-16.30" in testo
    assert "Macchine utilizzate nelle lavorazioni" in testo and "Prima dell'uso" in testo
    assert "DPI dell'operatore" in testo


def test_archivio_esteso_verificato():
    import importlib

    from sicurcantiere.archivio import MODULI_ATTREZZATURE, MODULI_FASI
    from sicurcantiere.archivio_dati.base import CATALOGO
    from sicurcantiere.archivio_dati.verifica import verifica

    for nome in MODULI_ATTREZZATURE + MODULI_FASI:
        assert verifica(importlib.import_module(f"sicurcantiere.archivio_dati.{nome}")) == [], nome
    assert set(CATALOGO) <= set(ATTREZZATURE)
    assert all(l["categoria"] in CATEGORIE for l in LAVORAZIONI)
    assert all(a in ATTREZZATURE for l in LAVORAZIONI for a in l["attrezzature"])


def test_migrazione_sostanze_su_piu_fasi(tmp_path):
    import sqlite3

    f = tmp_path / "v7.spo"
    conn = sqlite3.connect(f)
    conn.executescript(";".join(MIGRAZIONI[:7]) + ";PRAGMA user_version = 7;")
    conn.execute("INSERT INTO fasi (id, codice) VALUES (7, 'A')")
    conn.execute("INSERT INTO sostanze (nome, fase_id) VALUES ('Malta', 7)")
    conn.execute("INSERT INTO sostanze (nome) VALUES ('Vernice')")
    conn.commit()
    conn.close()
    p = Progetto(f)
    r = Repository(p)
    malta, vernice = r.elenco("sostanze")
    assert r.fasi_sostanza(malta["id"]) == [7] and r.fasi_sostanza(vernice["id"]) == []
    r.elimina("fasi", 7)
    assert r.fasi_sostanza(malta["id"]) == []
    p.chiudi()


def test_segnaletica_e_impaginazione_pos(repo, tmp_path):
    from docx import Document
    from docx.oxml.ns import qn

    from sicurcantiere import segnaletica
    from sicurcantiere.documenti.pos import genera_pos

    assert all(s.file.is_file() for s in segnaletica.SEGNALI.values())
    assert segnaletica.icone_dpi(["Casco con sottogola", "guanti antitaglio", "otoprotettori"]) == [
        "casco", "guanti", "udito"]
    chiavi = [s.chiave for s in segnaletica.segnali_cantiere(["Caduta dall'alto", "Elettrocuzione"], [])]
    assert chiavi[0] == "accesso" and {"caduta", "tensione", "imbracatura"} <= set(chiavi)
    assert chiavi.index("tensione") < chiavi.index("casco") < chiavi.index("soccorso")  # divieti, avvertimenti, obblighi

    a = _cantiere_completo(repo)
    repo.salva_testo("foto_cantiere", f"{segnaletica.SEGNALI['casco'].file}\n{tmp_path / 'manca.jpg'}")
    doc = Document(genera_pos(repo, a, tmp_path / "pos.docx"))
    testo = "\n".join(p.text for p in doc.paragraphs)
    assert "Foto 1 - area del cantiere" in testo and "Foto 2" not in testo  # foto mancante saltata
    titoli = [p.text for p in doc.paragraphs if p.style.name == "Heading 1"]
    # capitoli numerati in ordine: «1   Dati generali del lavoro», «2   Committente...»
    assert titoli[0] == "1   Dati generali del lavoro" and titoli[-1].endswith("Allegati e firme")
    assert [int(t.split()[0]) for t in titoli] == list(range(1, len(titoli) + 1))
    assert any(t.endswith("Segnaletica prevista nel cantiere") for t in titoli)
    # immagini: foto, cartelli e icone dei DPI
    assert len(doc.inline_shapes) > 10
    # la tabella dei rischi ha colonne di larghezza fissa e il livello colorato
    rischi = next(t for t in doc.tables if t.rows[0].cells[0].text == "Rischio")
    griglia = [int(c.get(qn("w:w"))) for c in rischi._tbl.tblGrid.findall(qn("w:gridCol"))]
    assert griglia[2] > griglia[0] > griglia[1]
    assert rischi.rows[1].cells[1]._tc.tcPr.find(qn("w:shd")) is not None


def test_valutazioni_su_piu_lavorazioni_da_file_vecchio(tmp_path):
    import sqlite3

    f = tmp_path / "v8.spo"
    conn = sqlite3.connect(f)
    conn.executescript(";".join(MIGRAZIONI[:8]) + ";PRAGMA user_version = 8;")
    conn.execute("INSERT INTO fasi (id, codice) VALUES (7, 'A')")
    conn.execute("INSERT INTO valutazioni (titolo, fase_id) VALUES ('Sacchi', 7)")
    conn.execute("INSERT INTO valutazioni (titolo) VALUES ('Rumore')")
    conn.commit()
    conn.close()
    p = Progetto(f)
    r = Repository(p)
    sacchi, rumore = r.elenco("valutazioni")
    assert r.fasi_valutazione(sacchi["id"]) == [7] and r.fasi_valutazione(rumore["id"]) == []
    r.elimina("fasi", 7)
    assert r.fasi_valutazione(sacchi["id"]) == []
    p.chiudi()


def test_valutazioni_automatiche_dai_rischi_e_dalle_attrezzature(repo, tmp_path):
    from sicurcantiere import controlli
    from sicurcantiere.archivio import ATTREZZATURE
    from sicurcantiere.documenti.pos import genera_pos
    from sicurcantiere.valutazioni import automatiche

    # Schede dell'archivio completate con i rischi specifici mancanti.
    assert "Vibrazioni al corpo intero" in [x["descrizione"] for x in ATTREZZATURE["Escavatore"]["rischi"]]
    assert "Rumore" in [x["descrizione"] for x in ATTREZZATURE["Autocarro"]["rischi"]]

    a = repo.inserisci("imprese", ragione_sociale="Alfa", datore_lavoro="Mario")
    getti = repo.inserisci("fasi", codice="A", descrizione="Getti", impresa_id=a)
    tagli = repo.inserisci("fasi", codice="B", descrizione="Tagli", impresa_id=a)
    repo.inserisci("rischi", fase_id=getti, descrizione="Movimentazione manuale dei carichi")
    # Attrezzo scritto a mano, fuori archivio: i rischi specifici si riconoscono dal nome.
    repo.inserisci("attrezzature", fase_id=tagli, tipo="Attrezzo", nome="Flessibile a batteria")
    messaggi = automatiche.sincronizza(repo)
    assert any("Flessibile a batteria" in m for m in messaggi)
    rischi_tagli = {(r["descrizione"], r["fonte"]) for r in repo.elenco("rischi", fase_id=tagli)}
    assert {("Rumore", "Flessibile a batteria"), ("Vibrazioni al sistema mano-braccio",
                                                    "Flessibile a batteria")} <= rischi_tagli

    valutazioni = {v["metodo"]: v for v in repo.elenco("valutazioni")}
    assert set(valutazioni) == {"niosh", "rumore", "vibrazioni"}
    assert repo.fasi_valutazione(valutazioni["niosh"]["id"]) == [getti]
    assert repo.fasi_valutazione(valutazioni["rumore"]["id"]) == [tagli]
    assert valutazioni["rumore"]["impresa_id"] == a
    assert all(r["valutazione_id"] for r in repo.elenco("rischi"))
    assert automatiche.sincronizza(repo) == []  # niente da rifare

    # Una nuova fase con lo stesso rischio si aggiunge alla valutazione esistente.
    massetti = repo.inserisci("fasi", codice="C", descrizione="Massetti", impresa_id=a)
    repo.inserisci("rischi", fase_id=massetti, descrizione="Movimentazione manuale dei carichi")
    assert automatiche.fasi_scoperte(repo) == {"mmc": [massetti]}
    automatiche.sincronizza(repo)
    assert repo.fasi_valutazione(valutazioni["niosh"]["id"]) == [getti, massetti]
    assert len(repo.elenco("valutazioni")) == 3
    # Togliendo la fase a mano, la sincronizzazione la rimette: nessuna fase resta scoperta.
    repo.salva_fasi_valutazione(valutazioni["niosh"]["id"], [getti])
    repo.aggiorna("rischi", repo.elenco("rischi", fase_id=massetti)[0]["id"], valutazione_id=None)
    automatiche.sincronizza(repo)
    assert massetti in repo.fasi_valutazione(valutazioni["niosh"]["id"])

    # Finché non si inseriscono i parametri la valutazione è «da eseguire».
    messaggi = [s.messaggio for s in controlli.controlla_pos(repo)]
    assert any("Movimentazione manuale dei carichi" in m and "da eseguire" in m for m in messaggi)
    testo = _testo_docx(genera_pos(repo, a, tmp_path / "pos.docx"))
    assert "Valutazione da eseguire" in testo and "A Getti; C Massetti" in testo
    repo.salva_dati_valutazione(valutazioni["niosh"]["id"], {"peso": 10})
    assert not automatiche.da_eseguire(repo, valutazioni["niosh"]["id"])


def test_rischio_non_presente_con_motivazione(repo, tmp_path):
    from sicurcantiere import controlli
    from sicurcantiere.documenti.pos import genera_pos
    from sicurcantiere.valutazioni import automatiche

    a = _cantiere_completo(repo)
    fase = repo.inserisci("fasi", codice="T", descrizione="Tagli", impresa_id=a)
    repo.inserisci("attrezzature", fase_id=fase, tipo="Attrezzo", nome="Flessibile a batteria")
    automatiche.sincronizza(repo)
    vibrazioni = next(r for r in repo.elenco("rischi", fase_id=fase) if r["descrizione"].startswith("Vibrazioni"))
    v = vibrazioni["valutazione_id"]
    assert fase in repo.fasi_valutazione(v)
    repo.aggiorna("rischi", vibrazioni["id"], non_presente=1, valutazione_id=None)
    repo.salva_fasi_valutazione(v, [f for f in repo.fasi_valutazione(v) if f != fase])
    automatiche.sincronizza(repo)
    assert fase not in repo.fasi_valutazione(v)  # non viene rimessa
    assert len([r for r in repo.elenco("rischi", fase_id=fase) if r["descrizione"].startswith("Vibrazioni")]) == 1
    assert vibrazioni["id"] not in {r["id"] for r in repo.rischi(fase_id=fase)}
    messaggi = [s.messaggio for s in controlli.controlla_pos(repo)]
    assert any("motivare l'esclusione del rischio" in m for m in messaggi)
    repo.aggiorna("rischi", vibrazioni["id"], motivazione="Utensile a basse vibrazioni, uso inferiore a 10 minuti")
    messaggi = [s.messaggio for s in controlli.controlla_pos(repo)]
    assert not any("motivare l'esclusione del rischio" in m for m in messaggi)
    testo = _testo_docx(genera_pos(repo, a, tmp_path / "pos.docx"))
    assert "Rischi valutati come non presenti" in testo and "uso inferiore a 10 minuti" in testo


def test_autoverifica(tmp_path, monkeypatch):
    pytest.importorskip("PySide6.QtWidgets")
    monkeypatch.setenv("QT_QPA_PLATFORM", "offscreen")
    from PySide6.QtWidgets import QApplication

    from sicurcantiere import autoverifica

    app = QApplication.instance() or QApplication([])  # noqa: F841
    assert autoverifica.esegui(tmp_path, "pos") == 0
    esito = (tmp_path / "autoverifica.txt").read_text(encoding="utf-8")
    assert "ESITO: OK" in esito and (tmp_path / "prova POS.docx").is_file()


def test_prepara_pos_dal_psc(repo, tmp_path):
    from sicurcantiere import controlli
    from sicurcantiere.documenti.pos import genera_pos
    from sicurcantiere.prepara_pos import CHIAVE_ORIGINE, prepara_pos

    a = _cantiere_completo(repo)
    b = repo.inserisci("imprese", ragione_sociale="Beta Impianti", tipo="Esecutrice", datore_lavoro="Ugo Blu")
    sub = repo.inserisci("imprese", ragione_sociale="Gamma Sub", tipo="Subappaltatrice", affidataria_id=a)
    repo.inserisci("lavoratori", impresa_id=b, nome="Ivo", cognome="Gialli", mansione="Elettricista")
    fase_a = repo.elenco("fasi", impresa_id=a)[0]
    fase_b = repo.inserisci("fasi", codice="IMP", descrizione="Impianto elettrico", impresa_id=b,
                            data_inizio=fase_a["data_inizio"], data_fine=fase_a["data_fine"])
    repo.inserisci("rischi", fase_id=fase_b, descrizione="Elettrocuzione", probabilita=2, danno=3, misure="x")
    repo.salva_prescrizione(*sorted((fase_a["id"], fase_b)), "Sfasare le lavorazioni sullo stesso piano.")
    repo.inserisci("costi", descrizione="Recinzione", quantita=1, prezzo_unitario=100)

    file = prepara_pos(repo, a, tmp_path / "pos alfa.spo")
    pos = Repository(Progetto(file))
    try:
        assert {i["ragione_sociale"] for i in pos.elenco("imprese")} == {"Alfa Costruzioni", "Gamma Sub"}
        assert all(f["impresa_id"] == a for f in pos.elenco("fasi"))
        assert not pos.elenco("costi") and not pos.prescrizioni()
        assert [l["cognome"] for l in pos.elenco("lavoratori")] == [l["cognome"] for l in repo.elenco(
            "lavoratori", impresa_id=a)]
        procedure = pos.leggi("fasi", fase_a["id"])["procedure_psc"]
        assert "Impianto elettrico" in procedure and "Beta Impianti" in procedure and "Sfasare" in procedure
        assert pos.testo(CHIAVE_ORIGINE).startswith("Preparato dal PSC")
        assert pos.cantiere()["nome"] == repo.cantiere()["nome"] and pos.elenco("soggetti")
        # la subappaltatrice è solo di riferimento: niente errori per lei nei controlli del POS
        assert not any("Gamma Sub" in s.messaggio for s in controlli.controlla_pos(pos))
        testo = _testo_docx(genera_pos(pos, a, tmp_path / "pos.docx"))
        assert "Sfasare le lavorazioni" in testo and "Gamma Sub" in testo
    finally:
        pos.progetto.chiudi()
    assert len(repo.elenco("imprese")) == 3 and repo.prescrizioni()  # il PSC resta com'era
    assert sub


def test_codice_automatico_delle_fasi(repo):
    from sicurcantiere import archivio

    assert repo.codice_automatico() == "UT.01"
    a = repo.inserisci("fasi", codice=repo.codice_automatico(), descrizione="Fase a mano")
    assert repo.codice_automatico() == "UT.02"
    gruppo = archivio.LAVORAZIONI[0]["gruppo"]  # «Preparazione delle aree di cantiere», sigla AL
    codice = repo.codice_automatico(gruppo, escludi=a)
    assert codice.startswith("AL.") and codice not in {l["codice"] for l in archivio.LAVORAZIONI}
    # lavorazione nuova: prende la sigla delle altre fasi con la stessa lavorazione
    repo.inserisci("fasi", codice="XY.03", descrizione="b", lavorazione="Opere speciali")
    assert repo.codice_automatico("Opere speciali") == "XY.04"
    assert repo.codice_automatico("Lavorazione mai vista") == "UT.02"


def test_psc_completo_allegato_xv(repo, tmp_path):
    import json

    from PySide6.QtGui import QColor, QImage

    from sicurcantiere import relazione

    _cantiere_completo(repo)
    repo.inserisci("soggetti", ruolo="Coordinatore per la progettazione (CSP)", nome="Geom. Bianchi",
                   indirizzo="Via Bologna 15, Lavello", telefono="333 1234567", data_incarico="2025-06-25")
    repo.salva_testo("contesto", "Condominio abitato in zona urbana.")
    repo.salva_testo("org_servizi", "Servizi igienici dell'appartamento.")
    repo.salva_testo(relazione.CHIAVE_ALLEGATO_XI, json.dumps(
        {"verificato": True, "voci": {"4": {"presente": True, "misure": "Distanza minima dalle linee aeree"}}}))
    repo.salva_testo(relazione.CHIAVE_ORDIGNI, json.dumps(
        {"scavi": "Sì", "esito": relazione.ESITI_ORDIGNI[1], "motivazione": "Area già scavata negli anni '80"}))
    repo.inserisci("costi", categoria="a", descrizione="Recinzione", unita="m", quantita=100, prezzo_unitario=10)
    repo.inserisci("costi", categoria="c", descrizione="Impianto di terra", unita="cad", quantita=1,
                   prezzo_unitario=300)
    immagine = QImage(400, 300, QImage.Format.Format_RGB32)
    immagine.fill(QColor("#88AACC"))
    immagine.save(str(tmp_path / "planimetria.png"))
    repo.inserisci("allegati_psc", tipo="Planimetria di cantiere", titolo="Layout", file=str(tmp_path / "planimetria.png"))
    f = repo.elenco("fasi")[0]
    a = repo.inserisci("attrezzature", fase_id=f["id"], tipo="Attrezzo", nome="Martello demolitore",
                       potenza_sonora=113)

    doc = Document(genera_psc(repo, tmp_path / "psc.docx"))
    testo = _testo_docx(tmp_path / "psc.docx")
    for atteso in ("Condominio abitato in zona urbana.", "Servizi igienici dell'appartamento.",
                   "Recinzione, accessi e segnalazioni", "Lavori comportanti rischi particolari",
                   "Distanza minima dalle linee aeree", "Area già scavata negli anni '80",
                   "a) Apprestamenti previsti nel PSC", "Totale della categoria", "€ 1.300,00",
                   "Organigramma del cantiere", "Documentazione da custodire in cantiere",
                   "Rischi individuati e lavorazioni in cui compaiono", "113,0 dB(A)", "Tavole allegate",
                   "Via Bologna 15, Lavello", "25/06/2025", "Procedure complementari"):
        assert atteso in testo, atteso
    assert "[Da compilare: planimetria" not in testo
    assert len(doc.inline_shapes) >= 3  # organigramma, Gantt, planimetria
    assert a


def test_migrazione_sezione_area(tmp_path):
    import sqlite3

    from sicurcantiere.db import MIGRAZIONI

    percorso = tmp_path / "vecchio.scs"
    conn = sqlite3.connect(percorso)
    n = next(k for k, m in enumerate(MIGRAZIONI) if "area_caratteristiche" in m)  # prima della migrazione
    conn.executescript(";".join(MIGRAZIONI[:n]) + f";PRAGMA user_version = {n};")
    conn.execute("INSERT INTO relazione (sezione, testo) VALUES ('area', 'Area pianeggiante.')")
    conn.commit()
    conn.close()
    r = Repository(Progetto(percorso))
    assert r.testo("area_caratteristiche") == "Area pianeggiante."


def test_tavola_pdf_nel_psc(repo, tmp_path):
    from PySide6.QtGui import QPageSize, QPainter, QPdfWriter

    from sicurcantiere.documenti.immagini import applicazione

    applicazione()
    pdf = tmp_path / "tavola.pdf"
    scrittore = QPdfWriter(str(pdf))
    scrittore.setPageSize(QPageSize(QPageSize.PageSizeId.A4))
    pittore = QPainter(scrittore)
    pittore.drawText(200, 200, "Planimetria")
    scrittore.newPage()
    pittore.drawText(200, 200, "Sezione")
    pittore.end()
    repo.inserisci("allegati_psc", tipo="Planimetria di cantiere", titolo="Planimetria", file=str(pdf))
    doc = Document(genera_psc(repo, tmp_path / "psc.docx"))
    testo = _testo_docx(tmp_path / "psc.docx")
    assert "tavola.pdf · pagina 2 di 2" in testo and len(doc.inline_shapes) >= 2


def test_patente_crediti_e_idoneita(repo, tmp_path):
    import json

    from sicurcantiere import anagrafiche, controlli
    from sicurcantiere.documenti.idoneita import genera_verifica_idoneita

    a = _cantiere_completo(repo)
    autonomo = repo.inserisci("imprese", ragione_sociale="Rossi Mario", tipo="Lavoratore autonomo")
    messaggi = [s.messaggio for s in controlli.controlla(repo)]
    assert any("Alfa Costruzioni: indicare la situazione della patente" in m for m in messaggi)
    assert any("Alfa Costruzioni: eseguire la verifica dell'idoneità" in m for m in messaggi)

    repo.aggiorna("imprese", a, patente_stato=anagrafiche.PATENTE_POSSEDUTA, patente_numero="PC-123",
                  patente_crediti=10)
    errori = [s.messaggio for s in controlli.controlla(repo) if s.livello == controlli.ERRORE]
    assert any("meno di 15 crediti" in m for m in errori)
    repo.aggiorna("imprese", a, patente_crediti=30, idoneita_verificata="2025-09-20", idoneita=json.dumps(
        {k: {"esito": "Acquisito", "nota": ""} for k, _ in anagrafiche.IDONEITA_IMPRESE}))
    impresa = repo.leggi("imprese", a)
    assert anagrafiche.patente_regolare(impresa) and anagrafiche.esito_idoneita(impresa) == "Idonea"
    assert anagrafiche.descrizione_patente(impresa) == "Posseduta: n. PC-123, 30 crediti"
    messaggi = [s.messaggio for s in controlli.controlla(repo)]
    assert not any(m.startswith("Alfa Costruzioni: indicare la situazione") or
                   m.startswith("Alfa Costruzioni: eseguire la verifica dell'idoneità") for m in messaggi)

    testo = _testo_docx(genera_verifica_idoneita(repo, a, tmp_path / "v.docx"))
    assert "Verifica dell'idoneità tecnico-professionale" in testo and "Idoneità tecnico-professionale verificata" in testo
    assert "Posseduta: n. PC-123, 30 crediti" in testo and "20/09/2025" in testo
    # lavoratore autonomo: le voci del punto 2 dell'Allegato XVII
    testo = _testo_docx(genera_verifica_idoneita(repo, autonomo, tmp_path / "a.docx"))
    assert "Allegato XVII punto 2" in testo and "Documentazione incompleta" not in testo
    assert "PC-123" in _testo_docx(genera_psc(repo, tmp_path / "psc.docx"))


def test_deroga_art_90_comma_11():
    assert analisi.deroga_art_90_c11({"importo_lavori": 60000, "titolo_abilitativo": "CILA"})
    assert not analisi.deroga_art_90_c11({"importo_lavori": 60000, "titolo_abilitativo": "Permesso di costruire"})
    assert not analisi.deroga_art_90_c11({"importo_lavori": 60000, "titolo_abilitativo": "PdC n. 12"})
    assert not analisi.deroga_art_90_c11({"importo_lavori": 150000, "titolo_abilitativo": "SCIA"})
    assert not analisi.deroga_art_90_c11({"importo_lavori": 60000, "lavori_pubblici": 1})
    assert not analisi.deroga_art_90_c11({"importo_lavori": 0})


def test_psc_rischi_2_2_3_trasmissione_e_aggiornamento(repo, tmp_path):
    import json

    from sicurcantiere import controlli, relazione

    a = _cantiere_completo(repo)
    repo.aggiorna("imprese", a, psc_ricevuto_data="2025-09-10")
    repo.inserisci("imprese", ragione_sociale="Giallo Luigi", tipo="Lavoratore autonomo")
    f = repo.elenco("fasi")[0]
    repo.inserisci("rischi", fase_id=f["id"], descrizione="Caduta dall'alto", probabilita=2, danno=3,
                   misure="Parapetti")
    messaggi = [x.messaggio for x in controlli.controlla(repo)]
    assert any("2.2.3" in m for m in messaggi)

    repo.salva_testo(relazione.CHIAVE_RISCHI_LAVORAZIONI, json.dumps(
        {"verificato": True, "voci": {"l": {"presente": True, "misure": ""}}}))
    messaggi = [x.messaggio for x in controlli.controlla(repo)]
    assert any("lettera c)" in m and "non è indicato" in m for m in messaggi)  # caduta dall'alto nelle fasi
    assert any("lettera l)" in m and "indicare le scelte" in m for m in messaggi)

    testo = _testo_docx(genera_psc(repo, tmp_path / "psc.docx"))
    for atteso in ("Rischi da analizzare nelle lavorazioni", "Seppellimento negli scavi", "Aggiornamento del piano",
                   "Trasmissione e presa visione del piano", "Lavoratore autonomo", "Giallo Luigi", "10/09/2025",
                   "proposte di integrazione"):
        assert atteso in testo, atteso


def test_psc_redatto_dal_cse_art_90_comma_11(repo, tmp_path):
    from sicurcantiere import controlli

    _cantiere_completo(repo)
    repo.aggiorna_cantiere(importo_lavori=80000, titolo_abilitativo="CILA")
    b = repo.inserisci("imprese", ragione_sociale="Beta Impianti")
    f = repo.importa_lavorazione(lav("MT.01"))
    repo.aggiorna("fasi", f, impresa_id=b, data_inizio="2025-10-06", data_fine="2025-10-20")
    messaggi = [x.messaggio for x in controlli.controlla(repo)]
    assert not any("coordinatore per la progettazione" in m for m in messaggi)
    testo = _testo_docx(genera_psc(repo, tmp_path / "psc.docx"))
    assert "art. 90, comma 11" in testo


def test_contenuti_planimetria(repo, tmp_path):
    from PySide6.QtGui import QColor, QImage

    from sicurcantiere import controlli, relazione

    _cantiere_completo(repo)
    immagine = QImage(100, 80, QImage.Format.Format_RGB32)
    immagine.fill(QColor("#88AACC"))
    immagine.save(str(tmp_path / "p.png"))
    repo.inserisci("allegati_psc", tipo="Planimetria di cantiere", titolo="Layout", file=str(tmp_path / "p.png"))
    assert any("planimetria di cantiere riporti" in x.messaggio for x in controlli.controlla(repo))
    essenziali = [v for v, e in relazione.CONTENUTI_PLANIMETRIA if e]
    repo.salva_testo(relazione.CHIAVE_CONTENUTI_PLANIMETRIA, ";".join(essenziali))
    assert not any("planimetria di cantiere riporti" in x.messaggio for x in controlli.controlla(repo))
    assert "La planimetria di cantiere riporta:" in _testo_docx(genera_psc(repo, tmp_path / "psc.docx"))
