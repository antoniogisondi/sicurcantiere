"""Generazione del Piano di Sicurezza e Coordinamento (art. 100 e Allegato XV, punto 2.1.2).

Stessa impaginazione del POS, nei colori del programma PSC. I capitoli seguono le lettere dei contenuti
minimi del punto 2.1.2: a) opera, b) soggetti, c-d) rischi e misure per area, organizzazione e
lavorazioni, e) interferenze, f) uso comune, g) cooperazione, h) emergenze, i) cronoprogramma,
l) costi della sicurezza.
"""

from __future__ import annotations

from datetime import date, timedelta
from pathlib import Path
from typing import Any

from docx import Document
from docx.shared import Pt

from sicurcantiere import anagrafiche, analisi, segnaletica
from sicurcantiere.documenti import stile
from sicurcantiere.documenti.comune import data_it, euro, indirizzo_cantiere, nome_fase, salva
from sicurcantiere.documenti.composizione import (
    CONSIGLIATO,
    FACOLTATIVO,
    OBBLIGATORIO,
    Manca,
    blocco,
    segnaposto,
)
from sicurcantiere.documenti.pos import (
    DOCUMENTI_IN_CANTIERE,
    NUMERI_EMERGENZA,
    TIPI_ATTREZZATURA,
    dpi_con_icone,
    foto_cantiere,
    schede_attrezzature,
    tabella_rischi,
)
from sicurcantiere.relazione import SEZIONI, SEZIONI_DETTAGLIO
from sicurcantiere.repository import Repository

MAX_COLONNE_CRONO = 22
PRESCRIZIONE_GENERICA = (
    "Sfasamento spaziale delle lavorazioni o, se non possibile, sfasamento temporale; "
    "verifica da parte del CSE prima dell'inizio delle lavorazioni."
)
# Sezioni dei programmi in cui si compilano i dati (per i collegamenti dei dati mancanti).
CANTIERE, SOGGETTI, IMPRESE, FASI, RISCHI, INTERFERENZE, COSTI, RELAZIONE = (
    "Dati generali", "Soggetti", "Imprese", "Fasi", "Rischi", "Interferenze", "Costi", "Relazione")
RISCHI_PARTICOLARI, TAVOLE = "Rischi particolari", "Tavole"
TIPO_FOTO = "Foto o mappa del cantiere"


def _o(valore, dato: str, sezione: str, livello: str = OBBLIGATORIO):
    return valore if str(valore or "").strip() and valore != "-" else Manca(dato, sezione, livello)


def _livello(prodotto: int) -> int:
    """Indice del colore del livello P x D (come nelle tabelle dei rischi)."""
    return 3 if prodotto >= 9 else 2 if prodotto >= 6 else 1 if prodotto >= 3 else 0


def _colonne_cronoprogramma(inizio: date, fine: date) -> list[tuple[date, date, str]]:
    """Divide il periodo in settimane o, se troppo lungo, in periodi di più settimane."""
    inizio = inizio - timedelta(days=inizio.weekday())  # lunedì
    settimane = ((fine - inizio).days // 7) + 1
    passo = max(1, -(-settimane // MAX_COLONNE_CRONO))
    colonne = []
    giorno = inizio
    while giorno <= fine:
        termine = giorno + timedelta(days=7 * passo - 1)
        colonne.append((giorno, termine, giorno.strftime("%d/%m")))
        giorno = termine + timedelta(days=1)
    return colonne


def _cronoprogramma(doc, fasi: list[dict[str, Any]], imprese: dict[int, str],
                    interferenze: list[analisi.Interferenza]) -> None:
    """Diagramma a barre settimanale: barre nel colore del programma, in rosso le fasi che interferiscono."""
    periodo = analisi.periodo_lavori(fasi)
    if not periodo:
        stile.testo(doc, "", "date di inizio e fine delle fasi, per il cronoprogramma dei lavori", FASI,
                    OBBLIGATORIO)
        return
    colonne = _colonne_cronoprogramma(*periodo)
    interferenti = {f["id"] for i in interferenze for f in (i.fase_a, i.fase_b)}
    t = doc.add_table(rows=1, cols=len(colonne) + 2)
    stile._bordi_tabella(t, top=(stile.ACCENTO, 8), bottom=(stile.LINEA_FORTE, 6), insideH=(stile.LINEA, 4),
                         insideV=(stile.LINEA, 2))
    stile._margini_celle(t, 0.04, 0.05)
    stile._ripeti_intestazione(t.rows[0])
    for cella, testo_cella in zip(t.rows[0].cells, ["Fase", "Impresa"] + [c[2] for c in colonne]):
        stile.ombreggia_cella(cella, stile.TINTA)
        p = cella.paragraphs[0]
        stile._spaziatura(p)
        stile._run(p, testo_cella, 6 if testo_cella[:1].isdigit() else 7.5, stile.ACCENTO_SCURO, grassetto=True)
    for f in fasi:
        inizio, fine = analisi.data(f["data_inizio"]), analisi.data(f["data_fine"])
        riga = t.add_row()
        stile._non_spezzare(riga)
        celle = riga.cells
        for cella, testo_cella in zip(celle[:2], (nome_fase(f), imprese.get(f["impresa_id"], ""))):
            p = cella.paragraphs[0]
            stile._spaziatura(p)
            stile._run(p, testo_cella, 7)
        if not (inizio and fine):
            continue
        colore = stile.LIVELLI[3][1] if f["id"] in interferenti else stile.ACCENTO
        for cella, (da, a, _) in zip(celle[2:], colonne):
            if inizio <= a and fine >= da:
                stile.ombreggia_cella(cella, colore)
    resto = stile.LARGHEZZA_UTILE - 4.6 - 2.8
    stile.larghezze_colonne(t, [4.6, 2.8] + [resto / len(colonne)] * len(colonne))
    stile._spazio_dopo(doc)
    stile.nota(doc, "Ogni colonna è una settimana (o un gruppo di settimane per i lavori lunghi), indicata con la "
                    "data del lunedì. In rosso le fasi che si sovrappongono a quelle di altre imprese "
                    "(capitolo «Interferenze tra le lavorazioni»).")


def _programma(doc, fasi: list[dict[str, Any]], imprese: dict[int, str]) -> None:
    """Programma dei lavori in tabella: tutte le fasi di tutte le imprese, dall'inizio alla fine."""
    stile.banda(doc, "Programma dei lavori")
    ordinate = sorted(fasi, key=lambda f: (f["data_inizio"] is None, f["data_inizio"] or "", f["id"]))
    righe = []
    for f in ordinate:
        da, a = analisi.data(f["data_inizio"]), analisi.data(f["data_fine"])
        righe.append([nome_fase(f) + (f"\n{f['lavorazione']}" if f["lavorazione"] else ""),
                      imprese.get(f["impresa_id"], "") or "da definire", f["zona"] or "—",
                      data_it(da) if da else "—", data_it(a) if a else "—",
                      str(analisi.giorni_lavorativi(da, a)) if da and a and a >= da else "—",
                      str(f["lavoratori"]) if f.get("lavoratori") else "—"])
    stile.tabella_colorata(doc, ["Fase / lavorazione", "Impresa", "Zona", "Inizio", "Fine", "gg lav.", "Lav."],
                           righe, 8, larghezze=[5.6, 3.2, 1.8, 2.0, 2.0, 1.4, 1.2])


def _diagramma_gantt(doc, fasi: list[dict[str, Any]], imprese: dict[int, str], livelli: dict[int, int],
                     interferenze: list[analisi.Interferenza]) -> None:
    """Diagramma di Gantt su pagine orizzontali; senza Qt (non dovrebbe capitare) resta la tabella."""
    import io

    from docx.shared import Cm

    from sicurcantiere.documenti import gantt

    if not analisi.periodo_lavori(fasi):
        _cronoprogramma(doc, fasi, imprese, interferenze)
        return
    try:
        pagine = gantt.immagini(fasi, imprese, livelli)
    except Exception:  # noqa: BLE001 - il documento si genera comunque, con la tabella settimanale
        pagine = []
    if not pagine:
        _cronoprogramma(doc, fasi, imprese, interferenze)
        return
    stile.pagine_orizzontali(doc)
    for n, immagine in enumerate(pagine):
        p = doc.add_paragraph()
        if n:
            p.paragraph_format.page_break_before = True
        stile._spaziatura(p, 0, 4)
        stile._bordo_paragrafo(p, ("bottom",), stile.ACCENTO, 12, 4)
        stile._run(p, "Diagramma di Gantt dei lavori" + (f" ({n + 1} di {len(pagine)})" if len(pagine) > 1 else ""),
                   13, stile.INCHIOSTRO, carattere=stile.CARATTERE_TITOLI)
        doc.add_picture(io.BytesIO(immagine), width=Cm(gantt.LARGHEZZA_CM))
    stile.nota(doc, "Le barre sono colorate per impresa; in rosso sono evidenziati i periodi in cui lavorazioni di "
                    "imprese diverse si sovrappongono (capitolo «Interferenze tra le lavorazioni»), in grigio i fine "
                    "settimana. Sotto le fasi sono riportati i lavoratori presenti e il rischio massimo di ciascun "
                    "giorno lavorativo.")
    stile.pagine_verticali(doc)


def _metodo(doc) -> None:
    """Scale di probabilità e danno con la matrice colorata."""
    stile.banda(doc, "Probabilità di accadimento (P)")
    stile.tabella_colorata(doc, ["Valore", "Livello", "Criterio"], [
        ["P1", "Improbabile", "Il danno può verificarsi solo per il concorso di più eventi poco probabili e "
                              "indipendenti; non sono noti episodi analoghi."],
        ["P2", "Poco probabile", "Il danno può verificarsi solo in circostanze particolari; sono noti rari "
                                 "episodi analoghi."],
        ["P3", "Probabile", "Il danno può verificarsi, anche se non in modo automatico; è noto qualche "
                            "episodio analogo."],
        ["P4", "Molto probabile", "Esiste un legame diretto tra la situazione e il danno; sono noti diversi "
                                  "episodi analoghi."]], 8.5, larghezze=[1.6, 3.2, 12.6])
    stile.banda(doc, "Entità del danno (D)")
    stile.tabella_colorata(doc, ["Valore", "Livello", "Criterio"], [
        ["D1", "Lieve", "Infortunio o esposizione con effetti rapidamente reversibili."],
        ["D2", "Medio", "Infortunio o esposizione con effetti reversibili a medio termine."],
        ["D3", "Grave", "Infortunio o esposizione con effetti irreversibili o invalidità parziale."],
        ["D4", "Gravissimo", "Infortunio o esposizione con effetti letali o invalidità totale."]], 8.5,
        larghezze=[1.6, 3.2, 12.6])
    stile.banda(doc, "Matrice del rischio R = P × D")
    _matrice_rischio(doc)


def _matrice_rischio(doc) -> None:
    """Matrice probabilità x danno con i colori dei livelli."""
    righe = [[f"P = {p}"] + [str(p * d) for d in range(1, 5)] for p in range(4, 0, -1)]
    t = stile.tabella_colorata(doc, ["Probabilità \\ Danno", "D = 1", "D = 2", "D = 3", "D = 4"], righe, 8.5,
                               larghezze=[4.2, 3.3, 3.3, 3.3, 3.3])
    for riga, p in zip(t.rows[1:], range(4, 0, -1)):
        for cella, d in zip(riga.cells[1:], range(1, 5)):
            stile.colora_livello(cella, _livello(p * d))
    stile.scheda(doc, [("1 - 2", "rischio basso: misure ordinarie e verifica periodica"),
                       ("3 - 4", "rischio medio: misure da programmare e verificare"),
                       ("6 - 8", "rischio alto: misure prima dell'inizio della fase, verifica del CSE"),
                       ("9 - 16", "rischio molto alto: la fase non inizia senza misure specifiche e verifica "
                                  "del CSE")], 3.0)


def _soggetto(doc, s: dict[str, Any]) -> None:
    stile.banda(doc, s["ruolo"])
    if s.get("nota"):
        stile.nota(doc, s["nota"])
    stile.scheda(doc, [
        ("Nominativo / ragione sociale", s["nome"]), ("Nella persona di", s["rappresentante"]),
        ("Qualifica", s["qualifica"]), ("Indirizzo", s["indirizzo"]),
        ("Telefono", s["telefono"]), ("Email", s["email"]), ("PEC", s["pec"]),
        ("Codice fiscale", s["codice_fiscale"]), ("Partita IVA", s["partita_iva"]),
        ("Incarico conferito il", data_it(s["data_incarico"]) if s.get("data_incarico") else ""),
    ])


def _sezioni_testo(doc, testi: dict[str, str], chiavi: list[str], sottocapitoli: bool = True) -> None:
    """Sezioni della relazione con il loro riferimento di legge; quelle vuote come dati da completare."""
    for chiave in chiavi:
        s = SEZIONI_DETTAGLIO[chiave]
        if sottocapitoli:
            stile.sezione(doc, s.titolo)
            p = doc.add_paragraph()
            stile._spaziatura(p, 0, 4)
            stile._run(p, s.riferimento, 8, stile.TESTO3)
        if (testi.get(chiave) or "").strip():
            stile.testo(doc, testi[chiave])
        else:  # nel riepilogo il titolo, nel documento la guida alla compilazione
            blocco(doc, Manca(s.titolo[:1].lower() + s.titolo[1:], RELAZIONE, s.livello),
                   f"{s.titolo}. {s.guida}")


_immagine_png = stile.immagine_png


def _tavole(doc, allegati: list[dict[str, Any]]) -> None:
    """Tavole allegate (immagini e pagine dei PDF), ognuna con il suo titolo."""
    tavole = [a for a in allegati if a["tipo"] != TIPO_FOTO]
    if not tavole:
        return
    stile.capitolo(doc, "Tavole allegate", "Allegato XV punto 2.1.4", nuova_pagina=True)
    for n, a in enumerate(tavole, 1):
        stile.banda(doc, f"Tavola {n} · {a['tipo']}" + (f" · {a['titolo']}" if a["titolo"] else ""))
        stile.file_allegato(doc, a["file"], "nella pagina «Relazione del PSC», scheda «Tavole, foto e logo»")


def _rischi_per_rischio(doc, rischi: list[dict[str, Any]], nomi_fasi: dict[int, str]) -> None:
    """Ogni rischio con il livello più alto, le fasi in cui compare e le misure, senza ripetizioni."""
    gruppi: dict[str, dict[str, Any]] = {}
    for r in rischi:
        chiave = r["descrizione"].strip().lower()
        if not chiave:
            continue
        g = gruppi.setdefault(chiave, {"nome": r["descrizione"].strip(), "livello": 0, "fasi": [], "misure": []})
        g["livello"] = max(g["livello"], analisi.livello_rischio(r["probabilita"], r["danno"]))
        if r["fase_id"] in nomi_fasi and nomi_fasi[r["fase_id"]] not in g["fasi"]:
            g["fasi"].append(nomi_fasi[r["fase_id"]])
        for riga in r["misure"].split("\n"):
            if riga.strip() and riga.strip() not in g["misure"]:
                g["misure"].append(riga.strip())
    if not gruppi:
        return
    ordinati = sorted(gruppi.values(), key=lambda g: (-g["livello"], g["nome"].lower()))
    t = stile.tabella_colorata(doc, ["Rischio", "Livello max", "Lavorazioni in cui compare", "Misure"], [
        [g["nome"], f"{g['livello']}\n{analisi.classe_rischio(g['livello'])}", "\n".join(g["fasi"]),
         "\n".join(m if len(m) <= 170 else m[:167].rsplit(" ", 1)[0] + "..." for m in g["misure"][:2])
         + ("\n(altre misure nelle schede delle fasi)" if len(g["misure"]) > 2 else "")] for g in ordinati],
        7.5, larghezze=[3.4, 1.9, 5.0, 7.1])
    for riga, g in zip(t.rows[1:], ordinati):
        stile.colora_livello(riga.cells[1], _livello(g["livello"]))


def _attrezzature(doc, repo: Repository, attrezzature: list[dict[str, Any]], rischi: list[dict[str, Any]],
                  nomi_fasi: dict[int, str]) -> None:
    """Schede di macchine, attrezzi e opere provvisionali e tabella della potenza sonora."""
    schede = schede_attrezzature(attrezzature)
    if not schede:
        stile.testo(doc, "", "macchine, attrezzature e opere provvisionali usate nelle lavorazioni", FASI,
                    CONSIGLIATO)
        return
    ordine = list(TIPI_ATTREZZATURA)
    stile.tabella_colorata(doc, ["Tipo", "Denominazione", "Lavorazioni in cui è utilizzata"], [
        [s["tipo"], s["nome"], "; ".join(nomi_fasi[x] for x in s["fasi"] if x in nomi_fasi)]
        for s in sorted(schede.values(), key=lambda s: (ordine.index(s["tipo"]) if s["tipo"] in ordine
                                                        else len(ordine), s["nome"].lower()))],
        larghezze=[3.2, 5.6, 8.6])
    for s in sorted(schede.values(), key=lambda s: s["nome"].lower()):
        stile.riquadro(doc, s["nome"], s["tipo"].lower(), stile.TESTO3, stile.SUPERFICIE, livello=3)
        for riga in s["descrizione"]:
            doc.add_paragraph(riga)
        propri = list(dict.fromkeys(r["descrizione"] for r in rischi if r["fonte"] == s["nome"]))
        if propri:
            p = doc.add_paragraph()
            stile._run(p, "Rischi generati dall'uso: ", grassetto=True)
            p.add_run("; ".join(propri) + ".")
        if s["misure"]:
            p = doc.add_paragraph()
            stile._spaziatura(p, 4, 2)
            stile._run(p, "Misure preventive e protettive", grassetto=True)
            stile.testo(doc, "\n".join(s["misure"]))
        if s["dpi"]:
            dpi_con_icone(doc, s["dpi"], "DPI dell'operatore:")
    rumorose: dict[str, dict[str, Any]] = {}
    for a in attrezzature:
        if a.get("potenza_sonora"):
            v = rumorose.setdefault(a["nome"], {"lwa": 0.0, "fasi": []})
            v["lwa"] = max(v["lwa"], float(a["potenza_sonora"]))
            if nomi_fasi.get(a["fase_id"]) and nomi_fasi[a["fase_id"]] not in v["fasi"]:
                v["fasi"].append(nomi_fasi[a["fase_id"]])
    stile.sezione(doc, "Potenza sonora delle attrezzature e delle macchine", "art. 190")
    if rumorose:
        stile.tabella_colorata(doc, ["Attrezzatura o macchina", "Lavorazioni", "Potenza sonora LWA"], [
            [nome, "; ".join(v["fasi"]), f"{v['lwa']:.1f} dB(A)".replace(".", ",")]
            for nome, v in sorted(rumorose.items(), key=lambda x: -x[1]["lwa"])], larghezze=[5.0, 9.4, 3.0])
        stile.nota(doc, "Valori dichiarati dai fabbricanti. Le imprese valutano l'esposizione dei propri lavoratori "
                        "al rumore nei rispettivi POS (art. 190).")
    else:
        stile.testo(doc, "", "potenza sonora delle attrezzature rumorose (colonna «Potenza sonora» delle "
                             "attrezzature nelle fasi)", FASI, CONSIGLIATO)


def _rischi_particolari(doc, repo: Repository, fasi: list[dict[str, Any]]) -> None:
    from sicurcantiere import relazione

    xi = relazione.leggi_allegato_xi(repo.testo(relazione.CHIAVE_ALLEGATO_XI))
    if not xi["verificato"]:
        segnaposto(doc, doc.add_paragraph(), Manca("verifica dei lavori comportanti rischi particolari "
                                                   "dell'Allegato XI", RISCHI_PARTICOLARI, OBBLIGATORIO))
    righe = []
    for numero, testo_voce in relazione.ALLEGATO_XI:
        voce = xi["voci"][numero]
        misure = voce["misure"] or (Manca(f"misure per i lavori del punto {numero} dell'Allegato XI",
                                          RISCHI_PARTICOLARI, CONSIGLIATO) if voce["presente"] else "")
        righe.append([numero, testo_voce, "Presente" if voce["presente"] else ("Non presente" if xi["verificato"]
                                                                               else "—"), misure])
    t = stile.tabella_colorata(doc, ["Punto", "Lavori", "Nel cantiere", "Misure specifiche"], righe, 7.5,
                               larghezze=[1.2, 8.4, 2.2, 5.6])
    for riga, (numero, _) in zip(t.rows[1:], relazione.ALLEGATO_XI):
        if xi["voci"][numero]["presente"]:
            stile.colora_livello(riga.cells[2], 2)

    stile.sezione(doc, "Rischio da ordigni bellici inesplosi", "art. 91, comma 2-bis")
    ordigni = relazione.leggi_ordigni(repo.testo(relazione.CHIAVE_ORDIGNI))
    con_scavi = relazione.fasi_con_scavi(fasi)
    if ordigni["scavi"] == "No" and not con_scavi:
        doc.add_paragraph("Non sono previsti scavi: il rischio di rinvenimento di ordigni bellici inesplosi non "
                          "sussiste per questo cantiere.")
        return
    stile.scheda(doc, [
        ("Scavi previsti", ordigni["scavi"] or ("Sì (fasi di scavo nel cronoprogramma)" if con_scavi else
                                                Manca("se sono previsti scavi", RISCHI_PARTICOLARI))),
        ("Esito della valutazione", ordigni["esito"] or Manca("esito della valutazione del rischio da ordigni "
                                                              "bellici", RISCHI_PARTICOLARI)),
        ("Motivazione e fonti", ordigni["motivazione"] or Manca("motivazione e fonti della valutazione del rischio "
                                                                "da ordigni bellici", RISCHI_PARTICOLARI,
                                                                CONSIGLIATO)),
    ], 4.2)
    stile.nota(doc, "La valutazione del rischio da ordigni bellici inesplosi è eseguita dal coordinatore per la "
                    "progettazione. Se si procede alla bonifica preventiva, il committente incarica un'impresa "
                    "specializzata iscritta all'albo istituito presso il Ministero della difesa (artt. 91, comma "
                    "2-bis, e 104, comma 4-bis).")


def _costi(doc, c: dict[str, Any], costi: list[dict[str, Any]]) -> None:
    from sicurcantiere.relazione import CATEGORIE_COSTI

    if not costi:
        stile.testo(doc, "", "stima analitica dei costi della sicurezza (Allegato XV punto 4)", COSTI, OBBLIGATORIO)
        return
    stile.scheda(doc, [("Prezzario di riferimento", c.get("prezzario") or
                        Manca("prezzario di riferimento dei costi della sicurezza", COSTI, CONSIGLIATO))])
    nomi = dict(CATEGORIE_COSTI)
    righe, sottototali = [], []
    for chiave in [k for k, _ in CATEGORIE_COSTI] + [""]:
        voci = [x for x in costi if (x.get("categoria") or "") == chiave]
        if not voci:
            continue
        righe.append([" ", f"{chiave}) {nomi[chiave]}" if chiave else "Altre voci", " ", " ", " ", " "])
        sottototali.append(len(righe))
        for x in voci:
            righe.append([x["codice"], x["descrizione"], x["unita"], f"{x['quantita']:g}".replace(".", ","),
                          euro(x["prezzo_unitario"]), euro(x["quantita"] * x["prezzo_unitario"])])
        righe.append([" ", "Totale della categoria", " ", " ", " ", euro(analisi.totale_costi(voci))])
        sottototali.append(-len(righe))
    righe.append([" ", "Totale dei costi della sicurezza (non soggetti a ribasso)", " ", " ", " ",
                  euro(analisi.totale_costi(costi))])
    t = stile.tabella_colorata(doc, ["Codice", "Descrizione", "U.M.", "Quantità", "Prezzo unit.", "Importo"],
                               righe, 8.5, larghezze=[1.8, 7.4, 1.4, 1.8, 2.4, 2.6])
    for indice in sottototali:
        riga = t.rows[abs(indice)]
        for cella in riga.cells:
            if indice > 0:
                stile.ombreggia_cella(cella, stile.SUPERFICIE)
            for run in cella.paragraphs[0].runs:
                run.bold = True
    ultima = t.rows[-1]
    for cella in ultima.cells:
        stile.ombreggia_cella(cella, stile.TINTA)
        for run in cella.paragraphs[0].runs:
            run.bold = True
            run.font.size = Pt(9)
    senza = [x for x in costi if not x.get("categoria")]
    if senza:
        segnaposto(doc, doc.add_paragraph(), Manca(f"categoria del punto 4.1.1 per {len(senza)} voci di costo",
                                                   COSTI, CONSIGLIATO))
    stile.nota(doc, "Stima analitica per voci singole, a corpo o a misura (Allegato XV punto 4.1.3). I costi della "
                    "sicurezza sono compresi nell'importo totale dei lavori e non sono soggetti a ribasso nelle "
                    "offerte delle imprese esecutrici (punto 4.1.4).")


def genera_psc(repo: Repository, destinazione: str | Path) -> Path:
    with stile.colori_documento(*stile.COLORI_PSC):
        return _genera(repo, destinazione)


def _genera(repo: Repository, destinazione: str | Path) -> Path:
    from sicurcantiere.documenti import organigramma

    c = repo.cantiere()
    soggetti = anagrafiche.soggetti(repo)
    per_ruolo = {s["ruolo"]: s for s in soggetti if s["nome"].strip()}
    imprese = repo.elenco("imprese")
    fasi = repo.elenco("fasi")
    costi = repo.elenco("costi")
    rischi = repo.rischi()
    attrezzature = [a for f in fasi for a in repo.elenco("attrezzature", fase_id=f["id"])]
    allegati = repo.elenco("allegati_psc")
    nomi_imprese = {i["id"]: i["ragione_sociale"] for i in imprese}
    nomi_fasi = {f["id"]: nome_fase(f) for f in fasi}
    livelli = analisi.livelli_fasi(rischi)
    elenco_interferenze = analisi.interferenze(fasi, livelli)
    prescrizioni = repo.prescrizioni()
    testi = {chiave: repo.testo(chiave) for chiave, _ in SEZIONI}
    periodo = analisi.periodo_lavori(fasi)
    revisioni = repo.elenco("revisioni", documento="PSC")
    ultima = revisioni[-1] if revisioni else None
    data_documento = data_it(ultima["data"]) if ultima and ultima["data"] else data_it(date.today())
    revisione = f"Rev. {ultima['numero']} del {data_documento}" if ultima else f"Emissione del {data_documento}"
    comune = f"{c['comune']} ({c['provincia']})" if c["provincia"] else c["comune"]
    committente = per_ruolo.get(anagrafiche.COMMITTENTE)
    csp = per_ruolo.get(anagrafiche.CSP)
    cse = per_ruolo.get(anagrafiche.CSE)
    responsabile = per_ruolo.get(anagrafiche.RESPONSABILE_LAVORI)
    logo = repo.testo("logo_studio")

    doc = Document()
    stile.imposta_stili(doc)
    stile.imposta_pagina(doc, " · ".join(v for v in (c["nome"], comune) if v),
                         (f"Piano di sicurezza e coordinamento · {c['nome']}".strip(" ·"), revisione))

    # --- copertina ---
    firme_copertina = [("Il coordinatore per la progettazione", csp["nome"] if csp else "")]
    if responsabile and responsabile["nome"] != (committente or {}).get("nome"):
        firme_copertina.append(("Per presa visione, il responsabile dei lavori", responsabile["nome"]))
    else:
        firme_copertina.append(("Per presa visione, il committente", committente["nome"] if committente else ""))
    stile.copertina(
        doc,
        comune=comune,
        titolo="Piano di sicurezza e coordinamento",
        riferimento="D.Lgs. 9 aprile 2008, n. 81 · art. 100 e Allegato XV punto 2",
        voci=[("Cantiere", _o(indirizzo_cantiere(c), "indirizzo del cantiere", CANTIERE)),
              ("Natura dell'opera", _o(c["natura_opera"], "natura dell'opera", CANTIERE, CONSIGLIATO)),
              ("Committente", _o(", ".join(s["nome"] for s in soggetti if s["ruolo"] == anagrafiche.COMMITTENTE
                                           and s["nome"].strip()), "committente", SOGGETTI)),
              ("Importo dei lavori", euro(c["importo_lavori"]) if c["importo_lavori"] else "")],
        data_luogo=f"{c['comune'] or ''}, {data_documento}".strip(", "),
        firme=firme_copertina,
        impresa=[c["nome"] or "Opera", c["descrizione"].strip().split("\n")[0][:160] if c["descrizione"] else ""],
        logo=logo,
        revisione=revisione,
    )
    if csp:
        p = doc.add_paragraph()
        stile._spaziatura(p, 18, 0)
        stile._bordo_paragrafo(p, ("top",), stile.LINEA, 4, 6)
        stile._run(p, "Coordinatore per la progettazione", 7.5, stile.TESTO3)
        p = doc.add_paragraph()
        stile._spaziatura(p, 1, 0)
        stile._run(p, " ".join(v for v in (csp["qualifica"], csp["nome"]) if v), 9.5, stile.ACCENTO_SCURO,
                   grassetto=True)
        for riga in (csp["indirizzo"], " · ".join(v for v in (f"tel. {csp['telefono']}" if csp["telefono"] else "",
                                                               csp["email"], csp["pec"]) if v)):
            if riga:
                p = doc.add_paragraph()
                stile._spaziatura(p)
                stile._run(p, riga, 8, stile.TESTO2)

    # --- indice e revisioni ---
    stile.indice(doc)
    stile.banda(doc, "Revisioni del documento")
    if revisioni:
        stile.tabella_colorata(doc, ["Rev.", "Data", "Descrizione"], [
            [str(r["numero"]), data_it(r["data"]), r["descrizione"]] for r in revisioni], larghezze=[1.2, 2.6, 13.6])
    else:
        doc.add_paragraph(f"Prima emissione del {data_documento}.")

    # ===== Parte 1 =====
    stile.parte(doc, "1", "L'opera e i soggetti",
                "Identificazione e descrizione dell'opera, soggetti con compiti di sicurezza, imprese e lavoratori "
                "autonomi, organigramma (Allegato XV punto 2.1.2, lettere a) e b)).")

    stile.capitolo(doc, "Dati generali del lavoro", "Allegato XV punto 2.1.2, lettera a), punto 1")
    stile.scheda(doc, [
        ("Opera", _o(c["nome"], "denominazione dell'opera", CANTIERE)),
        ("Natura dell'opera", _o(c["natura_opera"], "natura dell'opera", CANTIERE, CONSIGLIATO)),
        ("Indirizzo del cantiere", _o(indirizzo_cantiere(c), "indirizzo del cantiere", CANTIERE)),
        ("CAP", c["cap"]),
        ("Titolo abilitativo", _o(c["titolo_abilitativo"], "titolo abilitativo", CANTIERE, FACOLTATIVO)),
        ("Importo presunto dei lavori", euro(c["importo_lavori"]) if c["importo_lavori"] else
         Manca("importo presunto dei lavori", CANTIERE, CONSIGLIATO)),
        ("Numero di imprese e lavoratori autonomi previsti", str(len(imprese)) if imprese else ""),
        ("Numero massimo presunto di lavoratori", str(analisi.picco_lavoratori(fasi)) if fasi else ""),
        ("Entità presunta del cantiere", f"{analisi.uomini_giorno(fasi)} uomini-giorno" if fasi else ""),
        ("Inizio dei lavori (presunto)", data_it(periodo[0]) if periodo else ""),
        ("Fine dei lavori (presunta)", data_it(periodo[1]) if periodo else ""),
        ("Durata presunta", f"{analisi.giorni_lavorativi(*periodo)} giorni lavorativi, "
                            f"{(periodo[1] - periodo[0]).days + 1} giorni di calendario" if periodo else ""),
    ])

    stile.capitolo(doc, "Contesto in cui è collocata l'area di cantiere", SEZIONI_DETTAGLIO["contesto"].riferimento)
    _sezioni_testo(doc, testi, ["contesto"], sottocapitoli=False)
    foto = [a["file"] for a in allegati if a["tipo"] == TIPO_FOTO] + foto_cantiere(repo)
    for n, percorso in enumerate(dict.fromkeys(foto), 1):
        stile.immagine(doc, percorso, 13, f"Figura {n} - area del cantiere")

    stile.capitolo(doc, "Descrizione sintetica dell'opera", "Allegato XV punto 2.1.2, lettera a), punto 3")
    stile.testo(doc, c["descrizione"], "descrizione sintetica dell'opera, con le scelte progettuali, "
                                       "architettoniche, strutturali e tecnologiche", CANTIERE, OBBLIGATORIO)

    stile.capitolo(doc, "Soggetti con compiti di sicurezza", "Allegato XV punto 2.1.2, lettera b), e artt. 89-93")
    for s in soggetti:
        if s["nome"].strip():
            _soggetto(doc, s)
    for ruolo, dato, livello in ((anagrafiche.COMMITTENTE, "committente", OBBLIGATORIO),
                                 (anagrafiche.CSP, "coordinatore per la progettazione (CSP)", OBBLIGATORIO),
                                 (anagrafiche.CSE, "coordinatore per l'esecuzione (CSE), se già designato",
                                  CONSIGLIATO)):
        if ruolo not in per_ruolo:
            p = doc.add_paragraph()
            stile._run(p, f"{ruolo}: ", grassetto=True)
            segnaposto(doc, p, Manca(dato, SOGGETTI, livello))

    stile.capitolo(doc, "Imprese esecutrici e lavoratori autonomi", "Allegato XV punto 2.1.2, lettera b)")
    if imprese:
        doc.add_paragraph("Sono di seguito indicati le imprese e i lavoratori autonomi previsti, con i rispettivi "
                          "datori di lavoro. L'elenco è completato prima dell'inizio dei lavori, quando le imprese "
                          "sono individuate.")
        affidatarie = {i["id"]: i["ragione_sociale"] for i in imprese}
        stile.tabella_colorata(doc, ["Ragione sociale", "Ruolo", "Datore di lavoro", "Patente a crediti",
                                     "Lavori affidati"], [
            [i["ragione_sociale"] + (f"\nper {affidatarie[i['affidataria_id']]}" if i.get("affidataria_id") in
                                     affidatarie else "") + (f"\nP. IVA {i['partita_iva']}" if i["partita_iva"] else ""),
             i["tipo"],
             i["datore_lavoro"] or Manca("datore di lavoro di " + (i["ragione_sociale"] or "un'impresa"), IMPRESE,
                                         CONSIGLIATO),
             anagrafiche.descrizione_patente(i) or Manca("patente a crediti di " + (i["ragione_sociale"] or
                                                                                   "un'impresa"), IMPRESE, CONSIGLIATO),
             " · ".join(v for v in (i["tipologia_lavori"], euro(i["importo_lavori"])
                                    if i.get("importo_lavori") else "") if v)]
            for i in imprese], 8, larghezze=[4.2, 2.4, 3.2, 3.4, 4.2])
        stile.nota(doc, "Prima dell'inizio dei lavori il committente o il responsabile dei lavori verifica l'idoneità "
                        "tecnico-professionale delle imprese e dei lavoratori autonomi (art. 90, comma 9, e Allegato "
                        "XVII), compresa la patente a crediti (art. 27): le imprese e i lavoratori autonomi privi di "
                        "patente, o con un punteggio inferiore a 15 crediti, non possono operare in cantiere.")
    else:
        stile.testo(doc, "", "imprese esecutrici e lavoratori autonomi, da indicare prima dell'inizio dei lavori",
                    IMPRESE, CONSIGLIATO)

    stile.capitolo(doc, "Organigramma del cantiere")
    try:
        schema = organigramma.immagine(soggetti, imprese, "#" + stile.ACCENTO_SCURO, "#" + stile.TINTA)
    except Exception:  # noqa: BLE001 - senza Qt l'organigramma si omette
        schema = None
    if schema:
        _immagine_png(doc, schema, stile.LARGHEZZA_UTILE)
    else:
        doc.add_paragraph("Organigramma non disponibile: soggetti e imprese da indicare.")

    stile.capitolo(doc, "Documentazione da custodire in cantiere", "artt. 90, 96, 99 e 101 e Allegato XVII")
    doc.add_paragraph("Le imprese tengono in cantiere, a disposizione degli organi di vigilanza e del coordinatore "
                      "per l'esecuzione, la seguente documentazione:")
    stile.elenco(doc, DOCUMENTI_IN_CANTIERE)

    # ===== Parte 2 =====
    stile.parte(doc, "2", "Area, organizzazione e lavorazioni",
                "Valutazione dei rischi e scelte progettuali e organizzative per l'area di cantiere, "
                "l'organizzazione del cantiere e le lavorazioni (Allegato XV punto 2.1.2, lettere c) e d), e "
                "punto 2.2).")

    stile.capitolo(doc, "Metodo di valutazione dei rischi", "Allegato XV punto 2.1.2, lettera c)")
    doc.add_paragraph("I rischi concreti sono individuati con riferimento all'area, all'organizzazione del cantiere, "
                      "alle lavorazioni e alle loro interferenze. Ogni rischio è valutato come prodotto R = P × D "
                      "della probabilità di accadimento P e dell'entità del danno D, secondo le scale seguenti.")
    _metodo(doc)

    stile.capitolo(doc, "Area di cantiere", "Allegato XV punto 2.1.2, lettera d), punto 1, e punto 2.2.1")
    _sezioni_testo(doc, testi, ["area_caratteristiche", "area_esterni", "area_circostante"])

    stile.capitolo(doc, "Caratteristiche idrogeologiche e profilo del terreno",
                   SEZIONI_DETTAGLIO["idrogeologia"].riferimento)
    _sezioni_testo(doc, testi, ["idrogeologia"], sottocapitoli=False)

    stile.capitolo(doc, "Organizzazione del cantiere", "Allegato XV punto 2.1.2, lettera d), punto 2, e punto 2.2.2")
    if testi.get("organizzazione", "").strip():
        stile.testo(doc, testi["organizzazione"])
    _sezioni_testo(doc, testi, [k for k in SEZIONI_DETTAGLIO if k.startswith("org_")])

    stile.capitolo(doc, "Segnaletica di cantiere", "Titolo V e Allegati XXIV-XXXII del D.Lgs. 81/2008")
    stile.cartello_testo(doc, "Vietato l'accesso ai non addetti ai lavori")
    cartelli = segnaletica.segnali_cantiere([r["descrizione"] for r in rischi],
                                            [r["dpi"] for r in rischi] + [a["dpi"] for a in attrezzature])
    stile.griglia_immagini(doc, [(s.file, s.testo) for s in cartelli], colonne=5, lato_cm=2.3)
    stile.nota(doc, "Cartelli conformi alla UNI EN ISO 7010, collocati agli accessi e in prossimità dei pericoli; "
                    "il cartello di cantiere riporta i dati della notifica preliminare.")

    stile.capitolo(doc, SEZIONI_DETTAGLIO["scelte"].titolo, SEZIONI_DETTAGLIO["scelte"].riferimento)
    _sezioni_testo(doc, testi, ["scelte"], sottocapitoli=False)

    stile.capitolo(doc, "Lavorazioni: rischi e misure per fase", "Allegato XV punto 2.1.2, lettera d), punto 3, e "
                                                                "punto 2.2.3")
    if not fasi:
        stile.testo(doc, "", "fasi di lavoro con i loro rischi", FASI, OBBLIGATORIO)
    lavorazione_corrente = object()
    for f in fasi:
        if f["lavorazione"] != lavorazione_corrente:
            lavorazione_corrente = f["lavorazione"]
            gruppo = [g for g in fasi if g["lavorazione"] == lavorazione_corrente]
            stile.riquadro(doc, lavorazione_corrente or "Altre fasi di lavoro", "lavorazione")
            if lavorazione_corrente:
                p = doc.add_paragraph()
                stile._run(p, "La lavorazione è suddivisa nelle seguenti fasi: ", grassetto=True)
                p.add_run("; ".join(nome_fase(g) for g in gruppo) + ".")
        h = doc.add_heading(level=3)
        stile._run(h, nome_fase(f))
        attrezzature_fase = [a for a in attrezzature if a["fase_id"] == f["id"]]
        periodo_fase = (f"dal {data_it(f['data_inizio'])} al {data_it(f['data_fine'])}"
                        if f["data_inizio"] and f["data_fine"] else
                        Manca(f"date di inizio e fine della fase {nome_fase(f)}", FASI, OBBLIGATORIO))
        stile.scheda(doc, [
            ("Impresa", nomi_imprese.get(f["impresa_id"]) or
             Manca(f"impresa esecutrice della fase {nome_fase(f)}", FASI, CONSIGLIATO)),
            ("Periodo", periodo_fase),
            ("Zona del cantiere", f["zona"]),
            ("Addetto", f["addetto"]),
            ("Lavoratori previsti", str(f["lavoratori"]) if f.get("lavoratori") else ""),
            ("Macchine, attrezzi e opere provvisionali", ", ".join(a["nome"] for a in attrezzature_fase)),
            ("Procedure richieste ai POS", f["procedure_psc"]),
        ], 5.0)
        if f["dettaglio"]:
            stile.testo(doc, f["dettaglio"])
        rischi_fase = [r for r in rischi if r["fase_id"] == f["id"]]
        if rischi_fase:
            tabella_rischi(doc, rischi_fase)
            dpi_con_icone(doc, [r["dpi"] for r in rischi_fase], "DPI previsti:")
        else:
            stile.testo(doc, "", f"rischi della fase {nome_fase(f)} (nessun rischio valutato)", RISCHI, OBBLIGATORIO)

    stile.capitolo(doc, "Rischi individuati e lavorazioni in cui compaiono", "Allegato XV punti 2.2.3 e 2.2.4")
    doc.add_paragraph("Si riporta il riepilogo dei rischi delle lavorazioni, in ordine decrescente di livello, con le "
                      "fasi in cui compaiono e le misure previste.")
    _rischi_per_rischio(doc, rischi, nomi_fasi)

    stile.capitolo(doc, "Macchine, attrezzature e opere provvisionali", "artt. 70-73 e Allegati V-VII")
    _attrezzature(doc, repo, attrezzature, rischi, nomi_fasi)

    stile.capitolo(doc, "Lavori comportanti rischi particolari", "art. 100, comma 1, e Allegato XI")
    _rischi_particolari(doc, repo, fasi)

    # ===== Parte 3 =====
    stile.parte(doc, "3", "Coordinamento",
                "Interferenze tra le lavorazioni, procedure da sviluppare nei POS, uso comune di apprestamenti e "
                "servizi, cooperazione tra le imprese ed emergenze (Allegato XV punto 2.1.2, lettere e), f), g) e h), "
                "e punti 2.1.3 e 2.3).")

    stile.capitolo(doc, "Interferenze tra le lavorazioni e prescrizioni operative",
                   "Allegato XV punto 2.1.2, lettera e), e punti 2.3.1-2.3.3")
    if elenco_interferenze:
        doc.add_paragraph("Dal cronoprogramma risultano le seguenti lavorazioni contemporanee di imprese diverse, "
                          "con la relativa gravità (livello di rischio più alto tra le due fasi) e le prescrizioni "
                          "per eliminarle o ridurle.")
        t = stile.tabella_colorata(doc, ["Fase", "Fase interferente", "Periodo", "Gravità", "Prescrizioni operative"], [
            [nome_fase(i.fase_a), nome_fase(i.fase_b), f"{data_it(i.inizio)}\n{data_it(i.fine)}",
             f"{i.gravita}\n{analisi.classe_rischio(i.gravita)}" if i.gravita else "-",
             prescrizioni.get(i.chiave) or PRESCRIZIONE_GENERICA] for i in elenco_interferenze],
            8, larghezze=[3.6, 3.6, 2.0, 1.9, 6.3])
        for riga, i in zip(t.rows[1:], elenco_interferenze):
            if i.gravita:
                stile.colora_livello(riga.cells[3], _livello(i.gravita))
        for i in elenco_interferenze:
            if i.gravita >= 6 and not prescrizioni.get(i.chiave, "").strip():
                segnaposto(doc, doc.add_paragraph(), Manca(
                    f"prescrizione specifica per l'interferenza grave tra {nome_fase(i.fase_a)} e "
                    f"{nome_fase(i.fase_b)}", INTERFERENZE, CONSIGLIATO))
    else:
        doc.add_paragraph("Dal cronoprogramma non risultano lavorazioni contemporanee di imprese diverse. Le "
                          "eventuali variazioni del programma sono valutate dal coordinatore per l'esecuzione.")
    _sezioni_testo(doc, testi, ["verifica_prescrizioni"])

    stile.capitolo(doc, "Procedure complementari e di dettaglio da sviluppare nei POS", "Allegato XV punto 2.1.3")
    con_procedure = [f for f in fasi if f["procedure_psc"].strip()]
    if con_procedure:
        doc.add_paragraph("Per le fasi seguenti le imprese esecutrici esplicitano nel proprio POS le procedure "
                          "complementari e di dettaglio, connesse alle loro scelte autonome:")
        stile.tabella_colorata(doc, ["Fase", "Impresa", "Procedura richiesta"], [
            [nome_fase(f), nomi_imprese.get(f["impresa_id"], "") or "—", f["procedure_psc"]] for f in con_procedure],
            larghezze=[5.4, 3.6, 8.4])
    else:
        doc.add_paragraph("Non sono richieste procedure complementari: le imprese esecutrici seguono le misure del "
                          "presente piano e le integrano nel proprio POS.")

    stile.capitolo(doc, SEZIONI_DETTAGLIO["uso_comune"].titolo, SEZIONI_DETTAGLIO["uso_comune"].riferimento)
    _sezioni_testo(doc, testi, ["uso_comune"], sottocapitoli=False)

    stile.capitolo(doc, SEZIONI_DETTAGLIO["cooperazione"].titolo, SEZIONI_DETTAGLIO["cooperazione"].riferimento)
    _sezioni_testo(doc, testi, ["cooperazione"], sottocapitoli=False)

    stile.capitolo(doc, "Primo soccorso, antincendio ed evacuazione", "Allegato XV punto 2.1.2, lettera h)")
    _sezioni_testo(doc, testi, ["emergenze"], sottocapitoli=False)
    stile.banda(doc, "Numeri di emergenza")
    stile.scheda(doc, [(nome, numero) for nome, numero in NUMERI_EMERGENZA])
    stile.banda(doc, "Strutture del territorio")
    _sezioni_testo(doc, testi, ["telefoni"], sottocapitoli=False)

    # ===== Parte 4 =====
    stile.parte(doc, "4", "Tempi e costi",
                "Durata delle lavorazioni e cronoprogramma, stima dei costi della sicurezza (Allegato XV punto "
                "2.1.2, lettere i) e l), e punto 4).")

    stile.capitolo(doc, "Cronoprogramma dei lavori", "Allegato XV punto 2.1.2, lettera i)")
    if fasi:
        stile.scheda(doc, [
            ("Inizio e fine dei lavori", f"dal {data_it(periodo[0])} al {data_it(periodo[1])}" if periodo else ""),
            ("Durata presunta", f"{analisi.giorni_lavorativi(*periodo)} giorni lavorativi" if periodo else ""),
            ("Entità presunta del cantiere", f"{analisi.uomini_giorno(fasi)} uomini-giorno"),
            ("Numero massimo presunto di lavoratori", str(analisi.picco_lavoratori(fasi))),
        ])
        _programma(doc, fasi, nomi_imprese)
    _diagramma_gantt(doc, fasi, nomi_imprese, livelli, elenco_interferenze)

    stile.capitolo(doc, "Stima dei costi della sicurezza", "Allegato XV punto 2.1.2, lettera l), e punto 4")
    _costi(doc, c, costi)

    # --- allegati e firme ---
    stile.capitolo(doc, "Allegati e firme")
    tavole = [a for a in allegati if a["tipo"] != TIPO_FOTO]
    voci_allegati = [f"{a['tipo']}" + (f": {a['titolo']}" if a["titolo"] else "") + ";" for a in tavole]
    voci_allegati.append("fascicolo con le caratteristiche dell'opera, ai sensi dell'art. 91, comma 1, lettera b).")
    doc.add_paragraph("Fanno parte integrante del presente piano:")
    stile.elenco_lettere(doc, voci_allegati)
    if not any(a["tipo"] == "Planimetria di cantiere" for a in tavole):
        segnaposto(doc, doc.add_paragraph(), Manca("planimetria del cantiere (Allegato XV punto 2.1.4)", TAVOLE,
                                                   OBBLIGATORIO))
    doc.add_paragraph("Il coordinatore per la progettazione dichiara di aver redatto il presente piano di sicurezza "
                      "e coordinamento ai sensi dell'art. 100 e dell'Allegato XV del D.Lgs. 81/2008. Il committente "
                      "o il responsabile dei lavori lo trasmette a tutte le imprese invitate a presentare offerte; "
                      "le imprese esecutrici ne prendono visione e lo mettono a disposizione dei rappresentanti dei "
                      "lavoratori per la sicurezza almeno dieci giorni prima dell'inizio dei lavori (art. 101, "
                      "comma 1, e art. 100, comma 4).")
    p = doc.add_paragraph(f"{c['comune'] or ''}, {data_documento}".strip(", "))
    stile._spaziatura(p, 8, 8)
    firme = [("Il coordinatore per la progettazione", csp["nome"] if csp else ""),
             ("Per presa visione, il committente" if not responsabile else
              "Per presa visione, il responsabile dei lavori", (responsabile or committente or {}).get("nome", ""))]
    if cse:
        firme.append(("Per presa visione, il coordinatore per l'esecuzione", cse["nome"]))
    for i in imprese:
        if i["ragione_sociale"].strip() and not anagrafiche.autonomo(i):
            firme.append((f"Per presa visione e accettazione, il datore di lavoro di {i['ragione_sociale']}",
                          i["datore_lavoro"]))
    stile.firme(doc, firme)
    _tavole(doc, allegati)

    return salva(doc, destinazione)
