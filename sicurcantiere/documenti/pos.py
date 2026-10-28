"""Generazione del Piano Operativo di Sicurezza (Allegato XV, punto 3.2.1)."""

from __future__ import annotations

import re
from datetime import date
from pathlib import Path

from docx import Document
from docx.shared import Cm

from sicurcantiere import anagrafiche, analisi, segnaletica
from sicurcantiere.documenti import stile
from sicurcantiere.documenti.composizione import CONSIGLIATO, FACOLTATIVO, OBBLIGATORIO, Manca, segnaposto
from sicurcantiere.documenti.comune import data_it, euro, indirizzo_cantiere, nome_fase, salva
from sicurcantiere.documenti.sostanze import aggiungi_sostanze, sostanze_impresa
from sicurcantiere.documenti.valutazione import aggiungi_valutazione, lavorazioni_valutazione, valutazioni_impresa
from sicurcantiere.repository import Repository
from sicurcantiere.struttura import MACCHINA, struttura_fase
from sicurcantiere.valutazioni import METODI
from sicurcantiere.valutazioni.automatiche import da_eseguire

# Sezioni a testo libero del POS, salvate per impresa.
SEZIONI_POS: list[tuple[str, str]] = [
    ("attivita", "Attività svolte in cantiere, modalità organizzative e turni di lavoro"),
    ("sostanze", "Sostanze pericolose: note integrative all'elenco"),
    ("rumore", "Esito della valutazione del rumore"),
    ("integrative", "Misure preventive e protettive integrative rispetto al PSC"),
    ("procedure", "Procedure complementari e di dettaglio richieste dal PSC"),
    ("telefoni", "Telefoni e indirizzi utili locali (ospedale, Vigili del fuoco, Carabinieri, ASL...)"),
    ("emergenze", "Gestione delle emergenze: primo soccorso, antincendio ed evacuazione"),
]


def chiave_testo(impresa_id: int, sezione: str) -> str:
    return f"pos:{impresa_id}:{sezione}"


def elenco_dpi(testi: list[str]) -> list[str]:
    """Unisce gli elenchi di DPI delle schede di rischio eliminando i doppioni."""
    visti: dict[str, str] = {}
    for testo in testi:
        for voce in re.split(r",|;|\n", testo):
            voce = voce.strip().rstrip(".")
            if voce:
                visti.setdefault(voce.lower(), voce[0].upper() + voce[1:])
    return sorted(visti.values(), key=str.lower)


def _data_o_spazio(valore) -> str:
    return data_it(valore) if valore else "____/____/________"


def periodo_impresa(repo: Repository, impresa: dict) -> tuple:
    """(inizio, fine) dei lavori dell'impresa: date indicate o, in mancanza, quelle delle sue lavorazioni."""
    periodo_fasi = analisi.periodo_lavori(repo.elenco("fasi", impresa_id=impresa["id"]))
    inizio = analisi.data(impresa["data_inizio"]) or (periodo_fasi[0] if periodo_fasi else None)
    fine = analisi.data(impresa["data_fine"]) or (periodo_fasi[1] if periodo_fasi else None)
    return inizio, fine


def _periodo(impresa: dict, periodo_fasi: tuple | None) -> str:
    """Date dei lavori indicate per l'impresa o, in mancanza, ricavate dalle sue lavorazioni."""
    inizio = analisi.data(impresa["data_inizio"]) or (periodo_fasi[0] if periodo_fasi else None)
    fine = analisi.data(impresa["data_fine"]) or (periodo_fasi[1] if periodo_fasi else None)
    if not (inizio and fine):
        return ""
    return f"dal {data_it(inizio)} al {data_it(fine)} ({(fine - inizio).days + 1} giorni)"


# Numeri di emergenza nazionali (il numero unico 112 è attivo in tutta Italia).
NUMERI_EMERGENZA = [
    ("Numero unico di emergenza", "112"),
    ("Carabinieri", "112"),
    ("Polizia di Stato", "113"),
    ("Vigili del fuoco", "115"),
    ("Emergenza sanitaria", "118"),
]

# Documenti da tenere in cantiere (testo redatto per SicurCantiere dagli obblighi del D.Lgs. 81/2008
# e delle norme collegate; le voci «ove previsto» dipendono dal cantiere).
DOCUMENTI_IN_CANTIERE = [
    "copia della notifica preliminare, esposta in modo visibile (art. 99)",
    "piano di sicurezza e di coordinamento e fascicolo dell'opera, se previsti (artt. 100 e 91)",
    "piano operativo di sicurezza di ciascuna impresa esecutrice e relativi aggiornamenti (art. 96)",
    "titolo abilitativo dei lavori",
    "iscrizione alla Camera di commercio, DURC e documentazione dell'idoneità tecnico-professionale "
    "delle imprese e dei lavoratori autonomi (art. 90 e Allegato XVII)",
    "libro unico del lavoro o documentazione equivalente e tessere di riconoscimento dei lavoratori (art. 26, comma 8)",
    "attestati di formazione, addestramento e abilitazione all'uso delle attrezzature (artt. 37 e 73)",
    "giudizi di idoneità alla mansione rilasciati dal medico competente (art. 41)",
    "verbali di consegna dei DPI (art. 77)",
    "schede di dati di sicurezza delle sostanze e miscele pericolose utilizzate (Regolamento (CE) n. 1907/2006)",
    "libretti d'uso e manutenzione e dichiarazioni di conformità CE di macchine e attrezzature (art. 70)",
    "verbali delle verifiche periodiche delle attrezzature di sollevamento e registro delle verifiche "
    "trimestrali di funi e catene (art. 71 e Allegato VII), ove previsto",
    "PiMUS, autorizzazione ministeriale e, se richiesto, progetto del ponteggio (artt. 133 e 136), ove previsto",
    "dichiarazione di conformità dell'impianto elettrico e di terra del cantiere (D.M. 37/2008) e relativa "
    "denuncia (D.P.R. 462/2001), ove prevista",
    "piano di lavoro per la rimozione dell'amianto (art. 256), ove previsto",
    "verbali dei sopralluoghi e delle riunioni di coordinamento con il coordinatore per l'esecuzione",
]


def foto_cantiere(repo: Repository) -> list[str]:
    return [r for r in repo.testo("foto_cantiere").splitlines() if r.strip()]


def descrizione_rischio(r: dict) -> str:
    return f"{r['descrizione']} ({r['fonte']})" if r["fonte"] else r["descrizione"]


def tabella_rischi(doc, rischi: list[dict]) -> None:
    """Rischi della fase: probabilità, danno e livello colorato, misure e DPI del rischio."""
    righe = []
    for r in rischi:
        livello = analisi.livello_rischio(r["probabilita"], r["danno"])
        misure = r["misure"].strip()
        if r["dpi"].strip():
            misure = f"{misure}\nDPI: {r['dpi'].strip()}".strip()
        righe.append([descrizione_rischio(r),
                      f"P{r['probabilita']} × D{r['danno']} = {livello}\n{analisi.classe_rischio(livello)}", misure])
    if not righe:
        return
    t = stile.tabella_colorata(doc, ["Rischio", "Valutazione", "Misure preventive e protettive"], righe, 8.5,
                               larghezze=[5.0, 2.6, 9.8])
    for riga, r in zip(t.rows[1:], rischi):
        livello = analisi.livello_rischio(r["probabilita"], r["danno"])
        stile.colora_livello(riga.cells[1], 3 if livello >= 9 else 2 if livello >= 6 else 1 if livello >= 3 else 0)


def _esito_breve(repo: Repository, valutazione: dict) -> list[str]:
    """Indice e classe della valutazione, o «da eseguire» se i parametri mancano."""
    if da_eseguire(repo, valutazione["id"]):
        return ["-", "Valutazione da eseguire"]
    e = METODI[valutazione["metodo"]].esegui(repo.dati_valutazione(valutazione["id"]))
    return [e.indice_testo, e.classe]


def dpi_con_icone(doc, testi: list[str], etichetta: str) -> None:
    """DPI in chiaro e, dove c'è, il cartello di obbligo corrispondente."""
    voci = elenco_dpi(testi)
    if not voci:
        return
    p = doc.add_paragraph()
    stile._run(p, f"{etichetta} ", 8.5, stile.TESTO3)
    p.add_run("; ".join(voci) + ".")
    icone = segnaletica.icone_dpi(voci)
    if icone:
        stile.icone_in_riga(doc.add_paragraph(), [segnaletica.SEGNALI[i].file for i in icone], 0.8)


def schede_attrezzature(attrezzature: list[dict]) -> dict[str, dict]:
    """Una scheda per attrezzatura: se la stessa compare in più fasi i testi si uniscono senza doppioni."""
    schede: dict[str, dict] = {}
    for a in attrezzature:
        nome = a["nome"].strip()
        if not nome:
            continue
        s = schede.setdefault(nome, {"nome": nome, "tipo": a["tipo"], "descrizione": [], "misure": [], "dpi": [],
                                     "fasi": []})
        for campo in ("descrizione", "misure"):
            for riga in a[campo].split("\n"):
                if riga.strip() and riga.strip() not in s[campo]:
                    s[campo].append(riga.strip())
        if a["dpi"].strip():
            s["dpi"].append(a["dpi"])
        if a["fase_id"] not in s["fasi"]:
            s["fasi"].append(a["fase_id"])
    return schede


TIPI_ATTREZZATURA = ("Macchina", "Attrezzo", "Opera provvisionale", "Impianto")


def _o(valore, dato: str, sezione: str, livello: str = OBBLIGATORIO):
    """Il valore, o il dato mancante da segnalare nel documento."""
    return valore if str(valore or "").strip() and valore != "-" else Manca(dato, sezione, livello)


def _figure_mancanti(repo: Repository, impresa: dict) -> list[tuple[str, Manca]]:
    """Figure richieste dal D.Lgs. 81/2008 non indicate né dichiarate non presenti: (ruolo, dato mancante)."""
    elenco = anagrafiche.figure(repo, impresa)
    ruoli = {anagrafiche.ruolo_riconosciuto(f["ruolo"]) for f in elenco if f["nome"].strip()}
    ruoli |= {r for r, _ in anagrafiche.figure_non_presenti(impresa)}
    mancanti = []
    for richiesti, _, grave in anagrafiche.FIGURE_OBBLIGATORIE:
        if ruoli & set(richiesti):
            continue
        ruolo = "RLS o RLST" if richiesti[0] == "RLS" else richiesti[0]
        mancanti.append((ruolo, Manca(f"nominativo: {ruolo if ruolo.split()[0].isupper() else ruolo.lower()}"
                                      + ("" if grave else " (o dichiararlo non presente)"),
                                      "Figure della sicurezza", OBBLIGATORIO if grave else CONSIGLIATO)))
    return mancanti


def genera_pos(repo: Repository, impresa_id: int, destinazione: str | Path) -> Path:
    impresa = repo.leggi("imprese", impresa_id)
    if impresa is None:
        raise ValueError("Impresa inesistente")
    fasi = repo.elenco("fasi", impresa_id=impresa_id)
    nomi_fasi = {f["id"]: nome_fase(f) for f in fasi}
    lavoratori = repo.elenco("lavoratori", impresa_id=impresa_id)
    rischi = [r for f in fasi for r in repo.rischi(fase_id=f["id"])]
    attrezzature = [a for f in fasi for a in repo.elenco("attrezzature", fase_id=f["id"])]
    c = repo.cantiere()
    soggetti = anagrafiche.soggetti(repo)
    committente = next((s for s in soggetti if s["ruolo"] == anagrafiche.COMMITTENTE), None)
    periodo = analisi.periodo_lavori(fasi)
    sezioni = {chiave: repo.testo(chiave_testo(impresa_id, chiave)) for chiave, _ in SEZIONI_POS}
    datore = impresa["datore_lavoro"].strip()
    figure = [f for f in anagrafiche.figure(repo, impresa) if f["nome"].strip()]
    non_presenti = anagrafiche.figure_non_presenti(impresa)
    revisioni = repo.elenco("revisioni", documento=f"POS:{impresa_id}")
    ultima = revisioni[-1] if revisioni else None
    data_documento = data_it(ultima["data"]) if ultima and ultima["data"] else data_it(date.today())
    valutazioni = valutazioni_impresa(repo, impresa_id)
    sostanze_elenco = sostanze_impresa(repo, impresa_id)
    schede = schede_attrezzature(attrezzature)
    ponteggi = repo.elenco("ponteggi", impresa_id=impresa_id)

    doc = Document()
    stile.imposta_stili(doc)
    comune = f"{c['comune']} ({c['provincia']})" if c["provincia"] else c["comune"]
    revisione = f"Rev. {ultima['numero']} del {data_documento}" if ultima else f"Emissione del {data_documento}"
    stile.imposta_pagina(doc, " · ".join(v for v in (c["nome"], comune) if v),
                         (f"Piano operativo di sicurezza · {impresa['ragione_sociale']}".strip(" ·"), revisione))

    # --- copertina ---
    rspp = next((f for f in figure if f["ruolo"] == anagrafiche.RSPP), None)
    firme_copertina = [("Il datore di lavoro", datore)]
    if rspp and rspp["nome"] != datore:
        firme_copertina.append(("Il responsabile del servizio di prevenzione e protezione", rspp["nome"]))
    stile.copertina(
        doc,
        comune=comune,
        titolo="Piano operativo di sicurezza",
        riferimento="D.Lgs. 9 aprile 2008, n. 81 · art. 89, comma 1, lettera h), art. 96 e Allegato XV punto 3.2",
        voci=[("Cantiere", _o(indirizzo_cantiere(c), "indirizzo del cantiere", "Cantiere")),
              ("Opera", _o(c["nome"] or c["descrizione"], "denominazione dell'opera", "Cantiere")),
              ("Committente", _o(committente["nome"] if committente else "", "committente", "Committente")),
              ("Lavori affidati", _o(impresa["tipologia_lavori"], "tipologia dei lavori affidati", "Impresa",
                                     CONSIGLIATO))],
        data_luogo=f"{c['comune'] or ''}, {data_documento}".strip(", "),
        firme=firme_copertina,
        impresa=[impresa["ragione_sociale"], anagrafiche.sede(impresa),
                 " · ".join(v for v in (f"tel. {impresa['telefono']}" if impresa["telefono"] else "",
                                        impresa["pec"] or impresa["email"]) if v)],
        logo=impresa["logo"],
        revisione=revisione,
    )

    # --- indice e revisioni ---
    stile.indice(doc)
    stile.banda(doc, "Revisioni del documento")
    if revisioni:
        stile.tabella_colorata(doc, ["Rev.", "Data", "Descrizione"], [
            [str(r["numero"]), data_it(r["data"]), r["descrizione"]] for r in revisioni], larghezze=[1.2, 2.6, 13.6])
    else:
        doc.add_paragraph(f"Prima emissione del {data_documento}.")

    # ===== Parte 1 =====
    stile.parte(doc, "1", "Il cantiere e l'impresa",
                "Dati del lavoro, soggetti coinvolti, impresa esecutrice, organizzazione della sicurezza e "
                "lavoratori (Allegato XV punto 3.2.1, lettere a) e b)).")

    stile.capitolo(doc, "Dati generali del lavoro", "Allegato XV punto 3.2.1, lettera a), punto 1")
    stile.scheda(doc, [
        ("Opera", _o(c["nome"], "denominazione dell'opera", "Cantiere")),
        ("Natura dell'opera", _o(c["natura_opera"], "natura dell'opera", "Cantiere", CONSIGLIATO)),
        ("Indirizzo del cantiere", _o(indirizzo_cantiere(c), "indirizzo del cantiere", "Cantiere")),
        ("CAP", c["cap"]),
        ("Titolo abilitativo", _o(c["titolo_abilitativo"], "titolo abilitativo", "Cantiere", FACOLTATIVO)),
        ("Importo presunto dei lavori", euro(c["importo_lavori"]) if c["importo_lavori"] else
         Manca("importo presunto dei lavori", "Cantiere", FACOLTATIVO)),
        ("Inizio dei lavori", data_it(periodo[0]) if periodo else (data_it(impresa["data_inizio"])
                                                                  if impresa["data_inizio"] else "")),
        ("Fine dei lavori (presunta)", data_it(periodo[1]) if periodo else (data_it(impresa["data_fine"])
                                                                           if impresa["data_fine"] else "")),
    ])
    if c["descrizione"].strip():
        stile.banda(doc, "Descrizione sintetica dell'opera")
        stile.testo(doc, c["descrizione"])
    for n, foto in enumerate(foto_cantiere(repo), 1):
        stile.immagine(doc, foto, 11, f"Foto {n} - area del cantiere")

    stile.capitolo(doc, "Committente e soggetti con compiti di sicurezza", "Allegato XV punto 3.2.1, lettera a)")
    soggetti_validi = [s for s in soggetti if s["nome"].strip()]
    if soggetti_validi:
        for s in soggetti_validi:
            stile.banda(doc, s["ruolo"])
            if s.get("nota"):
                stile.nota(doc, s["nota"])
            stile.scheda(doc, [
                ("Nominativo / ragione sociale", s["nome"]), ("Nella persona di", s["rappresentante"]),
                ("Qualifica", s["qualifica"]), ("Indirizzo", s["indirizzo"]),
                ("Telefono", s["telefono"]), ("Email", s["email"]), ("PEC", s["pec"]),
                ("Codice fiscale", s["codice_fiscale"]), ("Partita IVA", s["partita_iva"]),
            ])
    else:
        stile.testo(doc, "", "committente e coordinatori da indicare nella pagina «Committente e CSE»",
                    "Committente", OBBLIGATORIO)

    stile.capitolo(doc, "Impresa esecutrice", "Allegato XV punto 3.2.1, lettera a), punto 1")
    affidataria = repo.leggi("imprese", impresa["affidataria_id"]) if impresa["affidataria_id"] else None
    stile.scheda(doc, titolo="Dati identificativi", voci=[
        ("Ragione sociale", _o(impresa["ragione_sociale"], "ragione sociale dell'impresa", "Impresa")),
        ("Ruolo nel cantiere", impresa["tipo"]),
        ("Forma giuridica", impresa["forma_giuridica"]),
        ("Datore di lavoro", _o(datore, "datore di lavoro", "Impresa")),
        ("Referente", impresa["referente"]),
        ("Sede legale", _o(anagrafiche.sede(impresa) if impresa["indirizzo"].strip() else "",
                           "indirizzo della sede legale", "Impresa")),
        ("Telefono / fax", _o(" · ".join(v for v in (impresa["telefono"], impresa["fax"]) if v),
                              "telefono dell'impresa", "Impresa", CONSIGLIATO)),
        ("Email / PEC", " · ".join(v for v in (impresa["email"], impresa["pec"]) if v)),
        ("Recapito in cantiere", _o(impresa["recapito_cantiere"], "recapito dell'ufficio di cantiere", "Impresa",
                                    FACOLTATIVO)),
        ("Codice fiscale", impresa["codice_fiscale"] if impresa["codice_fiscale"] or impresa["partita_iva"]
         else Manca("codice fiscale o partita IVA dell'impresa", "Impresa", CONSIGLIATO)),
        ("Partita IVA", impresa["partita_iva"]),
        ("Registro imprese / CCIAA", impresa["cciaa"]),
        ("Codice ATECO", impresa["ateco"]),
    ])
    stile.scheda(doc, titolo="Posizioni assicurative, contratto e qualificazione", voci=[
        ("Posizione INAIL", _o(impresa["inail"], "posizione INAIL", "Impresa", CONSIGLIATO)),
        ("Matricola INPS", _o(impresa["inps"], "matricola INPS", "Impresa", CONSIGLIATO)),
        ("Cassa edile", _o(impresa["cassa_edile"], "cassa edile", "Impresa", FACOLTATIVO)),
        ("Contratto collettivo applicato", _o(impresa["ccnl"], "contratto collettivo applicato", "Impresa",
                                              FACOLTATIVO)),
        ("Organico medio annuo", impresa["organico"]),
        ("DURC valido fino al", data_it(impresa["durc_scadenza"]) if impresa["durc_scadenza"] else
         Manca("scadenza del DURC", "Impresa", FACOLTATIVO)),
        ("Patente a crediti (art. 27)", anagrafiche.descrizione_patente(impresa) or
         Manca("patente a crediti dell'impresa", "Impresa", OBBLIGATORIO)),
        ("DURC di congruità", impresa.get("durc_congruita") or ""),
        ("Documento di valutazione dei rischi del", data_it(impresa["dvr_data"]) if impresa["dvr_data"] else ""),
        ("Attestazione SOA", impresa["soa"]),
        ("Certificazioni", impresa["certificazioni"]),
    ])
    stile.scheda(doc, titolo="Lavori affidati", voci=[
        ("Impresa affidataria", affidataria["ragione_sociale"] if affidataria else ""),
        ("Autorizzazione al subappalto", impresa["aut_subappalto"]),
        ("Tipologia dei lavori affidati", impresa["tipologia_lavori"]),
        ("Importo dei lavori da eseguire", euro(impresa["importo_lavori"]) if impresa["importo_lavori"] else ""),
        ("Periodo dei lavori dell'impresa", _periodo(impresa, periodo)),
        ("Orario di lavoro", _o(impresa["orario_lavoro"], "orario di lavoro", "Impresa", CONSIGLIATO)),
        ("Soggetto incaricato dei compiti dell'impresa affidataria (art. 97)",
         " · ".join(v for v in (impresa["incaricato_97"], impresa["incaricato_97_mansione"]) if v)),
    ])
    subaffidatari = repo.elenco("imprese", affidataria_id=impresa_id)
    if subaffidatari:
        stile.banda(doc, "Imprese subappaltatrici e lavoratori autonomi")
        stile.tabella_colorata(doc, ["Ragione sociale", "Ruolo", "Patente a crediti", "Lavori affidati"], [
            [i["ragione_sociale"] + (f"\nP. IVA {i['partita_iva']}" if i["partita_iva"] else ""), i["tipo"],
             anagrafiche.descrizione_patente(i) or "—", i["tipologia_lavori"]] for i in subaffidatari],
            larghezze=[5, 3.2, 3.6, 5.6])
        stile.nota(doc, "L'impresa affidataria verifica l'idoneità tecnico-professionale delle imprese a cui "
                        "subappalta e dei lavoratori autonomi, compresa la patente a crediti (art. 97, comma 2, e "
                        "Allegato XVII).")

    stile.capitolo(doc, "Organizzazione della sicurezza",
                   "Allegato XV punto 3.2.1, lettera a), punti 3-6, e lettera b)")
    doc.add_paragraph("Sono di seguito indicati il datore di lavoro e le figure della sicurezza dell'impresa per il "
                      "presente cantiere, con le specifiche mansioni inerenti alla sicurezza svolte da ciascuno. "
                      "Una stessa persona può ricoprire più ruoli, nei limiti previsti dal D.Lgs. 81/2008.")
    righe_figure = []
    if datore:
        righe_figure.append([anagrafiche.DATORE, datore, " · ".join(v for v in (impresa["telefono"],
                                                                                 impresa["email"]) if v), ""])
    for f in figure:
        formazione = f["formazione"]
        if f.get("formazione_scadenza"):
            formazione = f"{formazione}\naggiornamento entro il {data_it(f['formazione_scadenza'])}".strip()
        righe_figure.append([f["ruolo"], f["nome"] + (f"\n{f['qualifica']}" if f["qualifica"] and
                                                       f["qualifica"] != anagrafiche.DATORE else ""),
                             " · ".join(v for v in (f["telefono"], f.get("email", ""), f.get("codice_fiscale", ""))
                                        if v), formazione])
    if righe_figure:
        stile.tabella_colorata(doc, ["Ruolo", "Nominativo", "Recapiti", "Formazione"], righe_figure,
                               larghezze=[4.2, 4.2, 4.4, 4.6])
    for ruolo, motivazione in non_presenti:
        stile.nota(doc, f"{ruolo}: figura non presente. {motivazione}")
    for ruolo, manca in _figure_mancanti(repo, impresa):
        p = doc.add_paragraph()
        stile._run(p, f"{ruolo}: ", grassetto=True)
        segnaposto(doc, p, manca)
    stile.sezione(doc, "Specifiche mansioni inerenti la sicurezza")
    ruoli = ([anagrafiche.DATORE] if datore else []) + list(dict.fromkeys(f["ruolo"] for f in figure))
    for ruolo in ruoli:
        nomi = [datore] if ruolo == anagrafiche.DATORE else list(dict.fromkeys(
            f["nome"] for f in figure if f["ruolo"] == ruolo))
        p = doc.add_paragraph()
        stile._spaziatura(p, 6, 1)
        p.paragraph_format.keep_with_next = True
        stile._run(p, ruolo, grassetto=True)
        stile._run(p, f" · {', '.join(nomi)}", colore=stile.TESTO2)
        p = doc.add_paragraph(anagrafiche.compiti(ruolo) or "Compiti definiti nell'incarico ricevuto.")
        p.paragraph_format.left_indent = Cm(0.4)
        for nota_ruolo in dict.fromkeys(f["nota"] for f in figure if f["ruolo"] == ruolo and f.get("nota")):
            q = doc.add_paragraph()
            q.paragraph_format.left_indent = Cm(0.4)
            stile._run(q, nota_ruolo, 8.5, stile.TESTO2, corsivo=True)

    stile.capitolo(doc, "Lavoratori dipendenti e autonomi", "Allegato XV punto 3.2.1, lettera a), punto 7")
    dipendenti = [l for l in lavoratori if not l["autonomo"]]
    autonomi = [l for l in lavoratori if l["autonomo"]]
    if dipendenti:
        stile.banda(doc, f"Lavoratori dipendenti: {len(dipendenti)}")
        stile.tabella_colorata(doc, ["N.", "Cognome e nome", "Qualifica", "Matricola", "Assunto il",
                                     "Idoneità fino al"], [
            [str(n), anagrafiche.nominativo(l), l["mansione"], l["matricola"],
             data_it(l["data_assunzione"]) if l["data_assunzione"] else "",
             data_it(l["idoneita_scadenza"]) if l["idoneita_scadenza"] else ""]
            for n, l in enumerate(dipendenti, 1)], larghezze=[0.9, 4.6, 3.6, 2.4, 2.4, 2.6])
    if autonomi:
        stile.banda(doc, f"Lavoratori autonomi: {len(autonomi)}")
        stile.tabella_colorata(doc, ["N.", "Nominativo", "Qualifica", "C.F. / P. IVA", "Sede", "Periodo"], [
            [str(n), anagrafiche.nominativo(l), l["mansione"], l["partita_iva"] or l["codice_fiscale"],
             l["indirizzo"], " - ".join(data_it(d) for d in (l["data_inizio"], l["data_fine"]) if d)]
            for n, l in enumerate(autonomi, 1)], larghezze=[0.9, 3.6, 3, 3, 3.6, 3.3])
    if not lavoratori:
        stile.testo(doc, "", "numero e qualifiche dei lavoratori, da indicare nella scheda «Maestranze» "
                             "dell'impresa", "Lavoratori", OBBLIGATORIO)

    # ===== Parte 2 =====
    stile.parte(doc, "2", "Attività e valutazione dei rischi",
                "Lavorazioni svolte, modalità organizzative, macchine e opere provvisionali, sostanze, rumore "
                "e rischi specifici (Allegato XV punto 3.2.1, lettera a), punto 2, e lettere c), d), e) e f)).")

    stile.capitolo(doc, "Attività di cantiere, modalità organizzative e turni di lavoro",
                   "Allegato XV punto 3.2.1, lettera a), punto 2, e lettera c)")
    stile.testo(doc, sezioni["attivita"], "descrizione delle attività, delle modalità organizzative e dei turni "
                                          "(«Testi del POS» dell'impresa)", "Testi del POS", OBBLIGATORIO)
    if fasi:
        stile.banda(doc, "Lavorazioni e fasi svolte dall'impresa")
        righe, corrente = [], object()
        for f in fasi:
            righe.append([" " if f["lavorazione"] == corrente else (f["lavorazione"] or "—"), nome_fase(f),
                          f["addetto"]])
            corrente = f["lavorazione"]
        stile.tabella_colorata(doc, ["Lavorazione", "Fase", "Addetto"], righe, larghezze=[5.4, 7, 5])
    else:
        stile.testo(doc, "", "lavorazioni svolte dall'impresa (nessuna lavorazione assegnata)", "Lavorazioni",
                    OBBLIGATORIO)

    lavorazione_corrente = object()
    mansioni = {l["id"]: repo.mansioni(l["id"]) for l in lavoratori}
    nomi_attrezzature = {a["id"]: a["nome"] for a in attrezzature}
    for f in fasi:
        if f["lavorazione"] != lavorazione_corrente:
            lavorazione_corrente = f["lavorazione"]
            gruppo = [g for g in fasi if g["lavorazione"] == lavorazione_corrente]
            stile.riquadro(doc, lavorazione_corrente or "Altre fasi di lavoro", "lavorazione")
            if lavorazione_corrente:
                p = doc.add_paragraph()
                stile._run(p, "La lavorazione è suddivisa nelle seguenti fasi: ", grassetto=True)
                p.add_run("; ".join(nome_fase(g) for g in gruppo) + ".")
        h = doc.add_heading(level=3)
        stile._run(h, nome_fase(f))
        stile._run(h, " (fase)" if f["lavorazione"] else "", 8.5, stile.TESTO3, carattere=stile.CARATTERE)
        if f["dettaglio"]:
            stile.testo(doc, f["dettaglio"])
        struttura = struttura_fase(repo, f["id"])
        impegnati = []
        for l in repo.lavoratori_fase(f["id"]):
            usa = [nomi_attrezzature[a] for fase, a in mansioni.get(l["id"], ()) if fase == f["id"] and a
                   in nomi_attrezzature]
            voce = f"{anagrafiche.nominativo(l)} ({l['mansione']})" if l["mansione"] else anagrafiche.nominativo(l)
            impegnati.append(voce + (f", usa: {', '.join(usa)}" if usa else ""))
        stile.scheda(doc, [
            ("Addetto", f["addetto"]),
            ("Lavoratori impegnati", "\n".join(impegnati)),
            ("Macchine utilizzate", ", ".join(a["nome"] for a in struttura["macchine"])),
            ("Attrezzi e opere provvisionali", ", ".join(a["nome"] for a in struttura["attrezzi"])),
            ("Modalità organizzative", f["modalita"]),
            ("Turni di lavoro", f["turni"]),
            ("Procedure complementari richieste dal PSC", f["procedure_psc"]),
        ], 5.0)
        rischi_fase = [r for r in rischi if r["fase_id"] == f["id"]]
        if rischi_fase:
            tabella_rischi(doc, rischi_fase)
            collegate = [v for v in valutazioni if f["id"] in repo.fasi_valutazione(v["id"]) and v["metodo"] in METODI]
            if collegate:
                stile.nota(doc, "Rischi valutati con metodo specifico (capitolo «Valutazioni dei rischi specifici»): "
                                + "; ".join(f"{v['titolo'] or METODI[v['metodo']].nome} - {_esito_breve(repo, v)[1]}"
                                            for v in collegate) + ".")
        else:
            stile.testo(doc, "", f"rischi della fase {nome_fase(f)} (nessun rischio valutato)", "Rischi",
                        OBBLIGATORIO)
        dpi_con_icone(doc, [r["dpi"] for r in struttura["rischi_lavorazione"]], "DPI dell'addetto:")
        esclusi = repo.rischi_esclusi(fase_id=f["id"])
        if esclusi:
            stile.banda(doc, "Rischi valutati come non presenti")
            stile.tabella_colorata(doc, ["Rischio", "Motivazione"], [
                [descrizione_rischio(r), r["motivazione"] or "-"] for r in esclusi], 8.5, larghezze=[6, 11.4])

    stile.capitolo(doc, "Opere provvisionali, macchine e impianti utilizzati",
                   "Allegato XV punto 3.2.1, lettera d), e artt. 70-73 e Allegati V-VII")
    if schede or ponteggi:
        ordine = list(TIPI_ATTREZZATURA)
        righe = [[s["tipo"], s["nome"], "; ".join(nomi_fasi[x] for x in s["fasi"] if x in nomi_fasi)]
                 for s in sorted(schede.values(), key=lambda s: (ordine.index(s["tipo"]) if s["tipo"] in ordine
                                                                else len(ordine), s["nome"].lower()))]
        righe += [["Ponteggio", p_["nome"] or p_["tipo"] or "Ponteggio",
                   " · ".join(v for v in (p_["tipo"], f"altezza {p_['altezza']} m" if p_["altezza"] else "") if v)]
                  for p_ in ponteggi]
        stile.tabella_colorata(doc, ["Tipo", "Denominazione", "Lavorazioni in cui è utilizzata"], righe,
                               larghezze=[3.2, 5.6, 8.6])
        stile.nota(doc, "Macchine, attrezzature e opere provvisionali sono a uso dell'impresa esecutrice, salvo "
                        "diversa indicazione nel PSC o negli accordi con l'impresa affidataria.")
    else:
        stile.testo(doc, "", "macchine, attrezzature e opere provvisionali usate (nessuna indicata nelle "
                             "lavorazioni)", "Lavorazioni", CONSIGLIATO)
    for macchine, titolo in ((True, "Macchine utilizzate nelle lavorazioni"),
                             (False, "Attrezzature e opere provvisionali utilizzate nelle lavorazioni")):
        elenco_schede = [s for s in sorted(schede.values(), key=lambda s: s["nome"])
                         if (s["tipo"] == MACCHINA) == macchine]
        if not elenco_schede:
            continue
        stile.sezione(doc, titolo)
        for s in elenco_schede:
            stile.riquadro(doc, s["nome"], s["tipo"].lower(), stile.TESTO3, stile.SUPERFICIE, livello=3)
            for riga in s["descrizione"]:
                doc.add_paragraph(riga)
            propri = list(dict.fromkeys(r["descrizione"] for r in rischi if r["fonte"] == s["nome"]))
            if propri:
                p = doc.add_paragraph()
                stile._run(p, "Rischi generati dall'uso: ", grassetto=True)
                p.add_run("; ".join(propri) + ".")
            if s["misure"]:
                p = doc.add_paragraph()
                stile._spaziatura(p, 4, 2)
                stile._run(p, "Misure preventive e protettive", grassetto=True)
                stile.testo(doc, "\n".join(s["misure"]))
            if s["dpi"]:
                dpi_con_icone(doc, s["dpi"], "DPI dell'operatore:" if macchine else "DPI dell'utilizzatore:")
    for p_ in ponteggi:
        stile.riquadro(doc, p_["nome"] or "Ponteggio", "ponteggio", stile.TESTO3, stile.SUPERFICIE, livello=3)
        stile.scheda(doc, [("Tipo", p_["tipo"]), ("Fabbricante e modello", " ".join(v for v in (p_["fabbricante"],
                                                                                           p_["modello"]) if v)),
                           ("Autorizzazione ministeriale", p_["autorizzazione"]),
                           ("Altezza / lunghezza", " · ".join(v for v in (f"{p_['altezza']} m" if p_["altezza"] else "",
                                                                         f"{p_['lunghezza']} m" if p_["lunghezza"]
                                                                         else "") if v)),
                           ("Classe di carico", p_["classe_carico"]),
                           ("Montato secondo gli schemi tipo", "sì" if p_["conforme_schemi"] else "no, con progetto"),
                           ("Piano di appoggio", p_["piano_appoggio"]), ("Ancoraggi", p_["ancoraggi"])])
        stile.nota(doc, "Il montaggio, la trasformazione e lo smontaggio sono eseguiti secondo il PiMUS (art. 136).")

    stile.capitolo(doc, "Sostanze e miscele pericolose", "Allegato XV punto 3.2.1, lettera e), e Titolo IX")
    aggiungi_sostanze(doc, repo, sostanze_elenco, livello=3)
    if sezioni["sostanze"].strip():
        stile.testo(doc, sezioni["sostanze"])

    for metodo, titolo in (("rumore", dict(SEZIONI_POS)["rumore"]),
                           ("vibrazioni", "Esito della valutazione delle vibrazioni")):
        elenco_v = [v for v in valutazioni if v["metodo"] == metodo]
        if metodo == "vibrazioni" and not elenco_v:
            continue
        stile.capitolo(doc, titolo, "Allegato XV punto 3.2.1, lettera f), e Titolo VIII, Capo II"
                       if metodo == "rumore" else "Titolo VIII, Capo III del D.Lgs. 81/2008")
        if elenco_v:
            stile.tabella_colorata(doc, ["Mansione / gruppo omogeneo", "Livello di esposizione", "Esito"], [
                [v["mansione"] or v["titolo"], *_esito_breve(repo, v)] for v in elenco_v], larghezze=[6, 4, 7.4])
            stile.nota(doc, "I calcoli e le misure di dettaglio sono riportati nel capitolo «Valutazioni dei rischi "
                                "specifici».")
        if metodo == "rumore" and (sezioni["rumore"].strip() or not elenco_v):
            stile.testo(doc, sezioni["rumore"], "esito della valutazione del rumore (art. 190), oppure la "
                                                "valutazione nella pagina «Valutazioni»", "Testi del POS",
                        OBBLIGATORIO)

    eseguibili = [v for v in valutazioni if v["metodo"] in METODI]
    if eseguibili:
        stile.capitolo(doc, "Valutazioni dei rischi specifici", "Titoli VI, VIII e IX del D.Lgs. 81/2008")
        stile.tabella_colorata(doc, ["Valutazione", "Lavorazioni", "Metodo", "Indice", "Classe"], [
            [v["titolo"] or METODI[v["metodo"]].nome, lavorazioni_valutazione(repo, v) or "-",
             METODI[v["metodo"]].norma, *_esito_breve(repo, v)] for v in eseguibili],
            8, larghezze=[3.6, 4.6, 3.4, 2.2, 3.6])
        for v in eseguibili:
            stile.riquadro(doc, v["titolo"] or METODI[v["metodo"]].nome, METODI[v["metodo"]].nome)
            aggiungi_valutazione(doc, repo, v, livello=2, titolo=False)

    # ===== Parte 3 =====
    stile.parte(doc, "3", "Misure, protezione ed emergenze",
                "Misure integrative al PSC, procedure, dispositivi di protezione, formazione, emergenze, "
                "segnaletica e documenti (Allegato XV punto 3.2.1, lettere g), h), i) e l)).")

    stile.capitolo(doc, dict(SEZIONI_POS)["integrative"], "Allegato XV punto 3.2.1, lettera g)")
    stile.testo(doc, sezioni["integrative"], "misure integrative rispetto a quelle del PSC, se previste",
                "Testi del POS", CONSIGLIATO)
    stile.nota(doc, "Le misure di prevenzione specifiche di ciascuna fase sono riportate nelle tabelle dei rischi "
                    "del capitolo «Attività di cantiere, modalità organizzative e turni di lavoro».")

    stile.capitolo(doc, dict(SEZIONI_POS)["procedure"], "Allegato XV punto 3.2.1, lettera h)")
    stile.testo(doc, sezioni["procedure"], "procedure complementari e di dettaglio richieste dal PSC, se previste",
                "Testi del POS", CONSIGLIATO)
    con_procedure = [f for f in fasi if f["procedure_psc"].strip()]
    if con_procedure:
        stile.tabella_colorata(doc, ["Fase", "Procedura richiesta dal PSC"],
                               [[nome_fase(f), f["procedure_psc"]] for f in con_procedure], larghezze=[6, 11.4])

    stile.capitolo(doc, "Dispositivi di protezione individuale forniti ai lavoratori",
                   "Allegato XV punto 3.2.1, lettera i), e artt. 74-79")
    con_dpi = [l for l in lavoratori if anagrafiche.dpi_lavoratore(l)]
    if con_dpi:
        stile.tabella_colorata(doc, ["Lavoratore", "Qualifica", "DPI forniti"], [
            [anagrafiche.nominativo(l), l["mansione"], "; ".join(anagrafiche.dpi_lavoratore(l))] for l in con_dpi],
            larghezze=[4.4, 3.6, 9.4])
    dpi = elenco_dpi([r["dpi"] for r in rischi] + [a["dpi"] for a in attrezzature])
    if dpi:
        stile.banda(doc, "DPI previsti dalle schede di rischio delle lavorazioni e delle attrezzature")
        stile.griglia_immagini(doc, [
            (segnaletica.SEGNALI[i].file, segnaletica.SEGNALI[i].testo.replace(" obbligatori", "")
             .replace(" obbligatoria", "")) for i in segnaletica.icone_dpi(dpi)], colonne=6, lato_cm=1.8,
            dimensione_testo=6.5)
        stile.elenco_colonne(doc, dpi)
    elif not con_dpi:
        stile.testo(doc, "", "DPI forniti ai lavoratori, da indicare per i lavoratori o nelle schede di rischio",
                    "Lavoratori", OBBLIGATORIO)

    stile.capitolo(doc, "Informazione, formazione e addestramento", "Allegato XV punto 3.2.1, lettera l), e "
                                                                    "artt. 36, 37 e 77")
    doc.add_paragraph("I lavoratori ricevono informazione, formazione e addestramento ai sensi degli artt. 36, 37 e "
                      "77 del D.Lgs. 81/2008 e dell'Accordo Stato-Regioni vigente; la documentazione è conservata in "
                      "cantiere.")
    if dipendenti:
        stile.tabella_colorata(doc, ["Lavoratore", "Qualifica", "Informazione, formazione e addestramento"], [
            [anagrafiche.nominativo(l), l["mansione"], "; ".join(anagrafiche.formazione_lavoratore(l))
             or Manca(f"formazione di {anagrafiche.nominativo(l) or 'un lavoratore'}", "Lavoratori", CONSIGLIATO)]
            for l in dipendenti], larghezze=[4.4, 3.6, 9.4])

    stile.capitolo(doc, "Primo soccorso, antincendio ed evacuazione",
                   "artt. 18, 43-46 e Allegato XV punto 3.2.1, lettera a), punto 3")
    addetti = [f for f in figure if anagrafiche.ruolo_riconosciuto(f["ruolo"]) in
               (anagrafiche.PRIMO_SOCCORSO, anagrafiche.ANTINCENDIO) or f["ruolo"].startswith("Addetto")]
    if addetti:
        per_ruolo: dict[str, list[str]] = {}
        for f in addetti:
            per_ruolo.setdefault(f["ruolo"], []).append(f["nome"])
        stile.scheda(doc, [(r, ", ".join(dict.fromkeys(n))) for r, n in per_ruolo.items()])
    stile.banda(doc, "Numeri di emergenza")
    stile.scheda(doc, [(nome, numero) for nome, numero in NUMERI_EMERGENZA])
    stile.banda(doc, "Telefoni e indirizzi utili locali")
    stile.testo(doc, sezioni["telefoni"], "telefoni e indirizzi utili locali (pronto soccorso, Vigili del fuoco, "
                                          "ASL...)", "Testi del POS", CONSIGLIATO)
    stile.banda(doc, "Gestione delle emergenze")
    stile.testo(doc, sezioni["emergenze"], "gestione delle emergenze: primo soccorso, antincendio ed evacuazione",
                "Testi del POS", OBBLIGATORIO)
    stile.griglia_immagini(doc, [(segnaletica.SEGNALI[k].file, segnaletica.SEGNALI[k].testo)
                                 for k in ("soccorso", "estintore")], colonne=4, lato_cm=2.0)

    stile.capitolo(doc, "Segnaletica prevista nel cantiere", "Titolo V e Allegati XXIV-XXXII del D.Lgs. 81/2008")
    stile.cartello_testo(doc, "Vietato l'accesso ai non addetti ai lavori")
    cartelli = segnaletica.segnali_cantiere([r["descrizione"] for r in rischi],
                                            [r["dpi"] for r in rischi] + [a["dpi"] for a in attrezzature])
    stile.griglia_immagini(doc, [(s.file, s.testo) for s in cartelli], colonne=5, lato_cm=2.3)
    stile.nota(doc, "Rappresentazione indicativa: i cartelli installati in cantiere sono conformi alla norma UNI EN "
                    "ISO 7010 e collocati in prossimità dei pericoli e degli accessi.")

    stile.capitolo(doc, "Documentazione da custodire in cantiere")
    doc.add_paragraph("Presso il cantiere o gli uffici di cantiere l'impresa tiene a disposizione degli organi di "
                      "vigilanza e del coordinatore per l'esecuzione la seguente documentazione:")
    stile.elenco(doc, DOCUMENTI_IN_CANTIERE)

    stile.capitolo(doc, "Consegna, verifica e aggiornamento del piano", "artt. 100, comma 4, 101, comma 3, e 97")
    stile.tabella_colorata(doc, ["Passaggio", "Data"], [
        ["Copia messa a disposizione del rappresentante dei lavoratori per la sicurezza (almeno 10 giorni prima "
         "dell'inizio dei lavori)", _data_o_spazio(impresa.get("pos_rls_data"))],
        ["Trasmissione all'impresa affidataria" + (f" ({affidataria['ragione_sociale']})" if affidataria
                                                   else "") + " per la verifica di congruenza",
         _data_o_spazio(impresa.get("pos_affidataria_data"))],
        ["Trasmissione al coordinatore per l'esecuzione per la verifica di idoneità",
         _data_o_spazio(impresa.get("pos_cse_data"))],
        ["Esito positivo delle verifiche (i lavori iniziano solo dopo)", _data_o_spazio(impresa.get("pos_esito_data"))],
    ], 8.5, larghezze=[13.4, 4.0])
    doc.add_paragraph("Il piano è aggiornato, datato, firmato e trasmesso di nuovo all'impresa affidataria o al "
                      "coordinatore per l'esecuzione a ogni variazione rilevante: nuove lavorazioni, attrezzature o "
                      "sostanze, cambi della squadra o delle figure della sicurezza, modifiche del PSC.")

    # --- allegati e firme ---
    stile.capitolo(doc, "Allegati e firme")
    allegati = []
    if sostanze_elenco:
        schede_sds = [Path(s["file_sds"]).name for s in sostanze_elenco if s.get("file_sds")]
        allegati.append("schede di dati di sicurezza delle sostanze e miscele pericolose"
                        + (f" ({', '.join(schede_sds)})" if schede_sds else "") + ";")
    if lavoratori:
        allegati += ["attestati di informazione, formazione e addestramento dei lavoratori;",
                     "verbali di consegna dei DPI;"]
    if ponteggi:
        allegati.append("PiMUS dei ponteggi;")
    if allegati:
        allegati[-1] = allegati[-1][:-1] + "."
        doc.add_paragraph("Sono allegati al presente piano operativo di sicurezza, e ne fanno parte integrante:")
        stile.elenco_lettere(doc, allegati)
    doc.add_paragraph("Il datore di lavoro dichiara di aver redatto il presente piano operativo di sicurezza per il "
                      "cantiere indicato, ai sensi dell'art. 96 e dell'Allegato XV del D.Lgs. 81/2008, e di averlo "
                      "messo a disposizione del rappresentante dei lavoratori per la sicurezza"
                      + (f" il {data_it(impresa['pos_rls_data'])}" if impresa.get("pos_rls_data") else "")
                      + ", almeno dieci giorni prima dell'inizio dei lavori (art. 100, comma 4).")
    p = doc.add_paragraph(f"{c['comune'] or ''}, {data_documento}".strip(", "))
    stile._spaziatura(p, 8, 8)
    rls = next((f for f in figure if anagrafiche.ruolo_riconosciuto(f["ruolo"]) in ("RLS", "RLST")), None)
    medico = next((f for f in figure if anagrafiche.ruolo_riconosciuto(f["ruolo"]) == anagrafiche.MEDICO), None)
    firme = [("Il datore di lavoro", datore)]
    if rspp and rspp["nome"] != datore:
        firme.append(("Il responsabile del servizio di prevenzione e protezione", rspp["nome"]))
    if medico:
        firme.append(("Il medico competente", medico["nome"]))
    firme.append(("Per presa visione, il rappresentante dei lavoratori per la sicurezza"
                  + (" territoriale" if rls and anagrafiche.ruolo_riconosciuto(rls["ruolo"]) == "RLST" else ""),
                  rls["nome"] if rls else ""))
    stile.firme(doc, firme)

    return salva(doc, destinazione)
