"""Generazione del PiMUS (art. 136 e Allegato XXII del D.Lgs. 81/2008).

Stessa impaginazione del PSC e del POS, nel colore del programma: copertina, capitoli numerati secondo
l'Allegato XXII, dati mancanti evidenziati.
"""

from __future__ import annotations

from datetime import date
from pathlib import Path

from docx import Document

from sicurcantiere import anagrafiche, analisi, pimus
from sicurcantiere import elementi_ponteggio as E
from sicurcantiere import ponteggio as P
from sicurcantiere.documenti import stile
from sicurcantiere.documenti.comune import data_it, indirizzo_cantiere, salva
from sicurcantiere.documenti.composizione import CONSIGLIATO, OBBLIGATORIO, Manca
from sicurcantiere.repository import Repository

CANTIERE, IMPRESA, PONTEGGI = "Cantiere", "Impresa", "Ponteggi"


def documento_pimus(ponteggio_id: int) -> str:
    """Chiave del PiMUS di un ponteggio nelle revisioni e nei documenti composti."""
    return f"PIMUS:{ponteggio_id}"


def genera_pimus(repo: Repository, ponteggio_id: int, destinazione: str | Path) -> Path:
    with stile.colori_documento(*stile.COLORI_PIMUS):
        return _genera(repo, ponteggio_id, destinazione)


def _misura(valore: float) -> str:
    return f"{valore:g} m".replace(".", ",") if valore else ""


def _classe(livello: int) -> int:
    """Classe colorata del livello P×D: 0 basso, 1 medio, 2 alto, 3 molto alto."""
    return 0 if livello <= 2 else 1 if livello <= 4 else 2 if livello <= 8 else 3


def _data(valore) -> str:
    return data_it(valore) if valore else ""


def _genera(repo: Repository, ponteggio_id: int, destinazione: str | Path) -> Path:
    p = repo.leggi("ponteggi", ponteggio_id)
    if p is None:
        raise ValueError("Ponteggio inesistente")
    c = repo.cantiere()
    impresa = (repo.leggi("imprese", p["impresa_id"]) if p["impresa_id"] else None) or {}
    testi = {chiave: repo.testo(pimus.chiave_testo(ponteggio_id, chiave)) for chiave, _, _ in pimus.SEZIONI}
    titoli = {chiave: titolo for chiave, titolo, _ in pimus.SEZIONI}
    comune = f"{c['comune']} ({c['provincia']})" if c["provincia"] else c["comune"]
    revisioni = repo.elenco("revisioni", documento=documento_pimus(ponteggio_id))
    ultima = revisioni[-1] if revisioni else None
    data_documento = data_it(ultima["data"]) if ultima and ultima["data"] else data_it(date.today())
    revisione = f"Rev. {ultima['numero']} del {data_documento}" if ultima else data_documento
    g = P.geometria(p)
    lati = P.lati(p, g)
    altezza = max(p["altezza"], P.altezza_massima(p, g))
    progetto_richiesto = altezza > pimus.ALTEZZA_SENZA_PROGETTO or not p["conforme_schemi"]
    addetti = [a for a in repo.elenco("addetti_ponteggio", ponteggio_id=ponteggio_id) if a["nome"].strip()]
    preposti = [a["nome"] for a in addetti if a["ruolo"] == P.PREPOSTO]
    preposto = ", ".join(preposti) or p["preposto"]
    allegati = repo.elenco("allegati_ponteggio", ponteggio_id=ponteggio_id)

    doc = Document()
    stile.imposta_stili(doc)
    stile.imposta_pagina(doc, " · ".join(v for v in (c["nome"], comune) if v),
                         (f"PiMUS · {p['nome']}".strip(" ·"), revisione))
    stile.copertina(
        doc,
        comune=comune,
        titolo="Piano di montaggio, uso e smontaggio del ponteggio",
        riferimento="D.Lgs. 9 aprile 2008, n. 81 · art. 136 e Allegato XXII",
        voci=[("Cantiere", indirizzo_cantiere(c) or Manca("indirizzo del cantiere", CANTIERE)),
              ("Ponteggio", p["nome"] or Manca("ubicazione del ponteggio", PONTEGGI)),
              ("Tipologia", ", ".join(P.tipologie(p, g)))],
        data_luogo=f"{c['comune'] or ''}, {data_documento}".strip(", "),
        firme=_firme(impresa, preposto, p),
        impresa=[impresa.get("ragione_sociale") or c["nome"] or "Ponteggio", anagrafiche.sede(impresa)
                 if impresa else ""],
        logo=impresa.get("logo", ""),
        revisione=revisione if ultima else "",
    )
    if revisioni:
        stile.banda(doc, "Revisioni del documento")
        stile.tabella_colorata(doc, ["Rev.", "Data", "Descrizione"], [
            [str(r["numero"]), data_it(r["data"]), r["descrizione"]] for r in revisioni], larghezze=[1.2, 2.6, 13.6])

    avvisi = pimus.avvisi(p, addetti)
    if avvisi:
        stile.banda(doc, "Avvertenze")
        for avviso in avvisi:
            stile.nota(doc, avviso, avviso=True)

    stile.capitolo(doc, "Dati identificativi del luogo di lavoro", "Allegato XXII, punto 1", nuova_pagina=True)
    stile.scheda(doc, [
        ("Opera", c["nome"] or Manca("denominazione dell'opera", CANTIERE)),
        ("Indirizzo del cantiere", indirizzo_cantiere(c) or Manca("indirizzo del cantiere", CANTIERE)),
        ("Ubicazione del ponteggio", p["nome"] or Manca("ubicazione del ponteggio", PONTEGGI)),
        ("Opera servita", p.get("opera_servita", "") or Manca("opera servita dal ponteggio", PONTEGGI, CONSIGLIATO)),
        ("Piano di appoggio e terreno", p["piano_appoggio"]),
        ("Condizioni ambientali", p.get("condizioni_ambientali", "")),
        ("Vincoli per gli ancoraggi", p.get("vincoli_ancoraggi", "")),
    ])

    stile.capitolo(doc, "Datore di lavoro che esegue montaggio, trasformazione e smontaggio",
                   "Allegato XXII, punto 2")
    stile.scheda(doc, [
        ("Impresa", impresa.get("ragione_sociale") or Manca("impresa che monta il ponteggio", PONTEGGI)),
        ("Sede", anagrafiche.sede(impresa) if impresa else ""),
        ("Partita IVA", impresa.get("partita_iva", "")),
        ("Datore di lavoro", impresa.get("datore_lavoro") or Manca("datore di lavoro dell'impresa", IMPRESA)),
        ("Impresa responsabile in opera", p.get("impresa_utilizzatrice", "") or "la stessa impresa"),
        ("Persona competente che redige il piano", " · ".join(v for v in (
            p.get("redattore", ""), p.get("redattore_qualifica", "")) if v)
         or Manca("persona competente che redige il PiMUS (art. 136)", PONTEGGI)),
    ])

    stile.capitolo(doc, "Squadra di lavoratori e preposto", "Allegato XXII, punto 3, e art. 136, comma 6")
    if addetti:
        stile.tabella_colorata(
            doc, ["Nome e cognome", "Ruolo", "Corso (data ed ente)", "Ultimo aggiornamento", "Valido fino al"],
            [[a["nome"], a["ruolo"], " · ".join(v for v in (_data(a["corso_data"]), a["corso_ente"]) if v) or "—",
              _data(a["aggiornamento_data"]) or "—",
              data_it(P.scadenza_formazione(a)) if P.scadenza_formazione(a) else "da indicare"] for a in addetti],
            8.5, larghezze=[4.4, 3.4, 4.2, 2.8, 2.6])
        if not preposti:
            stile.testo(doc, "", "preposto alle operazioni", PONTEGGI, OBBLIGATORIO)
    else:
        stile.scheda(doc, [
            ("Preposto alle operazioni", p["preposto"] or Manca("preposto al montaggio", PONTEGGI)),
            ("Addetti al montaggio", p["squadra"] or Manca("addetti al montaggio", PONTEGGI)),
        ])
    doc.add_paragraph("Gli addetti al montaggio, alla trasformazione e allo smontaggio hanno frequentato il corso "
                      "di formazione previsto dall'Allegato XXI (art. 136, comma 6), di 28 ore, con aggiornamento "
                      "di almeno 4 ore ogni 4 anni, e operano sotto la diretta sorveglianza del preposto.")
    esperienze = [a for a in addetti if a["note"].strip()]
    if esperienze:
        stile.scheda(doc, [(a["nome"], a["note"]) for a in esperienze], titolo="Esperienza e note")

    stile.capitolo(doc, "Identificazione del ponteggio", "Allegato XXII, punto 4")
    scadenza = pimus.scadenza_autorizzazione(p)
    stile.scheda(doc, [
        ("Tipologia", ", ".join(P.tipologie(p, g)) or Manca("tipologia del ponteggio", PONTEGGI)),
        ("Fabbricante", p["fabbricante"] or Manca("fabbricante", PONTEGGI)),
        ("Sistema", " · ".join(v for v in (p.get("sistema", ""), p["modello"]) if v)),
        ("Autorizzazione ministeriale", p["autorizzazione"] or Manca("autorizzazione ministeriale (art. 131)",
                                                                     PONTEGGI)),
        ("Data e rinnovo", " · ".join(v for v in (
            f"rilasciata il {data_it(p['aut_data'])}" if p.get("aut_data") else "",
            f"rinnovata il {data_it(p['aut_rinnovo'])}" if p.get("aut_rinnovo") else "",
            f"valida fino al {data_it(scadenza)}" if scadenza else "") if v)),
        ("Schemi tipo utilizzati", p.get("schemi_tipo", "")),
        ("Altezza massima", _misura(altezza) or Manca("altezza massima", PONTEGGI)),
        ("Sviluppo in pianta", _misura(P.sviluppo(g))),
        ("Classe di carico degli impalcati", p["classe_carico"] or Manca("classe di carico", PONTEGGI, CONSIGLIATO)),
        ("Interpiano", _misura(g["interpiano"])),
        ("Mantovana parasassi", f"all'impalcato a {_misura(g['mantovana_livello'] * g['interpiano'])}: "
                                + ", ".join(l.nome for l in lati if l.mantovana) if any(l.mantovana for l in lati)
         else "non prevista"),
    ])
    if p.get("conformita_psc"):
        stile.scheda(doc, [("Rispetto al PSC", p["conformita_psc"]),
                           ("Motivazione", p.get("motivazione_psc", "") or (
                               Manca("motivazione della difformità dal PSC", PONTEGGI)
                               if p["conformita_psc"].startswith("Difforme") else ""))])
    stile.sezione(doc, "Tratti del ponteggio")
    doc.add_paragraph("Ogni parte del ponteggio ha un codice, riportato nei disegni, per il montaggio, lo smontaggio "
                      "e le verifiche: tratti T1, T2..., stilate S1, S2... di ogni tratto, impalcati numerati dal "
                      "basso, ancoraggi A<tratto>.<numero>.")
    stile.tabella_colorata(
        doc, ["Tratto", "Tipologia", "Facciata", "Campate", "Impalcati", "Larghezza", "Distanza", "Quota"],
        [[f"T{t.indice + 1} · {t.nome}", t.tipo, _misura(t.lunghezza), f"{t.campate} × {_misura(t.campata)}",
          f"{t.livelli} (h {_misura(t.livelli * t.interpiano)})", _misura(t.larghezza), _misura(t.distanza) or "0 m",
          f"{t.quota:+.2f} m".replace(".", ",")] for t in lati], 8, larghezze=[3.2, 2.8, 1.9, 2.3, 2.4, 1.7, 1.6, 1.5])
    inseriti = g["elementi"]
    if inseriti:
        stile.sezione(doc, "Elementi inseriti e personalizzazioni")
        stile.tabella_colorata(doc, ["Elemento", "Posizione", "Quantità", "Note"],
                               [[E.voce(e["tipo"]).nome, E.descrizione(e, lati[e["tratto"]].nome),
                                 f"{E.quantita(e, lati[e['tratto']].campata, g['interpiano']):g}".replace(".", ","),
                                 e["note"]] for e in inseriti], 8, larghezze=[5.4, 5.6, 1.8, 4.6])
    stile.sezione(doc, "Elementi del ponteggio")
    stile.tabella_colorata(doc, ["Elemento", "Quantità indicativa"],
                           [[nome, f"{q:g} {u}".replace(".", ",")] for nome, q, u in P.elementi(p, g)], 8.5,
                           larghezze=[11.4, 6.0])
    stile.nota(doc, "Quantità stimate dalla geometria: gli elementi effettivi e il loro numero sono quelli degli schemi "
                    "del libretto del fabbricante.")

    stile.capitolo(doc, "Disegno esecutivo", "Allegato XXII, punto 5, e art. 134")
    doc.add_paragraph("Il disegno esecutivo riporta le generalità e la firma del progettista o, quando il calcolo non è "
                      "obbligatorio, della persona competente (art. 136, comma 1), i sovraccarichi massimi per metro "
                      "quadrato di impalcato, gli appoggi e gli ancoraggi.")
    a = P.ancoraggi(p, g)
    s = P.appoggio(p, g)
    stile.scheda(doc, [
        ("Sovraccarico massimo", f"{P._n(s.sovraccarico, 2)} kN/m² per impalcato" + (
            "" if P.sovraccarico(p["classe_carico"]) else " (da confermare con la classe di carico)")),
        ("Tipo di ancoraggio", g["tipo_ancoraggio"]),
        ("Descrizione degli ancoraggi", p["ancoraggi"] or Manca("descrizione degli ancoraggi", PONTEGGI)),
        ("Schema degli ancoraggi", f"{a.numero} ancoraggi, {a.maglia}; {P._n(a.area_per_ancoraggio, 1)} m² di "
                                   f"facciata per ancoraggio"),
        ("Piano di appoggio", p["piano_appoggio"] or Manca("piano di appoggio", PONTEGGI)),
        ("Verifica dell'appoggio", f"carico stimato sul montante più caricato ({s.tratto}) "
                                   f"{P._n(s.carico_montante, 2)} kN; pressione {P._n(s.pressione)} kPa su "
                                   f"{P._n(g['area_appoggio'], 2)} m² di ripartizione, portanza assunta "
                                   f"{P._n(s.portanza)} kPa: "
                                   + ("verificata" if s.verificato else
                                      f"non verificata, ripartizione minima {P._n(s.area_minima, 2)} m²")),
    ])
    codici = P.elenco_ancoraggi(p, g)
    if codici:
        stile.sezione(doc, "Elenco degli ancoraggi")
        stile.tabella_colorata(doc, ["Codice", "Tratto", "Posizione", "Impalcato", "Quota", "Tipo"],
                               [[x["codice"], x["tratto"], x["stilata"], str(x["impalcato"]),
                                 f"{x['quota']:+.2f} m".replace(".", ","), x["tipo"]] for x in codici],
                               7.5, larghezze=[1.6, 3.2, 2.2, 1.7, 1.8, 6.9])
    _schemi(doc, p)
    _cartiglio(doc, c, p, impresa, progetto_richiesto)

    stile.capitolo(doc, "Progetto del ponteggio", "Allegato XXII, punto 6, e artt. 132-133")
    stile.scheda(doc, [("Caso", pimus.caso_progetto(p))])
    if progetto_richiesto:
        doc.add_paragraph("Per il ponteggio è richiesto il progetto (art. 133), comprendente calcolo di resistenza "
                          "e stabilità e disegno esecutivo, firmato da un ingegnere o architetto abilitato e "
                          "allegato al presente piano.")
        stile.scheda(doc, [("Progettista", p["progettista"] or Manca("progettista abilitato", PONTEGGI))])
    else:
        doc.add_paragraph("Il ponteggio è montato secondo gli schemi tipo del libretto del fabbricante, entro i 20 m "
                          "di altezza: il progetto non è richiesto (art. 133).")

    stile.capitolo(doc, "Indicazioni generali per montaggio, trasformazione e smontaggio",
                   "Allegato XXII, punto 7 (piano di applicazione generalizzata)")
    lettere = [("stoccaggio", "a"), ("appoggio", "b"), ("tracciamento", "c"), ("dpi", "d"), ("sollevamento", "e"),
               ("linee_elettriche", "f"), ("ancoraggi", "g"), ("meteo", "h"), ("caduta_materiali", "i")]
    for chiave, lettera in lettere:
        stile.sezione(doc, f"{lettera}) {titoli[chiave]}")
        stile.testo(doc, testi[chiave], titoli[chiave], PONTEGGI, OBBLIGATORIO)
        if chiave == "stoccaggio":
            planimetrie = [x for x in allegati if x["tipo"] == "Planimetria di stoccaggio e montaggio"]
            for x in planimetrie:
                stile.file_allegato(doc, x["file"], "nella pagina «Ponteggi e PiMUS», scheda «Allegati»")
            if not planimetrie:
                stile.testo(doc, "", "planimetria delle zone di stoccaggio e montaggio (scheda «Allegati»)",
                            PONTEGGI, CONSIGLIATO)
        if chiave == "dpi":
            stile.sezione(doc, "d.1) Sistema adottato contro le cadute durante il montaggio e lo smontaggio")
            sistema = p.get("sistema_anticaduta", "")
            if sistema:
                stile.scheda(doc, [("Sistema", sistema)])
                stile.testo(doc, pimus.descrizione_sistema(sistema) or "")
            else:
                stile.testo(doc, "", "sistema adottato per il montaggio e lo smontaggio (punto 7.1)", PONTEGGI,
                            OBBLIGATORIO)
        if chiave == "linee_elettriche" and p.get("linee_tensione"):
            stile.scheda(doc, [("Linea elettrica aerea", p["linee_tensione"]),
                               ("Distanza dal ponteggio", _misura(p.get("linee_distanza") or 0) or "da rilevare"),
                               ("Distanza minima (Allegato IX)", _misura(pimus.distanza_minima(p["linee_tensione"])))])
    for chiave in ("istruzioni", "soccorso"):
        stile.sezione(doc, titoli[chiave])
        stile.testo(doc, testi[chiave], titoli[chiave], PONTEGGI, CONSIGLIATO)
    stile.scheda(doc, [("DPI specifici", p["dpi"]), ("Attrezzature di sollevamento", p["sollevamento"])],
                 titolo="Dotazioni")

    stile.capitolo(doc, "Modalità di montaggio, trasformazione e smontaggio", "Allegato XXII, punto 8")
    for chiave in ("montaggio", "trasformazione", "smontaggio"):
        stile.sezione(doc, titoli[chiave])
        stile.testo(doc, testi[chiave], titoli[chiave], PONTEGGI, OBBLIGATORIO)
    stile.sezione(doc, titoli["interdizione"])
    stile.testo(doc, testi["interdizione"], titoli["interdizione"], PONTEGGI, CONSIGLIATO)
    situazioni = [x for x in (p.get("situazioni") or "").split(";") if x.strip()]
    speciali = [e for e in inseriti if e["note"].strip()]
    if situazioni or speciali or p.get("situazioni_note"):
        stile.sezione(doc, "Schemi particolari e istruzioni particolareggiate")
        if situazioni:
            stile.elenco(doc, situazioni)
        if p.get("situazioni_note"):
            stile.testo(doc, p["situazioni_note"])
        if speciali:
            stile.scheda(doc, [(f"{E.voce(e['tipo']).nome} ({E.descrizione(e, lati[e['tratto']].nome)})", e["note"])
                               for e in speciali], titolo="Istruzioni per gli elementi particolari")
    for x in [x for x in allegati if x["tipo"] == "Foto o schemi delle fasi di montaggio"]:
        stile.banda(doc, x["titolo"] or "Fasi di montaggio")
        stile.file_allegato(doc, x["file"], "nella pagina «Ponteggi e PiMUS», scheda «Allegati»")

    rischi = repo.elenco("rischi_ponteggio", ponteggio_id=ponteggio_id)
    stile.sezione(doc, "Valutazione dei rischi delle operazioni")
    if rischi:
        t = stile.tabella_colorata(
            doc, ["Rischio", "P", "D", "P×D", "Misure di prevenzione e protezione", "DPI"],
            [[r["descrizione"], str(r["probabilita"]), str(r["danno"]),
              str(analisi.livello_rischio(r["probabilita"], r["danno"])), r["misure"], r["dpi"]] for r in rischi],
            8, larghezze=[3.8, 0.7, 0.7, 1.0, 7.4, 3.6])
        for riga, r in zip(t.rows[1:], rischi):
            stile.colora_livello(riga.cells[3], _classe(analisi.livello_rischio(r["probabilita"], r["danno"])))
    else:
        stile.testo(doc, "", "valutazione dei rischi del montaggio", PONTEGGI, CONSIGLIATO)

    stile.capitolo(doc, "Regole da applicare durante l'uso", "Allegato XXII, punto 9")
    stile.testo(doc, testi["uso"], titoli["uso"], PONTEGGI, OBBLIGATORIO)
    for chiave in ("accesso", "manutenzione", "cessione"):
        stile.sezione(doc, titoli[chiave])
        stile.testo(doc, testi[chiave], titoli[chiave], PONTEGGI, CONSIGLIATO)

    stile.capitolo(doc, "Verifiche prima del montaggio e durante l'uso",
                   "Allegato XXII, punto 10, e Allegato XIX (già Circolare MLPS n. 46/2000)")
    verifiche = repo.elenco("verifiche_ponteggio", ponteggio_id=ponteggio_id)
    if verifiche:
        t = stile.tabella_colorata(doc, ["Momento", "Verifica", "Esito", "Note"],
                                   [[v["momento"], v["voce"], v["esito"] or "Da eseguire", v["note"]]
                                    for v in verifiche], 8, larghezze=[3.0, 8.6, 2.4, 3.4])
        for riga, v in zip(t.rows[1:], verifiche):
            stile.colora_livello(riga.cells[2], 0 if v["esito"] in ("Conforme", "Non applicabile")
                                 else 3 if v["esito"] == "Non conforme" else 1)
    else:
        stile.testo(doc, "", "elenco delle verifiche dell'Allegato XIX", PONTEGGI, OBBLIGATORIO)
    stile.nota(doc, "Durante l'uso il preposto controlla il ponteggio a intervalli periodici e dopo violente "
                    "perturbazioni atmosferiche o prolungata interruzione del lavoro (art. 137, comma 2); i controlli "
                    "sono annotati nel registro dei controlli del ponteggio, tenuto in cantiere con il presente "
                    "piano.")

    p_data = doc.add_paragraph(f"{c['comune'] or ''}, {data_documento}".strip(", "))
    stile._spaziatura(p_data, 14, 8)
    stile.firme(doc, _firme(impresa, preposto, p))
    stile.nota(doc, "Il PiMUS è redatto da persona competente e messo a disposizione del preposto e dei lavoratori "
                    "addetti; una copia, con il disegno esecutivo e il libretto del fabbricante, è tenuta in "
                    "cantiere a disposizione dell'organo di vigilanza (art. 136, comma 1).")
    _presa_visione(doc, addetti)
    _allegati(doc, [x for x in allegati if x["tipo"] not in ("Planimetria di stoccaggio e montaggio",
                                                            "Foto o schemi delle fasi di montaggio")])
    return salva(doc, destinazione)


def _firme(impresa, preposto, p) -> list[tuple[str, str]]:
    firme = [("Il datore di lavoro dell'impresa che esegue il montaggio", impresa.get("datore_lavoro", ""))]
    redattore = p.get("redattore", "")
    if redattore and redattore != impresa.get("datore_lavoro", ""):
        firme.append(("La persona competente che ha redatto il piano", redattore))
    firme.append(("Il preposto alle operazioni", preposto))
    return firme


def _cartiglio(doc, c, p, impresa, progetto_richiesto: bool) -> None:
    """Cartiglio del disegno esecutivo, con le firme richieste (Allegato XXII, punto 5)."""
    stile.sezione(doc, "Cartiglio del disegno esecutivo")
    stile.scheda(doc, [
        ("Cantiere", " · ".join(v for v in (c["nome"], indirizzo_cantiere(c)) if v)),
        ("Ponteggio", p["nome"]),
        ("Sistema", " · ".join(v for v in (p["fabbricante"], p.get("sistema", ""), p["autorizzazione"]) if v)),
        ("Sovraccarico ammesso", p["classe_carico"] or "da indicare"),
        ("Data", data_it(date.today())),
    ])
    firme = [("Il progettista" if progetto_richiesto else "La persona competente (art. 136, comma 1)",
              p["progettista"] if progetto_richiesto else p.get("redattore", "")),
             ("Il responsabile del cantiere", "")]
    stile.firme(doc, firme)


def _presa_visione(doc, addetti) -> None:
    """Il piano è messo a disposizione del preposto e dei lavoratori (art. 136, comma 1): presa visione firmata."""
    stile.capitolo(doc, "Presa visione del piano", "art. 136, comma 1", nuova_pagina=True)
    doc.add_paragraph("Il preposto e i lavoratori addetti dichiarano di aver ricevuto e letto il presente piano, il "
                      "disegno esecutivo e le istruzioni del libretto del fabbricante.")
    righe = [[a["nome"], a["ruolo"], "", ""] for a in addetti] + [["", "", "", ""] for _ in range(max(0, 4 - len(addetti)))]
    stile.tabella_colorata(doc, ["Nome e cognome", "Ruolo", "Data", "Firma"], righe, 9,
                           larghezze=[5.4, 4.0, 2.8, 5.2])


# Contenuto di ogni verbale: riferimento, titolo, controlli da spuntare, impegni o note finali.
VERBALI: dict[str, dict] = {
    pimus.FINE_MONTAGGIO: {
        "riferimento": "D.Lgs. 81/2008 · artt. 136 e 137 e Allegato XIX",
        "titolo": "Verbale di fine montaggio e messa in uso del ponteggio",
        "controlli": [
            "Montato secondo il PiMUS, il disegno esecutivo e gli schemi del libretto (o il progetto)",
            "Piano di appoggio, basette e tavole di ripartizione efficienti",
            "Montanti verticali, giunti e collegamenti serrati, diagonali e controventi completi",
            "Ancoraggi presenti nel numero e nelle posizioni del disegno esecutivo",
            "Impalcati completi e bloccati, parapetti e fermapiede su tutti i lati aperti",
            "Scale di accesso e botole funzionanti, distanza dalla facciata non superiore a 20 cm o protetta",
            "Mantovane, reti, teli e protezioni dei passaggi in opera dove previste",
            "Messa a terra eseguita se richiesta dalla valutazione contro le scariche atmosferiche",
            "Cartello con l'impresa, il carico ammesso e il divieto di modifiche esposto",
        ],
        "chiusura": [
            "Il ponteggio è dichiarato idoneo all'uso a partire dalla data del presente verbale.",
            "Durante l'uso il preposto ne verifica lo stato a intervalli periodici e dopo violente perturbazioni "
            "atmosferiche o prolungate interruzioni del lavoro, registrando i controlli (art. 137, comma 2).",
        ],
    },
    pimus.CONSEGNA: {
        "riferimento": "D.Lgs. 81/2008 · art. 136 e Allegato XXII, punto 9",
        "titolo": "Verbale di consegna del ponteggio",
        "controlli": [
            "Montato secondo il PiMUS e il disegno esecutivo, verificato dal preposto",
            "Ancoraggi presenti nel numero e nelle posizioni previste",
            "Parapetti, fermapiede e impalcati completi; botole e scale di accesso funzionanti",
            "Mantovane, reti e protezioni dei passaggi in opera dove previste",
            "Cartello con il carico ammesso esposto",
            "Copia del PiMUS, del disegno esecutivo e del libretto consegnata all'impresa che riceve",
        ],
        "titolo_chiusura": "Impegni dell'impresa che riceve",
        "chiusura": [
            "Usare il ponteggio solo per i carichi ammessi e secondo le regole d'uso del PiMUS.",
            "Non modificare, togliere o spostare elementi, ancoraggi, parapetti o tavole: ogni trasformazione è "
            "richiesta all'impresa di montaggio.",
            "Segnalare subito ogni anomalia e sospendere l'uso dopo eventi meteorologici violenti fino alla "
            "verifica.",
            "Far accedere solo dalle scale interne o dalle torri scala previste.",
        ],
    },
    pimus.TRASFORMAZIONE: {
        "riferimento": "D.Lgs. 81/2008 · artt. 133, 134 e 136",
        "titolo": "Verbale di trasformazione del ponteggio",
        "controlli": [
            "Trasformazione prevista dal PiMUS o PiMUS aggiornato prima di eseguirla",
            "Configurazione finale compresa negli schemi tipo del libretto, oppure progetto aggiornato e firmato",
            "Disegno esecutivo aggiornato e firmato dal responsabile del cantiere",
            "Area sottostante delimitata e accesso impedito ai non addetti durante i lavori",
            "Eseguita da lavoratori formati sotto la diretta sorveglianza del preposto",
            "Ancoraggi, parapetti e impalcati ripristinati e verificati a fine trasformazione",
        ],
        "chiusura": [
            "Al termine della trasformazione il ponteggio è verificato dal preposto prima di essere riutilizzato.",
        ],
    },
    pimus.SMONTAGGIO: {
        "riferimento": "D.Lgs. 81/2008 · art. 136, comma 6, e Allegato XXII, punto 8",
        "titolo": "Verbale di smontaggio del ponteggio",
        "controlli": [
            "Stato del ponteggio verificato prima di iniziare (ancoraggi, giunti, impalcati)",
            "Area sottostante delimitata e accesso impedito ai non addetti",
            "Smontaggio dall'alto verso il basso con la sequenza del PiMUS e i sistemi anticaduta previsti",
            "Ancoraggi rimossi solo man mano, al livello che si sta smontando",
            "Elementi calati con mezzi di sollevamento o passati a mano, mai lanciati",
            "Materiale accatastato, controllato e allontanato; elementi danneggiati separati",
            "Fori degli ancoraggi chiusi e area restituita libera e pulita",
        ],
        "chiusura": [
            "Lo smontaggio è stato eseguito da lavoratori formati sotto la diretta sorveglianza del preposto "
            "(art. 136, comma 6).",
        ],
    },
}


def genera_consegna(repo: Repository, ponteggio_id: int, destinazione: str | Path) -> Path:
    """Verbale di consegna del ponteggio all'impresa che lo utilizza (cessione a terzi, Allegato XXII punto 9)."""
    return genera_verbale(repo, ponteggio_id, destinazione, tipo=pimus.CONSEGNA)


def genera_verbale(repo: Repository, ponteggio_id: int, destinazione: str | Path, verbale_id: int | None = None,
                   tipo: str | None = None) -> Path:
    """Verbale di fine montaggio, consegna, trasformazione o smontaggio, dalla riga del registro dei verbali
    (data, preposto, impresa che riceve, descrizione) oppure in bianco per il tipo indicato."""
    with stile.colori_documento(*stile.COLORI_PIMUS):
        p = repo.leggi("ponteggi", ponteggio_id)
        if p is None:
            raise ValueError("Ponteggio inesistente")
        v = (repo.leggi("verbali_ponteggio", verbale_id) if verbale_id else None) or {}
        tipo = v.get("tipo") or tipo or pimus.FINE_MONTAGGIO
        modello = VERBALI.get(tipo, VERBALI[pimus.FINE_MONTAGGIO])
        c = repo.cantiere()
        impresa = (repo.leggi("imprese", p["impresa_id"]) if p["impresa_id"] else None) or {}
        g = P.geometria(p)
        tratti = P.tratti(p, g)
        data = data_it(v["data"]) if v.get("data") else "____________"
        preposti = [a["nome"] for a in repo.elenco("addetti_ponteggio", ponteggio_id=ponteggio_id)
                    if a["ruolo"] == P.PREPOSTO and a["nome"].strip()]
        preposto = v.get("preposto") or ", ".join(preposti) or p["preposto"]
        vuoto = "______________________________"
        doc = stile.documento_semplice(" · ".join(x for x in (c["nome"], c["comune"]) if x),
                                       (f"{tipo} del ponteggio · {p['nome']}", data_it(v["data"]) if v.get("data")
                                        else data_it(date.today())))
        stile.intestazione_modulo(doc, modello["riferimento"], modello["titolo"], p["nome"])
        righe = [
            ("Cantiere", " · ".join(x for x in (c["nome"], indirizzo_cantiere(c)) if x)),
            ("Ponteggio", f"{p['nome']} · {', '.join(P.tipologie(p, g))}"),
            ("Sistema e autorizzazione", " · ".join(x for x in (p["fabbricante"], p.get("sistema", ""),
                                                               p["autorizzazione"]) if x)),
            ("Impresa che ha montato il ponteggio", impresa.get("ragione_sociale", "")),
            ("Data", data),
            ("Preposto", preposto or vuoto),
        ]
        if tipo == pimus.CONSEGNA:
            righe.append(("Ricevuto da", v.get("ricevente") or p.get("impresa_utilizzatrice", "") or vuoto))
        righe.append(("Tratti", "; ".join(f"T{t.indice + 1} {t.nome} ({_misura(t.lunghezza)}, {t.livelli} impalcati)"
                                          for t in tratti)))
        if tipo != pimus.SMONTAGGIO:
            righe.append(("Carico ammesso sugli impalcati", p["classe_carico"] or vuoto))
        stile.scheda(doc, righe)
        if tipo == pimus.TRASFORMAZIONE or v.get("descrizione"):
            stile.sezione(doc, "Trasformazione eseguita" if tipo == pimus.TRASFORMAZIONE else "Descrizione")
            doc.add_paragraph(v.get("descrizione") or "\n".join([vuoto * 2] * 3))
        stile.sezione(doc, "Controlli")
        stile.tabella_colorata(doc, ["Controllo", "Sì", "No", "Note"], [[x, "", "", ""] for x in modello["controlli"]],
                               9, larghezze=[10.0, 1.2, 1.2, 5.0])
        stile.sezione(doc, modello.get("titolo_chiusura", "Esito"))
        stile.elenco(doc, modello["chiusura"])
        if v.get("note"):
            stile.nota(doc, v["note"])
        p_data = doc.add_paragraph(f"{c['comune'] or ''}, {data}".strip(", "))
        stile._spaziatura(p_data, 14, 8)
        firme = [("Il preposto alle operazioni", preposto or ""),
                 ("Per l'impresa che ha montato il ponteggio", impresa.get("datore_lavoro", ""))]
        if tipo == pimus.CONSEGNA:
            firme.append(("Per l'impresa che riceve", v.get("ricevente") or p.get("impresa_utilizzatrice", "")))
        stile.firme(doc, firme)
        return salva(doc, destinazione)


def genera_registro(repo: Repository, ponteggio_id: int, destinazione: str | Path) -> Path:
    """Registro dei controlli del ponteggio in uso (art. 137, comma 2), con righe in bianco da compilare in
    cantiere e le foto allegate ai controlli."""
    with stile.colori_documento(*stile.COLORI_PIMUS):
        p = repo.leggi("ponteggi", ponteggio_id)
        if p is None:
            raise ValueError("Ponteggio inesistente")
        c = repo.cantiere()
        registro = repo.elenco("registro_ponteggio", ponteggio_id=ponteggio_id)
        doc = stile.documento_semplice(" · ".join(x for x in (c["nome"], c["comune"]) if x),
                                       (f"Registro dei controlli · {p['nome']}", data_it(date.today())))
        stile.intestazione_modulo(doc, "D.Lgs. 81/2008 · art. 137, comma 2, e Allegato XIX",
                                  "Registro dei controlli del ponteggio", p["nome"])
        doc.add_paragraph("Il preposto, a intervalli periodici o dopo violente perturbazioni atmosferiche o prolungata "
                          "interruzione del lavoro, si assicura della verticalità dei montanti, del giusto serraggio "
                          "dei giunti, dell'efficienza degli ancoraggi e dei controventi, curando l'eventuale "
                          "sostituzione o il rinforzo degli elementi inefficienti. Legenda dei controlli:")
        stile.elenco(doc, [f"{n}. {voce}" for n, (_, voce) in enumerate(pimus.CONTROLLI_USO, 1)])
        righe = []
        for r in registro:
            esiti = " ".join(str(n) for n, (campo, _) in enumerate(pimus.CONTROLLI_USO, 1) if r[campo])
            righe.append([data_it(r["data"]) if r["data"] else "", r["motivo"], esiti or "—", r["esito"],
                          r["interventi"], r["eseguita_da"], ""])
        righe += [["", "", "", "", "", "", ""] for _ in range(max(6, 14 - len(righe)))]
        stile.tabella_colorata(doc, ["Data", "Motivo", "Controlli eseguiti", "Esito", "Interventi",
                                     "Preposto", "Firma"], righe, 8, larghezze=[1.9, 3.0, 1.9, 2.4, 3.6, 2.4, 2.2])
        con_foto = [r for r in registro if r["foto"]]
        for r in con_foto:
            stile.banda(doc, f"Foto del controllo del {data_it(r['data']) if r['data'] else '—'} · {r['motivo']}")
            stile.file_allegato(doc, r["foto"], "nella pagina «Ponteggi e PiMUS», scheda «Registro e verbali»")
        return salva(doc, destinazione)


def _schemi(doc, p) -> None:
    """Schema grafico generato dalla geometria: vista d'insieme, pianta, prospetti e sezione tipo."""
    from sicurcantiere import disegno_ponteggio

    try:
        tavole = disegno_ponteggio.tavole(p)
    except Exception:  # noqa: BLE001 - senza Qt il documento si genera lo stesso, senza disegni
        return
    stile.sezione(doc, "Schema grafico del ponteggio")
    stile.nota(doc, "Schema generato dai dati del ponteggio, con gli ancoraggi in rosso: va verificato sul libretto "
                    "del fabbricante e completato nel disegno esecutivo firmato dal responsabile del cantiere.")
    for tavola in tavole:
        stile.immagine_png(doc, tavola.png, stile.LARGHEZZA_UTILE if tavola.orizzontale else 11.5, tavola.titolo)


def _allegati(doc, allegati) -> None:
    if not allegati:
        return
    stile.capitolo(doc, "Allegati", "art. 134 e art. 136, comma 1", nuova_pagina=True)
    stile.tabella_colorata(doc, ["N.", "Tipo", "Titolo"], [[str(n), a["tipo"], a["titolo"]]
                                                          for n, a in enumerate(allegati, 1)], 8.5,
                           larghezze=[1.2, 6.0, 10.2])
    for n, a in enumerate(allegati, 1):
        stile.banda(doc, f"Allegato {n} · {a['tipo']}" + (f" · {a['titolo']}" if a["titolo"] else ""))
        stile.file_allegato(doc, a["file"], "nella pagina «Ponteggi e PiMUS», scheda «Allegati»")
